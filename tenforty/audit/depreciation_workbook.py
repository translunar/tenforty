"""The depreciation audit workbook: `depreciation_audit_<year>.xlsx`.

Emitted beside the forms on every return that computes MACRS depreciation
from a `depreciable_assets` block. Every computed figure is a spreadsheet
formula over cells visible in the workbook; the Tie-outs sheet compares the
formula results against the figures the engine printed on the forms.

This module computes no tax. It transcribes the depreciation resolver's audit
trail (`resolver.audit_trail`) into cells -- inputs and table rates as values,
the arithmetic as formulas -- and owns all layout.

THE MIRROR. No spreadsheet engine runs, at emit or in the tests. Instead the
formula text written to each cell is evaluated here in Python (`evaluate`, a
small evaluator for exactly the functions this workbook writes) over the
written cells, and compared with the engine's own figure for that cell. A
difference of half a cent or more means the workbook would not reproduce the
filing, and the whole emit is refused (`depreciation_audit_unbuildable`).

Two scenario shapes reach that refusal today, both a stated
`prior_depreciation` carrying cents on an asset whose basis ceiling binds.
The engine takes the prior in whole dollars; the workbook's formula reads the
stated cell, so the remaining basis differs by the cents and can round to a
different dollar:
  - ACKNOWLEDGED (`acknowledges_prior_depreciation_as_stated`): the engine
    takes the stated prior, rounded;
  - UNACKNOWLEDGED: the stated prior rounds to the table reconstruction, and
    the engine takes the reconstruction.
State the prior in whole dollars and the return emits. The refusal also fires
when the file cannot be written.

THREE ARITHMETICS, AND WHERE THEY AGREE. The engine rounds the exact decimal
product, half up (`rounding.irs_round_product`): that is the legally correct
figure and the one on the printed form. The mirror does the same arithmetic
-- exact decimal, each operand read as the decimal it was written as -- so
mirror and engine agree wherever the formula is the engine's rule. A real
spreadsheet multiplies in binary floating point and its ROUND works to 15
significant digits; that agrees with the exact figure as long as the exact
product fits in 15 digits. A dollars-and-cents basis times a table rate has
at most seven decimal places, so agreement is guaranteed while the product
is under $100,000,000.

At $100,000,000 and above the engine stays exact and a spreadsheet's ROUND
may display one dollar more. Observed: 1,400,008,872.42 x 0.07219 is
101,066,640.4999998 exactly; the engine gives 101,066,640 and LibreOffice
26.2.2.2 shows 101,066,641. That is the spreadsheet's precision ceiling,
not a defect in the figure, and the return still emits: every row holding
such a product carries a visible note (`PRECISION_NOTE`), so a tie-out a
filer sees FAIL in their spreadsheet is explained on the page where the
discrepancy arises. Below the threshold, agreement with LibreOffice was
checked by probe on 706 formula cells, including products exactly on the
half dollar (25,000 x 0.0197 = 492.50 -> 493); no spreadsheet runs here.

FAIL-CLOSED: the mirror evaluates every formula cell on every sheet. A
formula it cannot read (anything but arithmetic, comparisons, references and
ROUND, MIN, MAX, ABS, SUM, SUMIF, IF) is a
refusal, never a skipped cell. IF is lazy, as a spreadsheet's is: the branch
not taken is read but not computed.

TWO LAYOUTS. A SINGLE-ACTIVITY return (one activity lists assets or states a
depreciation figure) gets three sheets: Tie-outs, Assets, Year-by-year. A
MULTI-ACTIVITY return (two or more such activities, as counted by the
`ConventionInputs` handed in) gets Tie-outs, Mid-quarter test, then an
`Assets <recon key>` / `Years <recon key>` pair for each activity that lists
assets -- the single-activity pair, holding that activity's assets only.
Either ends with 481(a) when the scenario carries a `form_3115:` block.

On a multi-activity return each Form 4562 copy lists one activity, so no
printed copy shows why the return is (or is not) mid-quarter: the 40% test
pools every activity. The Mid-quarter test sheet lays that test out -- each
activity's last-three-months and full-year bases in whole cents, their
totals, and a DECIDING cell that is integer cross-multiplication
(`late * 100 > 40 * total`), never a fraction against a threshold. An
activity that states its depreciation as one figure is a row that says its
placements are not visible. Tie-outs then runs per copy (line 17, line 19,
line 22, each from that activity's own tab), per schedule line, and for the
whole return: every copy's line 22 plus the schedule depreciation line of
every activity with no copy (listed by recon key, with the reason) against
the sum of every activity's printed schedule line.

Three checks beyond the cell-by-cell mirror refuse a multi-activity emit:
the deciding cell, evaluated by the mirror, must equal the boolean the
orchestrator handed every copy; each copy's line 22 from its own tab must
equal that copy's printed line 22; and a trail of two or more activities
handed in without `ConventionInputs` is refused rather than laid out as a
single-activity workbook. The first is a check of this module's PYTHON
MIRROR of the formula: nothing in the emit path evaluates a spreadsheet, so
whether the rendered cell agrees belongs to the oracle tier and the visual
check. A spreadsheet computes the deciding formula in doubles, which is
exact while the pooled last-three-months bases stay under about 2^53 / 100
cents (roughly $900 billion); beyond that the engine's integer answer, not
the cell, is authoritative.

Assets columns: A Activity #, B Activity, C Description, D In service,
E Method, F Life, G Convention, H Rate, I Basis, J Prior accumulated,
K Deduction, L Accumulated after, M 4562 line, N Provenance. (The design
spec's example formula `=ROUND(MIN(G2*F2,G2-H2),0)` predates the leading
Activity # column and the zero floor: the deduction is
`=ROUND(MAX(0,MIN(I2*H2,I2-J2)),0)` in column K.) Activity tie-outs SUMIF on
the numeric Activity #, never on the activity label: a label is user text and
could carry SUMIF wildcards.
"""
import re
from collections.abc import Mapping
from dataclasses import dataclass, field
from datetime import date
from decimal import ROUND_HALF_UP, Context, Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font
from openpyxl.utils import get_column_letter, range_boundaries

from tenforty.attestations import raise_scoped_refusal
from tenforty.forms import f4562 as form_4562
from tenforty.forms.depreciation import macrs
from tenforty.forms.depreciation.resolver import (
    MID_QUARTER_THRESHOLD_PERCENT, MODE_ASSETS_OVERRIDDEN, ActivityAudit,
    ActivityBases, AssetAudit, audit_trail, mid_quarter_applies,
    mid_quarter_bases_by_activity,
)
from tenforty.rounding import irs_round

REFUSAL = "depreciation_audit_unbuildable"
EMITTED_KEY = "depreciation_audit"

TIE_OUTS = "Tie-outs"
ASSETS = "Assets"
YEAR_BY_YEAR = "Year-by-year"
SECTION_481A = "481(a)"
MID_QUARTER_TEST = "Mid-quarter test"


def assets_tab_name(recon_key: str) -> str:
    """One activity's Assets sheet on a multi-activity workbook."""
    return f"Assets {recon_key}"


def years_tab_name(recon_key: str) -> str:
    """One activity's year-by-year sheet on a multi-activity workbook."""
    return f"Years {recon_key}"

# How each Pub 946 Appendix A table is cited beside the rates read from it.
# A literal list on purpose: a table the engine learns to read must be added
# here before the workbook will cite it (tests enumerate the engine's tables
# against this).
_TABLE_CITATIONS: dict[str, str] = {
    "A-1": "Pub 946 Table A-1",
    "A-2": "Pub 946 Table A-2",
    "A-3": "Pub 946 Table A-3",
    "A-4": "Pub 946 Table A-4",
    "A-5": "Pub 946 Table A-5",
    "A-6": "Pub 946 Table A-6",
    "A-7a": "Pub 946 Table A-7a",
}
GRID_TABLES: frozenset[str] = frozenset(_TABLE_CITATIONS)

# Tie-out tolerance, in dollars: half a cent.
TOLERANCE = 0.005
VERDICT_FORMULA = '=IF(ABS(E{row})<0.005,"PASS","FAIL")'
NET_481A_FORMULA = "=SUM(E2:E{last})"

# A product this large can need 16 significant digits; a spreadsheet keeps 15
# (see the module docstring).
SPREADSHEET_EXACT_BELOW = Decimal(100_000_000)
PRECISION_NOTE = (
    "spreadsheet ROUND may show $1 above the engine's exact figure at this "
    "magnitude \u2014 the printed form is exact; see module docstring")

TIE_OUT_HEADERS = (
    "Claim", "Computed", "Printed on form", "Source of printed figure",
    "Difference", "Verdict")
ASSET_HEADERS = (
    "Activity #", "Activity", "Description", "In service", "Method", "Life",
    "Convention", "Rate", "Basis", "Prior accumulated", "Deduction",
    "Accumulated after", "4562 line", "Provenance")
YEAR_HEADERS = (
    "Asset", "Recovery year", "Tax year", "Rate", "Deduction", "Cumulative",
    "Source")
OVERRIDE_HEADERS = {
    "A": "Activity #", "B": "Activity", "I": "Override amount",
    "J": "Restated engine amount", "N": "Provenance"}
SECTION_481A_HEADERS = (
    "Asset", "In service", "Claimed", "Allowable", "Adjustment", "Note",
    "Provenance")

MID_QUARTER_HEADERS = (
    "Activity", "Label", "Last three months (cents)", "Full year (cents)",
    "Note")
STATED_ROW_NOTE = (
    "placements not visible; the return carries "
    "acknowledges_no_personal_property_behind_stated_depreciation")
DECIDING_LABEL = (
    "Mid-quarter convention applies (last three months x 100 > "
    f"{MID_QUARTER_THRESHOLD_PERCENT} x full year)")
# Integer cross-multiplication, never a fraction against a threshold: no
# division, so nothing to round and nothing to divide by when no personal
# property was placed in service at all.
DECIDING_FORMULA = (
    "={late}*100>" + str(MID_QUARTER_THRESHOLD_PERCENT) + "*{total}")
SHARE_LABEL = (
    "Last three months as a share of the full year (display only; the row "
    "above decides)")
SHARE_FORMULA = "=IF({total}=0,0,{late}/{total})"
PRECISION_BOUND_NOTE = (
    "A spreadsheet computes the deciding formula in doubles, which is exact "
    "while the pooled last-three-months bases stay under about 2^53 / 100 "
    "cents (roughly $900 billion). Beyond that the engine's integer answer, "
    "not this cell, is authoritative.")
NO_PLACEMENT = "no placement this year"
STATED_FIGURE = "stated figure"
WHOLE_RETURN_CLAIM = (
    "Whole return: Form 4562 copies' line 22 plus no-copy activities' "
    "schedule lines = total depreciation")

_SECTION_LINE = {
    "rental_properties": "Schedule E line 18",
    "schedule_c_businesses": "Schedule C line 13",
}
NO_MATCH_NOTE = "no matching depreciable asset"
INCOME_NOTE = (
    "tenforty does not carry the section 481(a) adjustment onto the "
    "return's income: the filer must include it in gross income, and the "
    "scenario's income figures must already include the portion taken into "
    "account this year.")


@dataclass(frozen=True)
class Form3115Figures:
    """What Form 3115 and its Schedule E statement print, as printed."""
    # Line 26, signed (positive = an increase in income).
    line_26: float
    year_of_change: int
    # One per statement row: (description, date placed in service,
    # depreciation claimed under the present method).
    assets: tuple[tuple[str, date, float], ...]


@dataclass(frozen=True)
class CopyFigures:
    """What one Form 4562 copy printed."""
    # None when the copy printed no line 17 (no earlier-year assets).
    line_17: int | None
    # Section B row letter AS PRINTED on that year's form
    # (`f4562.printed_row_label`) -> column (g) deduction.
    line_19: Mapping[str, int]
    line_22: int


@dataclass(frozen=True)
class ConventionInputs:
    """The return's mid-quarter 40% test, as the engine decided it."""
    # One row per depreciation activity (`mid_quarter_bases_by_activity`),
    # stated-mode rows flagged. Two or more rows make the return
    # MULTI-ACTIVITY and select the per-activity layout.
    activities: tuple[ActivityBases, ...]
    # The boolean the orchestrator handed every Form 4562 copy.
    mid_quarter: bool


@dataclass(frozen=True)
class PrintedFigures:
    """The depreciation figures the engine filled into the forms this emit."""
    # recon key (`rental_0`, `sch_c_0`) -> that activity's Form 4562 copy.
    # Keyed by recon key also when the return has one copy; an activity
    # that printed no copy is absent.
    f4562_copies: Mapping[str, CopyFigures] = field(default_factory=dict)
    # (section, activity index) -> Schedule E line 18 / Schedule C line 13.
    # A line that printed nothing is absent and compares as zero.
    activity_lines: Mapping[tuple[str, int], int] = field(
        default_factory=dict)
    # None only on a direct call for a single-activity trail; a trail of
    # two or more activities without it is refused.
    convention: ConventionInputs | None = None
    form_3115: Form3115Figures | None = None


# --- The mirror: a Python evaluation of the formulas this workbook writes ---

_TOKEN = re.compile(r"""
    \s*(?:
      (?P<number>\d+(?:\.\d+)?)
    | "(?P<string>[^"]*)"
    | (?P<sheet>'[^']+'|[A-Za-z_][A-Za-z0-9_]*)!
    | (?P<cell>\$?[A-Z]{1,3}\$?\d+)
    | (?P<name>[A-Z]+)(?=\()
    | (?P<op><=|>=|<>|[-+*/(),:<>=])
    )""", re.VERBOSE)


class FormulaError(ValueError):
    """A formula `evaluate` cannot read: not one this workbook writes."""


def _tokens(text: str) -> list[tuple[str, str]]:
    out, pos = [], 0
    while pos < len(text):
        match = _TOKEN.match(text, pos)
        if match is None or match.end() == pos:
            raise FormulaError(f"cannot read formula at {text[pos:]!r}")
        kind = match.lastgroup
        out.append((kind, match.group(kind)))
        pos = match.end()
    return out


# The mirror's arithmetic is EXACT DECIMAL, read the way the engine reads its
# own operands (`rounding.irs_round_product`): a float is the decimal it was
# written as (25000.0 -> 25000, 0.0197 -> 0.0197), products and sums are
# exact, and ROUND rounds once, half up. Wide enough that no operation here
# is itself rounded.
_CONTEXT = Context(prec=60)


def _is_number(value) -> bool:
    return (isinstance(value, (int, float, Decimal))
            and not isinstance(value, bool))


def _number(value) -> Decimal:
    # An empty cell is zero in arithmetic, as in a spreadsheet.
    if value is None:
        return Decimal(0)
    if isinstance(value, Decimal):
        return value
    if not _is_number(value):
        raise FormulaError(f"arithmetic on a non-number: {value!r}")
    return Decimal(str(value))


def _round_half_up(value: Decimal, digits) -> Decimal:
    return value.quantize(
        Decimal(1).scaleb(-int(digits)), rounding=ROUND_HALF_UP,
        context=_CONTEXT)


def _plain(value):
    """A mirror value as an ordinary Python number: 500, not Decimal('500')."""
    if isinstance(value, Decimal):
        return int(value) if value == value.to_integral_value() else float(value)
    return value


def _same(left, right) -> bool:
    if _is_number(left) and _is_number(right):
        return _number(left) == _number(right)
    return left == right


def _flatten(args) -> list:
    out = []
    for arg in args:
        out.extend(arg if isinstance(arg, list) else [arg])
    return out


_FUNCTIONS = frozenset({"ROUND", "MIN", "MAX", "ABS", "SUM", "SUMIF", "IF"})


class _Evaluator:
    """IF is LAZY, as a spreadsheet's is: only the branch it takes is
    evaluated. The other branch is still READ -- every token of it must be
    one this workbook writes -- but no cell is fetched and no arithmetic is
    done for it (``self.skipping``), so an error that could only arise in
    the untaken branch (a division by zero behind an `=0` guard) does not
    arise. An error in the branch taken is raised as before."""

    def __init__(self, workbook, sheet: str, text: str, cache: dict):
        self.workbook = workbook
        self.sheet = sheet
        self.cache = cache
        self.tokens = _tokens(text)
        self.pos = 0
        # Depth of untaken IF branches being read; nothing is computed
        # while it is above zero.
        self.skipping = 0

    def _peek(self):
        return self.tokens[self.pos] if self.pos < len(self.tokens) else None

    def _take(self, kind=None, value=None):
        token = self._peek()
        if (token is None or (kind is not None and token[0] != kind)
                or (value is not None and token[1] != value)):
            raise FormulaError(
                f"expected {value or kind}, found {token!r}")
        self.pos += 1
        return token

    def _op(self, *values) -> str | None:
        token = self._peek()
        if token is not None and token[0] == "op" and token[1] in values:
            self.pos += 1
            return token[1]
        return None

    def result(self):
        value = self._comparison()
        if self._peek() is not None:
            raise FormulaError(f"unexpected {self._peek()!r}")
        return value

    def _comparison(self):
        left = self._sum()
        op = self._op("<=", ">=", "<>", "<", ">", "=")
        if op is None:
            return left
        right = self._sum()
        if self.skipping:
            return None
        if op == "=":
            return _same(left, right)
        if op == "<>":
            return not _same(left, right)
        if _is_number(left) and _is_number(right):
            left, right = _number(left), _number(right)
        return {
            "<": lambda: left < right, ">": lambda: left > right,
            "<=": lambda: left <= right, ">=": lambda: left >= right,
        }[op]()

    def _sum(self):
        value = self._product()
        while (op := self._op("+", "-")) is not None:
            other = self._product()
            if self.skipping:
                continue
            other = _number(other)
            value = (_CONTEXT.add(_number(value), other) if op == "+"
                     else _CONTEXT.subtract(_number(value), other))
        return value

    def _product(self):
        value = self._unary()
        while (op := self._op("*", "/")) is not None:
            other = self._unary()
            if self.skipping:
                continue
            other = _number(other)
            value = (_CONTEXT.multiply(_number(value), other) if op == "*"
                     else _CONTEXT.divide(_number(value), other))
        return value

    def _unary(self):
        if self._op("-") is not None:
            value = self._unary()
            return None if self.skipping else -_number(value)
        return self._primary()

    def _primary(self):
        kind, text = self._take()
        if kind == "number":
            return Decimal(text)
        if kind == "string":
            return text
        if kind == "name":
            return self._call(text)
        if kind == "sheet":
            return self._reference(text.strip("'"), self._take("cell")[1])
        if kind == "cell":
            return self._reference(self.sheet, text)
        if kind == "op" and text == "(":
            value = self._comparison()
            self._take("op", ")")
            return value
        raise FormulaError(f"unexpected {text!r}")

    def _reference(self, sheet: str, first: str):
        first = first.replace("$", "")
        if self._op(":") is None:
            if self.skipping:
                return None
            return _cell(self.workbook, sheet, first, self.cache)
        last = self._take("cell")[1].replace("$", "")
        if self.skipping:
            return None
        min_col, min_row, max_col, max_row = range_boundaries(
            f"{first}:{last}")
        return [
            _cell(
                self.workbook, sheet, f"{get_column_letter(col)}{row}",
                self.cache)
            for row in range(min_row, max_row + 1)
            for col in range(min_col, max_col + 1)]

    def _branch(self, taken: bool):
        """One IF branch: evaluated when ``taken``, otherwise only read."""
        if taken and not self.skipping:
            return self._comparison()
        self.skipping += 1
        try:
            self._comparison()
        finally:
            self.skipping -= 1
        return None

    def _if(self):
        test = self._comparison()
        self._take("op", ",")
        when_true = self._branch(bool(test))
        self._take("op", ",")
        when_false = self._branch(not test)
        self._take("op", ")")
        return when_true if test else when_false

    def _call(self, name: str):
        if name not in _FUNCTIONS:
            raise FormulaError(f"unsupported function {name}")
        self._take("op", "(")
        if name == "IF":
            return self._if()
        args = [self._comparison()]
        while self._op(",") is not None:
            args.append(self._comparison())
        self._take("op", ")")
        if self.skipping:
            return None
        if name == "ROUND":
            value, digits = args
            return _round_half_up(_number(value), digits)
        if name == "MIN":
            return min(_number(v) for v in _flatten(args))
        if name == "MAX":
            return max(_number(v) for v in _flatten(args))
        if name == "ABS":
            (value,) = args
            return abs(_number(value))
        if name == "SUM":
            return sum(
                (_number(v) for v in _flatten(args)), start=Decimal(0))
        keys, criterion, values = args                      # SUMIF
        return sum(
            (_number(v) for k, v in zip(keys, values, strict=True)
             if _same(k, criterion)), start=Decimal(0))


def _cell(workbook, sheet: str, coordinate: str, cache: dict):
    key = (sheet, coordinate)
    if key not in cache:
        cell = workbook[sheet][coordinate]
        cache[key] = (
            _formula(workbook, sheet, cell.value, cache)
            if cell.data_type == "f" else cell.value)
    return cache[key]


def _formula(workbook, sheet: str, formula: str, cache: dict):
    if not formula.startswith("="):
        raise FormulaError(f"not a formula: {formula!r}")
    value = _Evaluator(workbook, sheet, formula[1:], cache).result()
    # A formula that is a bare reference to an empty cell shows zero.
    return Decimal(0) if value is None else value


def cell_value(workbook, sheet: str, coordinate: str,
               cache: dict | None = None):
    """The value of one cell: an input as written, a formula evaluated.

    ``cache`` holds the cells already evaluated, keyed ``(sheet,
    coordinate)``; pass one dict across calls to share it. It is what keeps
    the cost linear: each Year-by-year cumulative cell is referenced twice by
    the row below it, so without it the work doubles with every year row."""
    return _plain(_cell(
        workbook, sheet, coordinate, {} if cache is None else cache))


def evaluate(workbook, sheet: str, formula: str, cache: dict | None = None):
    """Evaluate ``formula`` (text beginning "=") as written on ``sheet`` of
    ``workbook``, following references to other cells and sheets.

    Reads exactly what this module writes: numbers, strings, cell and range
    references, + - * /, comparisons, and ROUND, MIN, MAX, ABS, SUM, SUMIF,
    IF. Anything else raises `FormulaError`. Arithmetic is exact decimal."""
    return _plain(_formula(
        workbook, sheet, formula, {} if cache is None else cache))


# --- Layout ----------------------------------------------------------------

@dataclass(frozen=True)
class _Expectation:
    """The engine's figure for one formula cell, and whom to name if the
    formula does not reproduce it."""
    sheet: str
    coordinate: str
    expected: float
    subject: str
    fields: str


def _header(sheet, headers) -> None:
    for column, title in enumerate(headers, start=1):
        cell = sheet.cell(row=1, column=column, value=title)
        cell.font = Font(bold=True)
    sheet.freeze_panes = "A2"


def _text(sheet, coordinate: str, value: str) -> None:
    """Write ``value`` as text, whatever it looks like: a description
    beginning "=" is a description, never a formula."""
    cell = sheet[coordinate]
    cell.value = value
    cell.data_type = "s"


def _convention_text(asset: AssetAudit) -> str:
    if asset.quarter is not None:
        return f"{asset.convention} (quarter {asset.quarter})"
    return asset.convention


def _form_line(asset: AssetAudit, year: int) -> str:
    if asset.date_placed_in_service.year == year:
        # The letter the row carries on THAT year's form.
        row = form_4562.printed_row_label(asset.recovery_class, year)
        return f"line 19{row}"
    return "line 17"


def _field_path(asset: AssetAudit, name: str) -> str:
    return (f"{asset.section}[{asset.activity_index}]."
            f"depreciable_assets[{asset.asset_index}].{name}")


def _asset_provenance(asset: AssetAudit) -> str:
    basis = f"Basis: {_field_path(asset, 'basis')}"
    if asset.prior_depreciation is None:
        return f"{basis}. Prior accumulated: none (placed in service this year)"
    prior = f"Prior accumulated: {_field_path(asset, 'prior_depreciation')}"
    if asset.prior_acknowledged:
        prior += (" (stated, differs from the tables: "
                  "acknowledges_prior_depreciation_as_stated)")
    return f"{basis}. {prior}"


def _precision_note(basis: float, rate) -> str:
    """The note for a row whose product ``basis`` x ``rate`` is too large
    for a spreadsheet to round reliably; empty below the threshold."""
    if rate is None:
        return ""
    product = _CONTEXT.multiply(_number(basis), _number(rate))
    if product < SPREADSHEET_EXACT_BELOW:
        return ""
    return f". Note: {PRECISION_NOTE}"


def _year_source(asset: AssetAudit, rate) -> str:
    table = _TABLE_CITATIONS[asset.table]
    if rate is None:
        return f"recovery period ended (no cell in {table})"
    if asset.convention == macrs.MID_MONTH:
        return f"{table}, month {asset.date_placed_in_service.month}"
    return table


def _subject(asset: AssetAudit) -> str:
    return f"asset {asset.description!r} on {asset.activity_label}"


_PLAIN_SHEET_NAME = re.compile(r"[A-Za-z_][A-Za-z0-9_]*\Z")


def _sheet_ref(title: str) -> str:
    """A sheet title as a formula writes it: quoted unless it is one plain
    word (`Assets`, but `'Year-by-year'` and `'Assets rental_0'`)."""
    return title if _PLAIN_SHEET_NAME.match(title) else f"'{title}'"


def _asset_key(asset: AssetAudit) -> tuple:
    return (asset.section, asset.activity_index, asset.asset_index)


class _Tab:
    """One Assets / Year-by-year sheet pair and the assets laid out on it:
    every asset on the return (single-activity layout) or one activity's
    (multi-activity layout)."""

    def __init__(self, workbook, assets_title: str, years_title: str,
                 activities):
        self.activities = tuple(activities)
        self.assets = [a for activity in self.activities
                       for a in activity.assets]
        self.assets_sheet = workbook.create_sheet(assets_title)
        self.years_sheet = workbook.create_sheet(years_title)
        self.assets_ref = _sheet_ref(assets_title)
        self.years_ref = _sheet_ref(years_title)
        self.last_asset_row = 1 + len(self.assets)
        # (section, activity index, asset index) -> row on each sheet.
        self.asset_rows: dict[tuple, int] = {}
        self.first_year_rows: dict[tuple, int] = {}
        year_row = 2
        for row, asset in enumerate(self.assets, start=2):
            self.asset_rows[_asset_key(asset)] = row
            self.first_year_rows[_asset_key(asset)] = year_row
            year_row += len(asset.years)

    def column(self, letter: str) -> str:
        """The Assets sheet's data rows in one column, as a range."""
        return (f"{self.assets_ref}!{letter}2:"
                f"{letter}{self.last_asset_row}")


class _Builder:
    def __init__(self, resolved: tuple[ActivityAudit, ...],
                 printed: PrintedFigures, year: int):
        self.resolved = resolved
        self.printed = printed
        self.year = year
        self.workbook = Workbook()
        self.expectations: list[_Expectation] = []
        # Why this hand-off cannot be laid out as it must be.
        self.reasons: list[str] = []
        self.override_rows: dict[tuple[str, int], int] = {}
        # recon key -> Tie-outs row of that copy's line 22 (multi-activity).
        self.line_22_rows: dict[str, int] = {}
        self.deciding_cell: str | None = None
        self.tie_outs = self.workbook.active
        self.tie_outs.title = TIE_OUTS
        self.numbers = {
            (activity.section, activity.index): n
            for n, activity in enumerate(resolved, start=1)}
        convention = printed.convention
        self.recon_keys: dict[tuple[str, int], str] = (
            {} if convention is None else {
                (row.section, row.index): row.recon_key
                for row in convention.activities})
        self.multi = self._is_multi_activity()
        if self.multi:
            self.pooled_sheet = self.workbook.create_sheet(MID_QUARTER_TEST)
            self.tabs = [
                _Tab(self.workbook,
                     assets_tab_name(self._recon_key(activity)),
                     years_tab_name(self._recon_key(activity)), (activity,))
                for activity in resolved]
        else:
            self.tabs = [_Tab(self.workbook, ASSETS, YEAR_BY_YEAR, resolved)]
            if len(printed.f4562_copies) > 1:
                self.reasons.append(
                    f"{len(printed.f4562_copies)} Form 4562 copies "
                    f"({', '.join(sorted(printed.f4562_copies))}) were "
                    f"handed to a single-activity workbook, which ties out "
                    f"one")
        self.tab_of = {
            (activity.section, activity.index): tab
            for tab in self.tabs for activity in tab.activities}
        self.assets = [a for tab in self.tabs for a in tab.assets]

    def _is_multi_activity(self) -> bool:
        """Whether this is a MULTI-ACTIVITY return: two or more depreciation
        activities (listing assets or stating a figure), as counted by the
        `ConventionInputs` handed in. FAIL-CLOSED: a trail of two or more
        activities with no `ConventionInputs` is a reason, never a quiet
        single-activity workbook."""
        convention = self.printed.convention
        if convention is None:
            if len(self.resolved) >= 2:
                self.reasons.append(
                    f"the return has {len(self.resolved)} activities with "
                    f"depreciable assets, but no ConventionInputs were "
                    f"handed to the workbook: the Mid-quarter test sheet is "
                    f"mandatory on a multi-activity workbook and cannot be "
                    f"laid out without them")
            return False
        on_trail = {(a.section, a.index) for a in self.resolved}
        missing = [a.label for a in self.resolved
                   if (a.section, a.index) not in self.recon_keys]
        unknown = [row.recon_key for row in convention.activities
                   if not row.stated
                   and (row.section, row.index) not in on_trail]
        if missing:
            self.reasons.append(
                f"the ConventionInputs handed to the workbook have no row "
                f"for {', '.join(missing)}")
        if unknown:
            self.reasons.append(
                f"the ConventionInputs handed to the workbook list "
                f"{', '.join(unknown)} as listing assets, but the "
                f"depreciation trail has no such activity")
        if missing or unknown:
            return False
        return len(convention.activities) >= 2

    def _recon_key(self, activity: ActivityAudit) -> str:
        return self.recon_keys[(activity.section, activity.index)]

    def _places_this_year(self, activity: ActivityAudit) -> bool:
        """Ruling A: the activity gets a Form 4562 copy."""
        return any(asset.date_placed_in_service.year == self.year
                   for asset in activity.assets)

    def _expect(self, sheet, coordinate, expected, subject, fields) -> None:
        self.expectations.append(_Expectation(
            sheet.title, coordinate, expected, subject, fields))

    def build(self):
        for tab in self.tabs:
            self._write_assets(tab)
            self._write_years(tab)
            self._write_overrides(tab)
        _header(self.tie_outs, TIE_OUT_HEADERS)
        if self.multi:
            self._write_pooled_test()
            self._write_tie_outs_multi()
        else:
            self._write_tie_outs_single()
        if self.printed.form_3115 is not None:
            self._write_481a(self.printed.form_3115)
        return self.workbook

    def _write_assets(self, tab: _Tab) -> None:
        sheet = tab.assets_sheet
        _header(sheet, ASSET_HEADERS)
        for asset in tab.assets:
            r = tab.asset_rows[_asset_key(asset)]
            current_year_row = (
                tab.first_year_rows[_asset_key(asset)] + len(asset.years) - 1)
            sheet[f"A{r}"] = self.numbers[
                (asset.section, asset.activity_index)]
            _text(sheet, f"B{r}", asset.activity_label)
            _text(sheet, f"C{r}", asset.description)
            sheet[f"D{r}"] = asset.date_placed_in_service
            sheet[f"E{r}"] = asset.method
            sheet[f"F{r}"] = asset.recovery_class
            sheet[f"G{r}"] = _convention_text(asset)
            sheet[f"H{r}"] = f"={tab.years_ref}!D{current_year_row}"
            sheet[f"I{r}"] = asset.basis
            sheet[f"J{r}"] = (
                0 if asset.prior_depreciation is None
                else asset.prior_depreciation)
            sheet[f"K{r}"] = f"=ROUND(MAX(0,MIN(I{r}*H{r},I{r}-J{r})),0)"
            sheet[f"L{r}"] = f"=J{r}+K{r}"
            sheet[f"M{r}"] = _form_line(asset, self.year)
            _text(sheet, f"N{r}", _asset_provenance(asset) + _precision_note(
                asset.basis, asset.current.rate))
            subject = _subject(asset)
            self._expect(sheet, f"H{r}", asset.current.rate or 0.0, subject,
                         "`recovery_class` / `date_placed_in_service`")
            self._expect(sheet, f"K{r}", asset.current.amount, subject,
                         "`basis` / `prior_depreciation`")

    def _write_years(self, tab: _Tab) -> None:
        sheet = tab.years_sheet
        _header(sheet, YEAR_HEADERS)
        for asset in tab.assets:
            r = tab.asset_rows[_asset_key(asset)]
            first = tab.first_year_rows[_asset_key(asset)]
            basis = f"{tab.assets_ref}!$I${r}"
            for n, year in enumerate(asset.years, start=first):
                sheet[f"A{n}"] = f"={tab.assets_ref}!C{r}"
                sheet[f"B{n}"] = year.recovery_year
                sheet[f"C{n}"] = year.tax_year
                sheet[f"D{n}"] = year.rate
                if n == first:
                    sheet[f"E{n}"] = (
                        f"=ROUND(MAX(0,MIN({basis}*D{n},{basis})),0)")
                    sheet[f"F{n}"] = f"=E{n}"
                else:
                    sheet[f"E{n}"] = (
                        f"=ROUND(MAX(0,MIN({basis}*D{n},"
                        f"{basis}-F{n - 1})),0)")
                    sheet[f"F{n}"] = f"=F{n - 1}+E{n}"
                sheet[f"G{n}"] = (
                    _year_source(asset, year.rate)
                    + _precision_note(asset.basis, year.rate))
                subject = f"{_subject(asset)}, tax year {year.tax_year}"
                self._expect(sheet, f"E{n}", year.amount, subject, "`basis`")
                self._expect(sheet, f"F{n}", year.taken_before + year.amount,
                             subject, "`basis`")

    def _write_overrides(self, tab: _Tab) -> None:
        overridden = [a for a in tab.activities
                      if a.mode == MODE_ASSETS_OVERRIDDEN]
        if not overridden:
            return
        sheet = tab.assets_sheet
        header_row = tab.last_asset_row + 2
        for column, title in OVERRIDE_HEADERS.items():
            sheet[f"{column}{header_row}"] = title
            sheet[f"{column}{header_row}"].font = Font(bold=True)
        for r, activity in enumerate(overridden, start=header_row + 1):
            self.override_rows[(activity.section, activity.index)] = r
            path = f"{activity.section}[{activity.index}].depreciation_override"
            sheet[f"A{r}"] = self.numbers[(activity.section, activity.index)]
            _text(sheet, f"B{r}", activity.label)
            sheet[f"I{r}"] = activity.override_amount
            sheet[f"J{r}"] = activity.restates_engine_amount
            _text(sheet, f"N{r}", (
                f"Override amount: {path}.amount (acknowledged; the forms "
                f"use this figure). Restated engine amount: "
                f"{path}.restates_engine_amount"))

    def _tie_out(self, claim: str, computed: str, printed, source: str,
                 expected: float, subject: str, fields: str) -> int:
        """Write one tie-out row; the row it landed on."""
        sheet = self.tie_outs
        r = sheet.max_row + 1
        _text(sheet, f"A{r}", claim)
        sheet[f"B{r}"] = computed
        sheet[f"C{r}"] = printed
        _text(sheet, f"D{r}", source)
        sheet[f"E{r}"] = f"=B{r}-C{r}"
        sheet[f"F{r}"] = VERDICT_FORMULA.format(row=r)
        self._expect(sheet, f"B{r}", expected, subject, fields)
        return r

    def _form_4562_rows(self, tab: _Tab, figures: CopyFigures | None, *,
                        prefix: str, where: str, subject: str,
                        always_line_22: bool) -> int | None:
        """The line 17, line 19 and line 22 tie-outs of one Form 4562,
        computed from ``tab``'s assets against the ``figures`` it printed.
        The row its line 22 landed on, if that row was written."""
        lines, deductions = tab.column("M"), tab.column("K")

        def line_total(line: str) -> int:
            return sum(a.current.amount for a in tab.assets
                       if _form_line(a, self.year) == line)

        if figures is not None and figures.line_17 is not None:
            self._tie_out(
                f"{prefix}Form 4562 line 17: MACRS deductions for assets "
                f"placed in service before {self.year}",
                f'=SUMIF({lines},"line 17",{deductions})',
                figures.line_17, f"{where} line 17 as printed",
                line_total("line 17"), subject, "Form 4562 line 17")
        for letter in sorted(figures.line_19 if figures is not None else ()):
            self._tie_out(
                f"{prefix}Form 4562 line 19{letter}, column (g): "
                f"depreciation deduction",
                f'=SUMIF({lines},"line 19{letter}",{deductions})',
                figures.line_19[letter],
                f"{where} line 19{letter} column (g) as printed",
                line_total(f"line 19{letter}"), subject,
                f"Form 4562 line 19{letter}")
        if figures is None and not always_line_22:
            return None
        return self._tie_out(
            f"{prefix}Form 4562 line 22: total depreciation",
            f"=SUM({deductions})",
            None if figures is None else figures.line_22,
            f"{where} line 22 as printed",
            sum(a.current.amount for a in tab.assets), subject,
            "Form 4562 line 22")

    def _schedule_line(self, section: str, index: int) -> str:
        line = _SECTION_LINE[section]
        if section == "schedule_c_businesses":
            line = f"{line} (business {index + 1})"
        return line

    def _schedule_rows(self) -> dict[tuple[str, int], int]:
        """One schedule-line tie-out per asset-mode activity (two for an
        overridden one). ``(section, index)`` -> the row whose Printed cell
        is that activity's schedule depreciation line."""
        rows: dict[tuple[str, int], int] = {}
        printed = self.printed
        for activity in self.resolved:
            key = (activity.section, activity.index)
            tab = self.tab_of[key]
            line = self._schedule_line(activity.section, activity.index)
            line_printed = printed.activity_lines.get(key, 0)
            asset_sum = (f"=SUMIF({tab.column('A')},{self.numbers[key]},"
                         f"{tab.column('K')})")
            if activity.mode != MODE_ASSETS_OVERRIDDEN:
                rows[key] = self._tie_out(
                    f"{line}: depreciation, {activity.label}", asset_sum,
                    line_printed, f"{line} as printed",
                    activity.engine_amount, activity.label, line)
                continue
            r = self.override_rows[key]
            rows[key] = self._tie_out(
                f"{line}: depreciation, {activity.label} "
                f"(depreciation_override in effect)",
                f"=ROUND({tab.assets_ref}!I{r},0)", line_printed,
                f"{line} as printed",
                irs_round(activity.used_amount), activity.label,
                "`depreciation_override.amount`")
            self._tie_out(
                f"Engine figure restated by the override, {activity.label}",
                asset_sum, f"=ROUND({tab.assets_ref}!J{r},0)",
                f"depreciation_override.restates_engine_amount "
                f"({tab.assets_ref}!J{r}); not printed on a form",
                activity.engine_amount, activity.label,
                "`depreciation_override.restates_engine_amount`")
        return rows

    def _lone_copy(self) -> CopyFigures | None:
        """The one Form 4562 a single-activity return printed, if any."""
        copies = list(self.printed.f4562_copies.values())
        return copies[0] if copies else None

    def _write_tie_outs_single(self) -> None:
        [tab] = self.tabs
        self._form_4562_rows(
            tab, self._lone_copy(), prefix="", where="Form 4562",
            subject="the return's depreciable assets", always_line_22=False)
        self._schedule_rows()

    def _write_tie_outs_multi(self) -> None:
        """Per copy, then per schedule line, then the whole return."""
        convention = self.printed.convention
        copy_activities = [
            a for a in self.resolved if self._places_this_year(a)]
        for activity in copy_activities:
            key = self._recon_key(activity)
            self.line_22_rows[key] = self._form_4562_rows(
                self.tab_of[(activity.section, activity.index)],
                self.printed.f4562_copies.get(key),
                prefix=f"Form 4562 copy {key}: ",
                where=f"Form 4562 copy {key},",
                subject=f"the depreciable assets of {activity.label}",
                always_line_22=True)
        schedule_rows = self._schedule_rows()

        # The whole return: every copy's line 22, plus the schedule
        # depreciation line of every activity that has no copy.
        sheet = self.tie_outs
        no_copy_rows: dict[tuple[str, int], int] = {}
        for row in convention.activities:
            if row.recon_key in self.line_22_rows:
                continue
            key = (row.section, row.index)
            line = self._schedule_line(row.section, row.index)
            reason = STATED_FIGURE if row.stated else NO_PLACEMENT
            r = sheet.max_row + 1
            _text(sheet, f"A{r}", (
                f"Whole return: {row.recon_key} has no Form 4562 copy "
                f"({reason}): {line}: depreciation, {row.label}"))
            sheet[f"C{r}"] = self.printed.activity_lines.get(key, 0)
            _text(sheet, f"D{r}", f"{line} as printed")
            no_copy_rows[key] = r
        no_copy_cells = [f"C{r}" for r in no_copy_rows.values()]
        computed = [f"B{r}" for r in self.line_22_rows.values()]
        printed = [f"C{schedule_rows[(a.section, a.index)]}"
                   for a in copy_activities]
        r = sheet.max_row + 1
        _text(sheet, f"A{r}", WHOLE_RETURN_CLAIM)
        sheet[f"B{r}"] = "=" + "+".join(computed + no_copy_cells)
        sheet[f"C{r}"] = "=" + "+".join(printed + no_copy_cells)
        _text(sheet, f"D{r}", (
            "every activity's schedule depreciation line as printed, "
            "summed"))
        sheet[f"E{r}"] = f"=B{r}-C{r}"
        sheet[f"F{r}"] = VERDICT_FORMULA.format(row=r)
        self._expect(
            sheet, f"B{r}",
            sum(a.engine_amount for a in copy_activities) + sum(
                self.printed.activity_lines.get(key, 0)
                for key in no_copy_rows),
            "the whole return",
            "every Form 4562 copy's line 22 and every other activity's "
            "schedule depreciation line")

    def _write_pooled_test(self) -> None:
        """The taxpayer-wide 40% test, laid out from the per-activity cents
        the engine decided on. No printed copy shows why the return is (or
        is not) mid-quarter: the bases span activities."""
        sheet = self.pooled_sheet
        _header(sheet, MID_QUARTER_HEADERS)
        rows = self.printed.convention.activities
        for r, row in enumerate(rows, start=2):
            _text(sheet, f"A{r}", row.recon_key)
            _text(sheet, f"B{r}", row.label)
            if row.stated:
                _text(sheet, f"E{r}", STATED_ROW_NOTE)
                continue
            sheet[f"C{r}"] = row.last_quarter_cents
            sheet[f"D{r}"] = row.full_year_cents
        last = 1 + len(rows)
        total = last + 1
        _text(sheet, f"A{total}", "Total")
        sheet[f"C{total}"] = f"=SUM(C2:C{last})"
        sheet[f"D{total}"] = f"=SUM(D2:D{last})"
        decide = total + 1
        _text(sheet, f"A{decide}", DECIDING_LABEL)
        sheet[f"C{decide}"] = DECIDING_FORMULA.format(
            late=f"C{total}", total=f"D{total}")
        self.deciding_cell = f"C{decide}"
        share = total + 2
        _text(sheet, f"A{share}", SHARE_LABEL)
        sheet[f"C{share}"] = SHARE_FORMULA.format(
            late=f"C{total}", total=f"D{total}")
        _text(sheet, f"A{total + 4}", PRECISION_BOUND_NOTE)
        subject = "the Mid-quarter test totals"
        fields = "`basis` / `date_placed_in_service` of property placed " \
                 "this year"
        self._expect(sheet, f"C{total}",
                     sum(row.last_quarter_cents for row in rows),
                     subject, fields)
        self._expect(sheet, f"D{total}",
                     sum(row.full_year_cents for row in rows),
                     subject, fields)

    def fail_closed_reasons(self, printed_is_final: bool) -> list[str]:
        """Why a multi-activity workbook does not stand behind the return,
        beyond what the cell-by-cell mirror reports.

        1. The deciding cell, evaluated by the mirror, must equal the
           boolean the orchestrator handed every Form 4562 copy.
        2. Once the forms are printed (``printed_is_final``): each copy's
           line 22 from its own tab must equal that copy's printed line 22,
           and the copies handed in must be exactly the activities that
           place property this year."""
        if not self.multi:
            return []
        reasons = []
        where = f"{_sheet_ref(MID_QUARTER_TEST)}!{self.deciding_cell}"
        answer = self.printed.convention.mid_quarter
        try:
            decided = cell_value(
                self.workbook, MID_QUARTER_TEST, self.deciding_cell)
        except _UNREADABLE:
            decided = answer        # unreadable: the mirror reports it
        if decided is not answer:
            reasons.append(
                f"the Mid-quarter test deciding cell ({where}) gives "
                f"{decided!r} where the return's 40% answer, as handed to "
                f"every Form 4562 copy, is {answer!r}")
        if not printed_is_final:
            return reasons
        copies = self.printed.f4562_copies
        for key, row in self.line_22_rows.items():
            figures = copies.get(key)
            if figures is None:
                reasons.append(
                    f"Form 4562 copy {key}: no figures for it were handed "
                    f"to the workbook, though that activity places property "
                    f"in service this year")
                continue
            try:
                computed = cell_value(self.workbook, TIE_OUTS, f"B{row}")
            except _UNREADABLE:
                continue            # unreadable: the mirror reports it
            if abs(computed - figures.line_22) >= TOLERANCE:
                reasons.append(
                    f"Form 4562 copy {key}: line 22 computed from its own "
                    f"tab ({TIE_OUTS}!B{row}) is {_figure(computed)} where "
                    f"the copy prints {_figure(figures.line_22)}")
        for key in sorted(set(copies) - set(self.line_22_rows)):
            reasons.append(
                f"Form 4562 copy {key}: figures for it were handed to the "
                f"workbook, but no activity {key} on the depreciation trail "
                f"places property in service this year")
        return reasons

    def _write_481a(self, figures: Form3115Figures) -> None:
        sheet = self.workbook.create_sheet(SECTION_481A)
        _header(sheet, SECTION_481A_HEADERS)
        by_identity = {
            (a.description, a.date_placed_in_service): (a, tab)
            for tab in self.tabs for a in tab.assets}
        last_prior_year = figures.year_of_change - 1
        r = 1
        for n, (description, placed, claimed) in enumerate(figures.assets):
            r += 1
            _text(sheet, f"A{r}", description)
            sheet[f"B{r}"] = placed
            sheet[f"C{r}"] = claimed
            _text(sheet, f"G{r}", (
                f"Claimed: form_3115.assets[{n}]."
                f"depreciation_claimed_present_method"))
            match = by_identity.get((description, placed))
            if match is None:
                _text(sheet, f"F{r}", NO_MATCH_NOTE)
                continue
            asset, tab = match
            rows = [
                tab.first_year_rows[_asset_key(asset)] + i
                for i, year in enumerate(asset.years)
                if year.tax_year == last_prior_year]
            if rows:
                sheet[f"D{r}"] = f"={tab.years_ref}!F{rows[0]}"
            else:
                sheet[f"D{r}"] = 0
                _text(sheet, f"F{r}", (
                    f"no depreciation allowable before "
                    f"{figures.year_of_change}: placed in service in "
                    f"{placed.year}"))
            sheet[f"E{r}"] = f"=C{r}-D{r}"
        last = r
        net = last + 2
        _text(sheet, f"A{net}", "Net section 481(a) adjustment")
        sheet[f"E{net}"] = NET_481A_FORMULA.format(last=last)
        _text(sheet, f"F{net}", (
            "claimed less allowable, matched assets only; positive is an "
            "increase in income"))
        stated = net + 1
        _text(sheet, f"A{stated}", "Stated section 481(a) adjustment")
        sheet[f"E{stated}"] = figures.line_26
        _text(sheet, f"F{stated}", INCOME_NOTE)
        _text(sheet, f"G{stated}", "form_3115.section_481a_adjustment")

        tie = self.tie_outs
        t = tie.max_row + 1
        _text(tie, f"A{t}", "Form 3115 line 26: net section 481(a) adjustment")
        tie[f"B{t}"] = f"='{SECTION_481A}'!E{net}"
        tie[f"C{t}"] = figures.line_26
        _text(tie, f"D{t}", (
            "Form 3115 line 26 as printed (stated: "
            "form_3115.section_481a_adjustment; tenforty does not compute "
            "or check it)"))
        tie[f"E{t}"] = f"=B{t}-C{t}"
        tie[f"F{t}"] = VERDICT_FORMULA.format(row=t)


def _figure(value) -> str:
    """A number as a reader writes it: 500, not 500.0."""
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


def _year_of(resolved: tuple[ActivityAudit, ...]) -> int:
    return resolved[0].assets[0].current.tax_year


def formula_cells(workbook) -> list[tuple[str, str]]:
    """``(sheet, coordinate)`` of every formula cell in ``workbook``."""
    return [
        (sheet.title, cell.coordinate)
        for sheet in workbook.worksheets for row in sheet.iter_rows()
        for cell in row if cell.data_type == "f"]


# What evaluating a formula this module did not write can raise.
_UNREADABLE = (
    FormulaError, TypeError, ValueError, KeyError, ArithmeticError,
    RecursionError)


def _mirror_mismatches(workbook, expectations) -> list[str]:
    """Why the workbook does not reproduce the engine. FAIL-CLOSED: every
    formula cell on every sheet is evaluated, whether or not the engine has
    a figure for it, and one the mirror cannot read is a reason -- never a
    cell it skips."""
    reasons, values, cache = [], {}, {}
    for sheet, coordinate in formula_cells(workbook):
        try:
            values[(sheet, coordinate)] = cell_value(
                workbook, sheet, coordinate, cache)
        except _UNREADABLE as error:
            reasons.append(
                f"the workbook formula at {sheet}!{coordinate} cannot be "
                f"checked: {error}")
    for e in expectations:
        if (e.sheet, e.coordinate) not in values:
            continue        # unreadable: already a reason above
        got = values[(e.sheet, e.coordinate)]
        if (isinstance(got, (int, float)) and not isinstance(got, bool)
                and abs(got - e.expected) < TOLERANCE):
            continue
        reasons.append(
            f"{e.subject}: the workbook formula at "
            f"{e.sheet}!{e.coordinate} gives {_figure(got)} where the "
            f"engine's figure is {_figure(e.expected)} (inputs: "
            f"{e.fields})")
    return reasons


def _build(resolved, printed, *, printed_is_final: bool = False):
    """The workbook, and every reason it must not be emitted.

    ``printed_is_final`` says ``printed`` holds what the forms really
    printed (the write, after the forms are prepared) rather than a stand-in
    (the refusal's pre-check, before any form exists); only then are a
    multi-activity workbook's copies checked against their printed figures."""
    builder = _Builder(resolved, printed, _year_of(resolved))
    workbook = builder.build()
    reasons = (builder.reasons
               + _mirror_mismatches(workbook, builder.expectations)
               + builder.fail_closed_reasons(printed_is_final))
    return workbook, reasons


def unbuildable_reasons(scenario) -> list[str]:
    """Why this return's audit workbook cannot be built; empty when it can
    (and for a return with no asset-mode depreciation, which has none).

    The offenders predicate of the `depreciation_audit_unbuildable` refusal.
    The reasons do not depend on the printed figures, so it runs before any
    form is prepared."""
    resolved = audit_trail(scenario)
    if not resolved:
        return []
    convention = ConventionInputs(
        activities=mid_quarter_bases_by_activity(scenario),
        mid_quarter=mid_quarter_applies(scenario))
    return _build(resolved, PrintedFigures(convention=convention))[1]


def write_depreciation_audit(
        resolved: tuple[ActivityAudit, ...], printed: PrintedFigures,
        out_path: Path) -> Path:
    """Write the audit workbook for ``resolved`` (a non-empty
    `resolver.audit_trail`) and the figures ``printed`` on the forms.

    Refuses, through the registry, when a formula does not reproduce the
    engine's figure, when a multi-activity workbook's deciding cell or a
    copy's line 22 disagrees with what was handed in, when two or more
    activities arrive without `ConventionInputs`, or when the file cannot
    be written; nothing is left at ``out_path`` then."""
    workbook, reasons = _build(resolved, printed, printed_is_final=True)
    if reasons:
        raise_scoped_refusal(REFUSAL, reasons)
    out_path = Path(out_path)
    try:
        workbook.save(out_path)
    except OSError as error:
        out_path.unlink(missing_ok=True)
        raise_scoped_refusal(
            REFUSAL, [f"{out_path.name} could not be written ({error})"])
    return out_path
