"""W-2s on the workbook path: every W-2 reaches the 'W-2s' sheet, or the
return refuses.

The 1040 workbook holds up to six employers per person in columns C..H of
the 'W-2s' sheet: the taxpayer's block, and the spouse's block at a fixed
row offset below it. The flattener numbers W-2s per owner
(`w2_<field>_<n>` for the taxpayer, `w2_spouse_<field>_<n>` for a W-2 with
owner "spouse"); the year's F1040.INPUTS maps them to cells. A flattened
W-2 key the year does not map refuses instead of being dropped silently
(it once was: only employer #1 was mapped, in every year, so a spouse's or
a second employer's W-2 never reached the sheet).

Addresses are verified against the committed workbook with openpyxl (no
recalc). Synthetic data only.
"""
import dataclasses
import re
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

from tenforty.mappings.f1040 import F1040
from tenforty.models import FilingStatus, Scenario
from tenforty.oracle.flattener import flatten_scenario
from tenforty.orchestrator import ReturnOrchestrator
from tests.fixtures.spine_battery import _battery_config, _w2
from tests.helpers import REPO_ROOT

FIELDS = ("wages", "fed_withheld", "ss_wages", "ss_withheld",
          "medicare_wages", "medicare_withheld", "state_wages",
          "state_withheld")
EMPLOYERS = range(1, 7)
COLUMNS = "CDEFGH"


def _workbook(year: int):
    return openpyxl.load_workbook(
        REPO_ROOT / "spreadsheets" / "federal" / str(year) / "1040.xlsx")


def _row(address: str) -> int:
    return int(re.fullmatch(r"[A-Z]+(\d+)", address).group(1))


def _scenario(year, status, owners):
    config = _battery_config(year, status)
    w2s = [dataclasses.replace(
               _w2(year, f"Synthetic Employer {i}", 10_000.0 + i, 1_000.0,
                   state_tax_withheld=500.0),
               owner=owner, state_wages=10_000.0 + i)
           for i, owner in enumerate(owners, start=1)]
    return Scenario(config=config, w2s=w2s)


class FlattenerOwnerRoutingTests(unittest.TestCase):
    def test_spouse_w2_flattens_to_the_spouse_keys(self):
        s = _scenario(2025, FilingStatus.MARRIED_FILING_JOINTLY,
                      ["taxpayer", "spouse"])
        flat = flatten_scenario(s)
        self.assertEqual(flat["w2_wages_1"], s.w2s[0].wages)
        self.assertEqual(flat["w2_spouse_wages_1"], s.w2s[1].wages)
        self.assertNotIn("w2_wages_2", flat)
        for field in FIELDS:
            with self.subTest(field=field):
                self.assertIn(f"w2_spouse_{field}_1", flat)

    def test_taxpayer_w2s_number_within_the_owner(self):
        s = _scenario(2025, FilingStatus.SINGLE, [None, None, None])
        flat = flatten_scenario(s)
        self.assertEqual([flat[f"w2_wages_{n}"] for n in (1, 2, 3)],
                         [w.wages for w in s.w2s])
        self.assertFalse([k for k in flat if k.startswith("w2_spouse_")])

    def test_spouse_numbering_is_independent(self):
        s = _scenario(2025, FilingStatus.MARRIED_FILING_JOINTLY,
                      ["spouse", "taxpayer", "spouse"])
        flat = flatten_scenario(s)
        self.assertEqual(flat["w2_wages_1"], s.w2s[1].wages)
        self.assertEqual(flat["w2_spouse_wages_1"], s.w2s[0].wages)
        self.assertEqual(flat["w2_spouse_wages_2"], s.w2s[2].wages)


class W2Mapping2025Tests(unittest.TestCase):
    """Six employers per person, both blocks, every field, 2025."""

    YEAR = 2025

    def setUp(self):
        self.inputs = F1040.INPUTS[self.YEAR]
        self.sheet_map = F1040.SHEET_MAP[self.YEAR]

    def test_every_taxpayer_field_has_its_spouse_twin_for_six_employers(self):
        for field in FIELDS:
            for n in EMPLOYERS:
                with self.subTest(field=field, employer=n):
                    tp = self.inputs[f"w2_{field}_{n}"]
                    sp = self.inputs[f"w2_spouse_{field}_{n}"]
                    self.assertEqual(tp[0], COLUMNS[n - 1])
                    self.assertEqual(sp[0], COLUMNS[n - 1])
                    self.assertEqual(
                        _row(tp), _row(self.inputs[f"w2_{field}_1"]))
                    self.assertGreater(_row(sp), _row(tp))
                    self.assertEqual(self.sheet_map[f"w2_{field}_{n}"], "W-2s")
                    self.assertEqual(
                        self.sheet_map[f"w2_spouse_{field}_{n}"], "W-2s")

    def test_employer_one_addresses_unchanged(self):
        # The taxpayer's employer-#1 cells are the ones mapped before.
        expected = {"wages": "C3", "fed_withheld": "C4", "ss_wages": "C5",
                    "ss_withheld": "C6", "medicare_wages": "C7",
                    "medicare_withheld": "C8", "state_wages": "C28",
                    "state_withheld": "C29"}
        for field, address in expected.items():
            with self.subTest(field=field):
                self.assertEqual(self.inputs[f"w2_{field}_1"], address)

    def test_addresses_match_the_workbook(self):
        ws = _workbook(self.YEAR)["W-2s"]
        for field in FIELDS:
            tp_row = _row(self.inputs[f"w2_{field}_1"])
            sp_row = _row(self.inputs[f"w2_spouse_{field}_1"])
            with self.subTest(field=field):
                # The spouse row carries the taxpayer row's own label.
                self.assertEqual(ws.cell(row=sp_row, column=2).value,
                                 ws.cell(row=tp_row, column=2).value)
                for n in EMPLOYERS:
                    for key in (f"w2_{field}_{n}", f"w2_spouse_{field}_{n}"):
                        # Writable input cells: blank, not a formula.
                        self.assertIsNone(ws[self.inputs[key]].value, key)

    def test_spouse_block_is_the_spouses(self):
        ws = _workbook(self.YEAR)["W-2s"]
        sp_wages_row = _row(self.inputs["w2_spouse_wages_1"])
        caption = ws.cell(row=sp_wages_row - 1, column=2).value
        self.assertIn("Name_1st_Sp", caption)
        for n in EMPLOYERS:
            self.assertIn("Employer #",
                          ws.cell(row=sp_wages_row - 1, column=2 + n).value)


class UnmappedW2RefusalTests(unittest.TestCase):
    """A W-2 the year's workbook map cannot place refuses, before the
    workbook runs."""

    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.orch = ReturnOrchestrator(
            spreadsheets_dir=REPO_ROOT / "spreadsheets",
            work_dir=Path(tmp.name) / "work")

    def _refuses(self, scenario):
        with mock.patch.object(
                self.orch.engine, "compute",
                side_effect=AssertionError("workbook ran")):
            with self.assertRaises(NotImplementedError) as ctx:
                self.orch._compute_1040_via_workbook(scenario)
        return str(ctx.exception)

    def test_seventh_employer_refuses(self):
        s = _scenario(2025, FilingStatus.MARRIED_FILING_SEPARATELY,
                      [None] * 7)
        message = self._refuses(s)
        self.assertIn("w2_wages_7", message)
        self.assertIn("2025", message)

    def test_seventh_spouse_employer_refuses(self):
        s = _scenario(2025, FilingStatus.MARRIED_FILING_JOINTLY,
                      ["taxpayer"] + ["spouse"] * 7)
        self.assertIn("w2_spouse_wages_7", self._refuses(s))

    def test_unmapped_year_spouse_w2_refuses(self):
        # 2024's spouse block is not mapped (2025-only this round).
        s = _scenario(2024, FilingStatus.MARRIED_FILING_JOINTLY,
                      ["taxpayer", "spouse"])
        message = self._refuses(s)
        self.assertIn("w2_spouse_wages_1", message)
        self.assertIn("2024", message)

    def test_six_employers_each_reach_the_engine(self):
        s = _scenario(2025, FilingStatus.MARRIED_FILING_JOINTLY,
                      ["taxpayer"] * 6 + ["spouse"] * 6)
        with mock.patch.object(
                self.orch.engine, "compute",
                side_effect=RuntimeError("engine reached")) as compute:
            with self.assertRaisesRegex(RuntimeError, "engine reached"):
                self.orch._compute_1040_via_workbook(s)
        written = compute.call_args.kwargs["inputs"]
        self.assertEqual(
            sorted(k for k in written if k.startswith("w2_wages_")),
            sorted(f"w2_wages_{n}" for n in EMPLOYERS))
        self.assertEqual(
            sorted(k for k in written if k.startswith("w2_spouse_wages_")),
            sorted(f"w2_spouse_wages_{n}" for n in EMPLOYERS))
