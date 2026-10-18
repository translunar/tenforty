"""The audit workbook on a return with more than one depreciation activity.

A MULTI-ACTIVITY return has two or more activities that list assets or state
a depreciation figure. Its workbook is laid out per activity:

  Tie-outs, Mid-quarter test, then `Assets <recon key>` / `Years <recon key>`
  for each activity that lists assets, then 481(a) when there is one.

A single-activity return keeps the three-sheet workbook unchanged
(tests/test_depreciation_audit_emit.py pins it cell for cell).

THE TWO SHAPES (tests/test_f4562_per_activity_emit.py). All 5-year personal
property; first-year rates from Publication 946: half-year 20%, mid-quarter
first quarter 35%, fourth quarter 5%.

  under  rental 9,000 March; business 1,000 November
         pooled 1,000 of 10,000 -> half-year: rental 1,800, business 200
  over   rental 6,000 March + 3,000 November; business 3,000 December
         pooled 6,000 of 12,000 -> mid-quarter: rental 2,100 + 150 = 2,250,
         business 150

PDF field paths are 2025 literals (Form 4562 line 22), as in
tests/test_f4562_emit_gate.py.
"""
import dataclasses
import tempfile
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

import openpyxl
from pypdf import PdfReader

from tenforty.audit import depreciation_workbook as dw
from tenforty.forms.depreciation import resolver
from tenforty.models import DepreciableAsset
from tenforty.orchestrator import ReturnOrchestrator
from tests import _audit_workbook_fixtures as fx
from tests.helpers import SPREADSHEETS_DIR

F4562_LINE_22 = "topmostSubform[0].Page2[0].f2_2[0]"
RENTAL = ("rental_properties", 0)
BUSINESS = ("schedule_c_businesses", 0)

TIE = "Tie-outs"
POOLED = "Mid-quarter test"
STATED_NOTE = (
    "placements not visible; the return carries "
    "acknowledges_no_personal_property_behind_stated_depreciation")

# 5-year, placed 2023, 1,920 this year (tests/_audit_workbook_fixtures.py).
PRIOR_EQUIPMENT = dataclasses.replace(fx.APPLIANCE, description="Equipment")


def under():
    return fx.multi_scenario(
        rental_assets=(fx.equipment(9_000.0, 3),),
        businesses=[fx.business(fx.equipment(1_000.0, 11))])


def over():
    return fx.multi_scenario(
        rental_assets=(fx.equipment(6_000.0, 3, "Furniture"),
                       fx.equipment(3_000.0, 11, "Appliances")),
        businesses=[fx.business(fx.equipment(3_000.0, 12, "Laptop"))])


def prior_only_business():
    """The rental places property; the business holds a prior-year asset."""
    return fx.multi_scenario(
        rental_assets=(fx.equipment(9_000.0, 3),),
        businesses=[fx.business(PRIOR_EQUIPMENT)])


def stated_businesses(*stated: float):
    """The rental places property; each business states one figure. (20,000
    of receipts apiece keeps two businesses under the Form 8995 threshold.)"""
    return fx.multi_scenario(
        rental_assets=(fx.equipment(9_000.0, 3),),
        businesses=[fx.business(stated=amount, gross_receipts=20_000.0)
                    for amount in stated],
        acknowledges_stated=True)


def all_real_property():
    """Each activity places only real property this year."""
    return fx.multi_scenario(
        rental_assets=(DepreciableAsset(
            description="Rental building",
            date_placed_in_service=date(fx.YEAR, 1, 15), basis=200_000.0,
            recovery_class="27.5-year"),),
        businesses=[fx.business(DepreciableAsset(
            description="Shop", date_placed_in_service=date(fx.YEAR, 6, 10),
            basis=100_000.0, recovery_class="39-year"))])


def convention(scenario) -> dw.ConventionInputs:
    return dw.ConventionInputs(
        activities=resolver.mid_quarter_bases_by_activity(scenario),
        mid_quarter=resolver.mid_quarter_applies(scenario))


def _printed(pdf_path, field_path) -> int:
    fields = PdfReader(str(pdf_path)).get_fields() or {}
    return int(str(fields[field_path].get("/V")).replace(",", ""))


class _Case(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.tmp = Path(tmp.name)
        self._n = 0

    def emit_with_handoff(self, scenario):
        """Emit the whole return; ``(emitted, resolved, printed)``, the last
        two being exactly what the orchestrator handed the workbook writer."""
        self._n += 1
        handed = []
        real = dw.write_depreciation_audit

        def capture(resolved, printed, out_path):
            handed.append((resolved, printed))
            return real(resolved, printed, out_path)

        orchestrator = ReturnOrchestrator(
            spreadsheets_dir=SPREADSHEETS_DIR,
            work_dir=self.tmp / f"work{self._n}")
        with mock.patch.object(dw, "write_depreciation_audit", capture):
            _results, emitted = orchestrator.run_full_return(
                scenario, self.tmp / f"out{self._n}")
        [(resolved, printed)] = handed
        return emitted, resolved, printed

    def emit(self, scenario):
        emitted, _resolved, _printed_figures = self.emit_with_handoff(scenario)
        return emitted, openpyxl.load_workbook(emitted["depreciation_audit"])

    def claims(self, workbook) -> list[str]:
        sheet = workbook[TIE]
        return [sheet[f"A{r}"].value for r in range(2, sheet.max_row + 1)]

    def value(self, workbook, sheet: str, coordinate: str):
        return dw.cell_value(workbook, sheet, coordinate)


class LayoutTests(_Case):
    def test_sheet_order_on_a_multi_activity_return(self):
        _emitted, workbook = self.emit(over())
        self.assertEqual(workbook.sheetnames, [
            "Tie-outs", "Mid-quarter test", "Assets rental_0",
            "Years rental_0", "Assets sch_c_0", "Years sch_c_0"])

    def test_each_activity_tab_holds_only_its_own_assets(self):
        _emitted, workbook = self.emit(over())
        rental, business = workbook["Assets rental_0"], workbook[
            "Assets sch_c_0"]
        self.assertEqual(
            [rental[f"C{r}"].value for r in range(2, rental.max_row + 1)],
            ["Furniture", "Appliances"])
        self.assertEqual(
            [business[f"C{r}"].value
             for r in range(2, business.max_row + 1)], ["Laptop"])
        self.assertEqual(
            [rental[f"B{r}"].value for r in (2, 3)],
            ["rental property #0 ('100 Example Street')"] * 2)
        self.assertEqual(
            business["B2"].value, "Schedule C business #0 ('Consulting')")
        # Each tab's deductions, under the mirror: 35% and 5% of the
        # rental's two, 5% of the business's one.
        self.assertEqual(
            [self.value(workbook, "Assets rental_0", f"K{r}")
             for r in (2, 3)], [2_100, 150])
        self.assertEqual(self.value(workbook, "Assets sch_c_0", "K2"), 150)
        # A tab's formulas stay on its own pair of sheets.
        self.assertEqual(rental["H2"].value, "='Years rental_0'!D2")
        self.assertEqual(business["H2"].value, "='Years sch_c_0'!D2")
        self.assertEqual(
            workbook["Years sch_c_0"]["A2"].value, "='Assets sch_c_0'!C2")

    def test_single_activity_return_keeps_todays_three_sheets(self):
        _emitted, workbook = self.emit(fx.scenario())
        self.assertEqual(
            workbook.sheetnames, ["Tie-outs", "Assets", "Year-by-year"])

    def test_the_481a_sheet_stays_last_and_reads_the_activitys_own_tab(self):
        _emitted, resolved, printed = self.emit_with_handoff(
            prior_only_business())
        figures = dw.Form3115Figures(
            line_26=700.0, year_of_change=fx.YEAR, assets=(
                ("Equipment", date(2023, 6, 1), 5_900.0),))
        path = dw.write_depreciation_audit(
            resolved, dataclasses.replace(printed, form_3115=figures),
            self.tmp / "with_481a.xlsx")
        workbook = openpyxl.load_workbook(path)
        self.assertEqual(workbook.sheetnames, [
            "Tie-outs", "Mid-quarter test", "Assets rental_0",
            "Years rental_0", "Assets sch_c_0", "Years sch_c_0", "481(a)"])
        # Allowable through 2024 for the business's 2023 asset: 2,000 +
        # 3,200, the 2024 cumulative on ITS year-by-year tab (row 3).
        self.assertEqual(
            workbook["481(a)"]["D2"].value, "='Years sch_c_0'!F3")
        self.assertEqual(self.value(workbook, "481(a)", "D2"), 5_200)
        self.assertEqual(self.value(workbook, "481(a)", "E2"), 700)


class TieOutTests(_Case):
    def test_one_line_22_row_per_copy_against_its_printed_value(self):
        emitted, workbook = self.emit(over())
        sheet = workbook[TIE]
        self.assertEqual(self.claims(workbook), [
            "Form 4562 copy rental_0: Form 4562 line 19b, column (g): "
            "depreciation deduction",
            "Form 4562 copy rental_0: Form 4562 line 22: total depreciation",
            "Form 4562 copy sch_c_0: Form 4562 line 19b, column (g): "
            "depreciation deduction",
            "Form 4562 copy sch_c_0: Form 4562 line 22: total depreciation",
            "Schedule E line 18: depreciation, rental property #0 "
            "('100 Example Street')",
            "Schedule C line 13 (business 1): depreciation, Schedule C "
            "business #0 ('Consulting')",
            "Whole return: Form 4562 copies' line 22 plus no-copy "
            "activities' schedule lines = total depreciation",
        ])
        # Each copy's line 22 is computed from its own tab...
        self.assertEqual(sheet["B3"].value, "=SUM('Assets rental_0'!K2:K3)")
        self.assertEqual(sheet["B5"].value, "=SUM('Assets sch_c_0'!K2:K2)")
        self.assertEqual(
            [self.value(workbook, TIE, c) for c in ("B3", "B5")],
            [2_250, 150])
        # ...and set against the value printed on THAT copy.
        self.assertEqual(sheet["C3"].value, _printed(
            emitted["f4562_rental_0"], F4562_LINE_22))
        self.assertEqual(sheet["C5"].value, _printed(
            emitted["f4562_sch_c_0"], F4562_LINE_22))
        self.assertEqual([sheet["C3"].value, sheet["C5"].value], [2_250, 150])
        self.assertEqual(
            [self.value(workbook, TIE, f"F{r}") for r in range(2, 9)],
            ["PASS"] * 7)
        self.assertEqual(sheet.max_row, 8)

    def test_no_copy_activities_are_listed_by_recon_key_with_a_reason(self):
        with self.subTest(shape="prior-year assets only"):
            _emitted, workbook = self.emit(prior_only_business())
            sheet = workbook[TIE]
            self.assertEqual(
                sheet["A6"].value,
                "Whole return: sch_c_0 has no Form 4562 copy (no placement "
                "this year): Schedule C line 13 (business 1): depreciation, "
                "Schedule C business #0 ('Consulting')")
            self.assertEqual(sheet["C6"].value, 1_920)
            self.assertEqual(
                [c for c in self.claims(workbook) if "no Form 4562 copy" in c],
                [sheet["A6"].value])
            # The rental has a copy and is not listed as lacking one.
            self.assertFalse(any(
                "rental_0 has no Form 4562 copy" in c
                for c in self.claims(workbook)))
        with self.subTest(shape="stated figure"):
            _emitted, workbook = self.emit(stated_businesses(1_200.0))
            sheet = workbook[TIE]
            self.assertEqual(
                sheet["A5"].value,
                "Whole return: sch_c_0 has no Form 4562 copy (stated "
                "figure): Schedule C line 13 (business 1): depreciation, "
                "Schedule C business #0 ('Consulting')")
            self.assertEqual(sheet["C5"].value, 1_200)
        with self.subTest(shape="two stated businesses, one name"):
            # Phase K of the f8995 per-source-rows lane: Form 8995 refuses
            # two Schedule C businesses until phase C (the half-SE-tax
            # allocation ruling). This assertion goes red when phase C lands;
            # then replace it with
            # self._assert_two_stated_businesses_tie_out(
            #     self.emit(stated_businesses(1_200.0, 800.0))[1]).
            with self.assertRaises(NotImplementedError) as ctx:
                self.emit(stated_businesses(1_200.0, 800.0))
            self.assertIn("Form 8995 per-source rows", str(ctx.exception))

    def _assert_two_stated_businesses_tie_out(self, workbook):
        # The two-business shape's tie-out, held until phase C (above).
        sheet = workbook[TIE]
        listed = [(sheet[f"A{r}"].value, sheet[f"C{r}"].value)
                  for r in (5, 6)]
        self.assertEqual(listed, [
            ("Whole return: sch_c_0 has no Form 4562 copy (stated "
             "figure): Schedule C line 13 (business 1): depreciation, "
             "Schedule C business #0 ('Consulting')", 1_200),
            ("Whole return: sch_c_1 has no Form 4562 copy (stated "
             "figure): Schedule C line 13 (business 2): depreciation, "
             "Schedule C business #1 ('Consulting')", 800)])

    def test_an_asset_rental_beside_a_stated_business_is_multi_activity_with_one_tab_pair(
            self):
        _emitted, workbook = self.emit(stated_businesses(1_200.0))
        self.assertEqual(workbook.sheetnames, [
            "Tie-outs", "Mid-quarter test", "Assets rental_0",
            "Years rental_0"])
        pooled = workbook[POOLED]
        self.assertEqual(
            [pooled[f"A{r}"].value for r in (2, 3)], ["rental_0", "sch_c_0"])
        self.assertEqual(pooled["E3"].value, STATED_NOTE)

    def test_an_asset_rental_beside_a_no_depreciation_business_is_single_activity(
            self):
        _emitted, workbook = self.emit(fx.multi_scenario(
            rental_assets=(fx.equipment(9_000.0, 3),),
            businesses=[fx.business()]))
        self.assertEqual(
            workbook.sheetnames, ["Tie-outs", "Assets", "Year-by-year"])

    def test_whole_return_row_ties_copies_plus_no_copy_lines_to_total_depreciation(
            self):
        _emitted, resolved, printed = self.emit_with_handoff(
            prior_only_business())

        def whole_return_row(figures, name):
            workbook = openpyxl.load_workbook(dw.write_depreciation_audit(
                resolved, figures, self.tmp / name))
            sheet = workbook[TIE]
            self.assertEqual(sheet.max_row, 7)
            self.assertEqual(sheet["B7"].value, "=B3+C6")
            self.assertEqual(sheet["C7"].value, "=C4+C6")
            return [self.value(workbook, TIE, f"{c}7") for c in "BCEF"]

        # The rental copy's line 22 (1,800) plus the business's Schedule C
        # line (1,920) is the return's total depreciation.
        self.assertEqual(
            whole_return_row(printed, "true.xlsx"),
            [3_720, 3_720, 0, "PASS"])
        # Twin: Schedule E prints one dollar more than the copy's line 22.
        lines = dict(printed.activity_lines)
        lines[RENTAL] += 1
        self.assertEqual(
            whole_return_row(
                dataclasses.replace(printed, activity_lines=lines),
                "off.xlsx"),
            [3_720, 3_721, -1, "FAIL"])


class PooledTestSheetTests(_Case):
    def test_rows_carry_each_activitys_integer_cents(self):
        _emitted, workbook = self.emit(over())
        sheet = workbook[POOLED]
        self.assertEqual(
            [[sheet[f"{c}{r}"].value for c in "ABCD"] for r in (2, 3)], [
                ["rental_0", "rental property #0 ('100 Example Street')",
                 300_000, 900_000],
                ["sch_c_0", "Schedule C business #0 ('Consulting')",
                 300_000, 300_000]])
        for coordinate in ("C2", "D2", "C3", "D3"):
            with self.subTest(cell=coordinate):
                self.assertIs(type(sheet[coordinate].value), int)
        self.assertEqual(sheet["C4"].value, "=SUM(C2:C3)")
        self.assertEqual(sheet["D4"].value, "=SUM(D2:D3)")
        self.assertEqual(
            [self.value(workbook, POOLED, c) for c in ("C4", "D4")],
            [600_000, 1_200_000])
        self.assertIs(self.value(workbook, POOLED, "C5"), True)

    def test_the_deciding_cell_is_integer_cross_multiplication(self):
        _emitted, workbook = self.emit(over())
        formula = workbook[POOLED]["C5"].value
        self.assertEqual(formula, "=C4*100>40*D4")
        self.assertIn("*100>40*", formula)
        self.assertNotIn("/", formula)
        # The percentage beside it is display only, and says so.
        self.assertEqual(workbook[POOLED]["C6"].value, "=IF(D4=0,0,C4/D4)")
        self.assertIn("display only", workbook[POOLED]["A6"].value)

    def test_the_deciding_cell_agrees_with_the_engine_on_exact_boundary_pairs(
            self):
        """The exact-40% cent pairs of tests/test_mid_quarter_convention.py
        (late = 2k cents, early = 3k cents), the early basis on the rental
        and the late one on the business, and one cent either side."""
        for k in range(1_000_001, 1_000_400):
            for extra, trips in ((-1, False), (0, False), (1, True)):
                with self.subTest(k=k, extra=extra):
                    scenario = fx.multi_scenario(
                        rental_assets=(fx.equipment(3 * k / 100, 3),),
                        businesses=[fx.business(
                            fx.equipment((2 * k + extra) / 100, 11))])
                    workbook, reasons = dw._build(
                        resolver.audit_trail(scenario),
                        dw.PrintedFigures(convention=convention(scenario)))
                    self.assertEqual(reasons, [])
                    self.assertEqual(
                        [self.value(workbook, POOLED, c)
                         for c in ("C4", "D4")],
                        [2 * k + extra, 5 * k + extra])
                    decided = self.value(workbook, POOLED, "C5")
                    self.assertIs(decided, trips)
                    self.assertIs(
                        decided, resolver.mid_quarter_applies(scenario))

    def test_all_real_property_placements_decide_false_without_dividing(self):
        scenario = all_real_property()
        self.assertEqual(dw.unbuildable_reasons(scenario), [])
        _emitted, workbook = self.emit(scenario)
        sheet = workbook[POOLED]
        self.assertEqual(
            [[sheet[f"{c}{r}"].value for c in "CD"] for r in (2, 3)],
            [[0, 0], [0, 0]])
        self.assertEqual(
            [self.value(workbook, POOLED, c) for c in ("C4", "D4")], [0, 0])
        self.assertIs(self.value(workbook, POOLED, "C5"), False)
        self.assertEqual(self.value(workbook, POOLED, "C6"), 0)

    def test_a_stated_mode_activity_is_a_row_that_says_so(self):
        _emitted, workbook = self.emit(stated_businesses(1_200.0))
        sheet = workbook[POOLED]
        self.assertEqual(
            [sheet[f"{c}3"].value for c in "ABCDE"],
            ["sch_c_0", "Schedule C business #0 ('Consulting')", None, None,
             STATED_NOTE])
        # Twin: the asset-mode row carries integers and no such note.
        self.assertEqual(
            [sheet[f"{c}2"].value for c in "ABCDE"],
            ["rental_0", "rental property #0 ('100 Example Street')", 0,
             900_000, None])
        self.assertEqual(
            [self.value(workbook, POOLED, c) for c in ("C4", "D4")],
            [0, 900_000])
        self.assertIs(self.value(workbook, POOLED, "C5"), False)

    def test_the_precision_bound_is_stated_once(self):
        _emitted, workbook = self.emit(over())
        note = workbook[POOLED]["A8"].value
        self.assertIn("2^53", note)
        self.assertEqual(dw.__doc__.count("2^53"), 1)
        self.assertEqual(
            [cell.coordinate for row in workbook[POOLED].iter_rows()
             for cell in row
             if isinstance(cell.value, str) and "2^53" in cell.value],
            ["A8"])


class FailClosedTests(_Case):
    REFUSAL = "depreciation audit workbook cannot be built"

    def setUp(self):
        super().setUp()
        _emitted, self.resolved, self.printed = self.emit_with_handoff(over())
        self.out = self.tmp / "audit.xlsx"

    def assert_refuses(self, printed, pattern, resolved=None):
        with self.assertRaisesRegex(NotImplementedError, pattern) as caught:
            dw.write_depreciation_audit(
                self.resolved if resolved is None else resolved, printed,
                self.out)
        self.assertIn(self.REFUSAL, str(caught.exception))
        self.assertFalse(self.out.exists())

    def assert_writes(self, printed, resolved=None):
        path = dw.write_depreciation_audit(
            self.resolved if resolved is None else resolved, printed,
            self.out)
        self.assertTrue(path.is_file())

    def test_a_boolean_that_disagrees_with_the_deciding_cell_refuses(self):
        self.assertIs(self.printed.convention.mid_quarter, True)
        flipped = dataclasses.replace(
            self.printed, convention=dataclasses.replace(
                self.printed.convention, mid_quarter=False))
        self.assert_refuses(
            flipped, r"Mid-quarter test deciding cell \('Mid-quarter test'!C5\)")
        # Twin: the true boolean writes the file.
        self.assert_writes(self.printed)

    def test_a_copy_whose_printed_line_22_disagrees_with_its_tab_refuses(
            self):
        copies = dict(self.printed.f4562_copies)
        copies["sch_c_0"] = dataclasses.replace(
            copies["sch_c_0"], line_22=copies["sch_c_0"].line_22 + 1)
        self.assert_refuses(
            dataclasses.replace(self.printed, f4562_copies=copies),
            r"Form 4562 copy sch_c_0: .*151")
        self.assert_writes(self.printed)

    def test_a_placing_activity_with_no_copy_handed_off_refuses(self):
        copies = dict(self.printed.f4562_copies)
        del copies["sch_c_0"]
        self.assert_refuses(
            dataclasses.replace(self.printed, f4562_copies=copies),
            r"Form 4562 copy sch_c_0: .*no figures")
        self.assert_writes(self.printed)

    def test_a_copy_for_an_activity_that_places_nothing_refuses(self):
        copies = dict(self.printed.f4562_copies)
        copies["sch_c_7"] = copies["sch_c_0"]
        self.assert_refuses(
            dataclasses.replace(self.printed, f4562_copies=copies),
            r"Form 4562 copy sch_c_7: ")
        self.assert_writes(self.printed)

    def test_a_multi_activity_trail_without_convention_inputs_refuses(self):
        """Two activities on the trail and no `ConventionInputs`: the
        mandatory pooled sheet cannot be laid out, and nothing falls back to
        the single-activity layout."""
        self.assert_refuses(
            dataclasses.replace(self.printed, convention=None),
            r"ConventionInputs")
        # Twin: a single-activity trail without them writes today's layout.
        single = resolver.audit_trail(fx.scenario())
        self.assert_writes(dw.PrintedFigures(), resolved=single)
        self.assertEqual(
            openpyxl.load_workbook(self.out).sheetnames,
            ["Tie-outs", "Assets", "Year-by-year"])


class LazyIfTests(unittest.TestCase):
    """The mirror's IF evaluates only the branch it takes, as a spreadsheet
    does. A division by zero in the branch NOT taken is not an error; one in
    the branch taken still is."""

    def _workbook(self, formula: str):
        workbook = openpyxl.Workbook()
        sheet = workbook.active
        sheet.title = "S"
        sheet["A1"] = 0
        sheet["A2"] = 7
        sheet["B1"] = formula
        return workbook

    def test_the_untaken_branch_is_not_evaluated(self):
        workbook = self._workbook("=IF(A1=0,0,A2/A1)")
        self.assertEqual(dw.cell_value(workbook, "S", "B1"), 0)
        self.assertEqual(dw._mirror_mismatches(workbook, []), [])
        # ...in either position.
        workbook = self._workbook("=IF(A1<>0,A2/A1,5)")
        self.assertEqual(dw.cell_value(workbook, "S", "B1"), 5)
        self.assertEqual(dw._mirror_mismatches(workbook, []), [])

    def test_a_division_by_zero_in_the_taken_branch_still_refuses(self):
        for formula in ("=IF(A1=0,A2/A1,0)", "=IF(A1<>0,0,A2/A1)"):
            with self.subTest(formula=formula):
                workbook = self._workbook(formula)
                with self.assertRaises(ArithmeticError):
                    dw.cell_value(workbook, "S", "B1")
                reasons = dw._mirror_mismatches(workbook, [])
                self.assertEqual(len(reasons), 1)
                self.assertIn(
                    "the workbook formula at S!B1 cannot be checked",
                    reasons[0])

    def test_a_skipped_branch_is_still_read_as_a_formula(self):
        """Laziness skips evaluation, not parsing: text the mirror cannot
        read is refused even in the branch not taken."""
        workbook = self._workbook("=IF(A1=0,0,NOSUCH(A2))")
        self.assertEqual(len(dw._mirror_mismatches(workbook, [])), 1)

    def test_a_nested_if_inside_a_skipped_branch_is_skipped_whole(self):
        workbook = self._workbook("=IF(A1=0,3,IF(A2/A1>1,A2/A1,A2/A1))")
        self.assertEqual(dw.cell_value(workbook, "S", "B1"), 3)

    def test_the_taken_branch_value_is_returned(self):
        workbook = self._workbook("=IF(A2>A1,A2*2,A2/A1)")
        self.assertEqual(dw.cell_value(workbook, "S", "B1"), 14)


class EmitTests(_Case):
    def test_the_hand_off_carries_one_bundle_per_copy_and_the_loops_boolean(
            self):
        for shape, answer in ((over, True), (under, False)):
            with self.subTest(shape=shape.__name__):
                emitted, _resolved, printed = self.emit_with_handoff(shape())
                self.assertEqual(
                    sorted(printed.f4562_copies), ["rental_0", "sch_c_0"])
                self.assertIs(printed.convention.mid_quarter, answer)
                self.assertEqual(
                    [row.recon_key for row in printed.convention.activities],
                    ["rental_0", "sch_c_0"])
                for key in ("rental_0", "sch_c_0"):
                    self.assertEqual(
                        printed.f4562_copies[key].line_22,
                        _printed(emitted[f"f4562_{key}"], F4562_LINE_22))
        # The over shape's bundles, in full.
        _emitted, _resolved, printed = self.emit_with_handoff(over())
        self.assertEqual(printed.f4562_copies, {
            "rental_0": dw.CopyFigures(
                line_17=None, line_19={"b": 2_250}, line_22=2_250),
            "sch_c_0": dw.CopyFigures(
                line_17=None, line_19={"b": 150}, line_22=150)})
        self.assertEqual(
            dict(printed.activity_lines), {RENTAL: 2_250, BUSINESS: 150})

    def test_a_single_copy_is_keyed_by_its_recon_key(self):
        emitted, _resolved, printed = self.emit_with_handoff(fx.scenario())
        self.assertEqual(list(printed.f4562_copies), ["rental_0"])
        self.assertEqual(
            printed.f4562_copies["rental_0"], dw.CopyFigures(
                line_17=fx.LINE_17, line_19={"c": fx.LINE_19C},
                line_22=fx.LINE_22))
        self.assertEqual(
            printed.f4562_copies["rental_0"].line_22,
            _printed(emitted["f4562"], F4562_LINE_22))
        self.assertEqual(
            [row.recon_key for row in printed.convention.activities],
            ["rental_0"])
        self.assertIs(printed.convention.mid_quarter, False)


if __name__ == "__main__":
    unittest.main()
