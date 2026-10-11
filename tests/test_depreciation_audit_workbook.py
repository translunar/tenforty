"""The depreciation audit workbook: structure, formula text, and the mirror.

No spreadsheet engine runs anywhere here. The workbook is written with
openpyxl, read back with openpyxl (formulas as text), and its formulas are
evaluated by the module's own Python evaluator against the hand figures in
tests/_audit_workbook_fixtures.py.

Fixture layout: Assets rows 2 (building), 3 (appliance), 4 (furniture);
Year-by-year rows 2-8 (building, 2019-2025), 9-11 (appliance, 2023-2025),
12 (furniture, 2025); Tie-outs rows 2 (line 17), 3 (line 19c), 4 (line 22),
5 (Schedule E line 18).
"""
import dataclasses
import tempfile
import unittest
import unittest.mock
from datetime import date
from pathlib import Path

import openpyxl

from tenforty.audit import depreciation_workbook as dw
from tenforty.forms.depreciation import resolver
from tenforty.models import DepreciableAsset
from tests import _audit_workbook_fixtures as fx

RENTAL = ("rental_properties", 0)


def true_printed(**changes) -> dw.PrintedFigures:
    """The figures the forms print for the default fixture."""
    return dataclasses.replace(dw.PrintedFigures(
        f4562_line_17=fx.LINE_17, f4562_line_19={"c": fx.LINE_19C},
        f4562_line_22=fx.LINE_22,
        activity_lines={RENTAL: fx.SCH_E_LINE_18}), **changes)


class _WorkbookCase(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.tmp = Path(tmp.name)

    def write(self, scenario=None, printed=None, name="audit.xlsx"):
        """Write the workbook and read it back from disk."""
        scenario = fx.scenario() if scenario is None else scenario
        path = dw.write_depreciation_audit(
            resolver.audit_trail(scenario),
            true_printed() if printed is None else printed,
            self.tmp / name)
        return openpyxl.load_workbook(path)

    def formulas(self, workbook, sheet: str) -> dict[str, str]:
        return {cell.coordinate: cell.value
                for row in workbook[sheet].iter_rows() for cell in row
                if cell.data_type == "f"}


class StructureTests(_WorkbookCase):
    def test_sheet_names_and_order(self):
        self.assertEqual(
            self.write().sheetnames, ["Tie-outs", "Assets", "Year-by-year"])

    def test_tie_outs_is_first_sheet(self):
        workbook = self.write()
        self.assertEqual(workbook.worksheets[0].title, "Tie-outs")
        self.assertEqual(workbook.active.title, "Tie-outs")

    def test_returns_the_path_it_wrote(self):
        path = dw.write_depreciation_audit(
            resolver.audit_trail(fx.scenario()), true_printed(),
            self.tmp / "depreciation_audit_2025.xlsx")
        self.assertEqual(path, self.tmp / "depreciation_audit_2025.xlsx")
        self.assertTrue(path.is_file())

    def test_assets_header_row(self):
        row = [c.value for c in self.write()["Assets"][1]]
        self.assertEqual(row, [
            "Activity #", "Activity", "Description", "In service", "Method",
            "Life", "Convention", "Rate", "Basis", "Prior accumulated",
            "Deduction", "Accumulated after", "4562 line", "Provenance"])

    def test_year_by_year_header_row(self):
        row = [c.value for c in self.write()["Year-by-year"][1]]
        self.assertEqual(row, [
            "Asset", "Recovery year", "Tax year", "Rate", "Deduction",
            "Cumulative", "Source"])

    def test_tie_outs_header_row(self):
        row = [c.value for c in self.write()["Tie-outs"][1]]
        self.assertEqual(row, [
            "Claim", "Computed", "Printed on form",
            "Source of printed figure", "Difference", "Verdict"])

    def test_one_assets_row_per_fixture_asset(self):
        sheet = self.write()["Assets"]
        self.assertEqual(sheet.max_row, 4)
        self.assertEqual(
            [sheet[f"C{r}"].value for r in (2, 3, 4)],
            ["Rental building", "Appliance set", "Furniture"])

    def test_assets_inputs_are_transcribed(self):
        sheet = self.write()["Assets"]
        self.assertEqual(
            [sheet[f"I{r}"].value for r in (2, 3, 4)],
            [200_000, 10_000, 7_000])
        self.assertEqual(
            [sheet[f"J{r}"].value for r in (2, 3, 4)], [43_330, 5_200, 0])
        self.assertEqual(
            [sheet[f"A{r}"].value for r in (2, 3, 4)], [1, 1, 1])
        self.assertEqual(
            [sheet[f"E{r}"].value for r in (2, 3, 4)],
            ["S/L", "200 DB", "200 DB"])
        self.assertEqual(
            [sheet[f"F{r}"].value for r in (2, 3, 4)],
            ["27.5-year", "5-year", "7-year"])
        self.assertEqual(
            [sheet[f"G{r}"].value for r in (2, 3, 4)],
            ["mid-month", "half-year", "half-year"])
        self.assertEqual(
            [sheet[f"M{r}"].value for r in (2, 3, 4)],
            ["line 17", "line 17", "line 19c"])
        self.assertEqual(sheet["D2"].value.date(), date(2019, 1, 15))

    def test_year_by_year_rows_and_rates(self):
        sheet = self.write()["Year-by-year"]
        self.assertEqual(sheet.max_row, 12)
        self.assertEqual(
            [sheet[f"C{r}"].value for r in range(2, 13)],
            [2019, 2020, 2021, 2022, 2023, 2024, 2025, 2023, 2024, 2025,
             2025])
        self.assertEqual(
            [sheet[f"B{r}"].value for r in range(2, 13)],
            [1, 2, 3, 4, 5, 6, 7, 1, 2, 3, 1])
        self.assertEqual(
            [sheet[f"D{r}"].value for r in range(2, 13)],
            [0.03485] + [0.03636] * 6 + [0.2, 0.32, 0.192] + [0.1429])

    def test_mid_quarter_convention_names_its_quarter(self):
        late = dataclasses.replace(
            fx.FURNITURE, date_placed_in_service=date(2025, 11, 3))
        workbook = self.write(
            fx.scenario(rental_assets=(late,)),
            dw.PrintedFigures(
                f4562_line_19={"c": 250}, f4562_line_22=250,
                activity_lines={RENTAL: 250}))
        self.assertEqual(
            workbook["Assets"]["G2"].value, "mid-quarter (quarter 4)")
        self.assertEqual(
            workbook["Year-by-year"]["G2"].value, "Pub 946 Table A-5")
        self.assertEqual(workbook["Year-by-year"]["D2"].value, 0.0357)


class FormulaStringTests(_WorkbookCase):
    def test_assets_formula_strings(self):
        self.assertEqual(self.formulas(self.write(), "Assets"), {
            "H2": "='Year-by-year'!D8",
            "K2": "=ROUND(MAX(0,MIN(I2*H2,I2-J2)),0)",
            "L2": "=J2+K2",
            "H3": "='Year-by-year'!D11",
            "K3": "=ROUND(MAX(0,MIN(I3*H3,I3-J3)),0)",
            "L3": "=J3+K3",
            "H4": "='Year-by-year'!D12",
            "K4": "=ROUND(MAX(0,MIN(I4*H4,I4-J4)),0)",
            "L4": "=J4+K4",
        })

    def test_year_by_year_formula_strings(self):
        formulas = self.formulas(self.write(), "Year-by-year")
        # First recovery year of a block: no cumulative above it.
        self.assertEqual(formulas["A2"], "=Assets!C2")
        self.assertEqual(
            formulas["E2"],
            "=ROUND(MAX(0,MIN(Assets!$I$2*D2,Assets!$I$2)),0)")
        self.assertEqual(formulas["F2"], "=E2")
        # Later years: the ceiling is basis less the cumulative above.
        self.assertEqual(
            formulas["E3"],
            "=ROUND(MAX(0,MIN(Assets!$I$2*D3,Assets!$I$2-F2)),0)")
        self.assertEqual(formulas["F3"], "=F2+E3")
        self.assertEqual(
            formulas["E8"],
            "=ROUND(MAX(0,MIN(Assets!$I$2*D8,Assets!$I$2-F7)),0)")
        self.assertEqual(formulas["F8"], "=F7+E8")
        # The next asset's block starts over, against its own basis.
        self.assertEqual(formulas["A9"], "=Assets!C3")
        self.assertEqual(
            formulas["E9"],
            "=ROUND(MAX(0,MIN(Assets!$I$3*D9,Assets!$I$3)),0)")
        self.assertEqual(formulas["F9"], "=E9")
        self.assertEqual(
            formulas["E10"],
            "=ROUND(MAX(0,MIN(Assets!$I$3*D10,Assets!$I$3-F9)),0)")
        self.assertEqual(
            formulas["E12"],
            "=ROUND(MAX(0,MIN(Assets!$I$4*D12,Assets!$I$4)),0)")
        self.assertEqual(
            sorted(formulas),
            sorted(f"{c}{r}" for c in "AEF" for r in range(2, 13)))

    def test_tie_out_formula_strings(self):
        self.assertEqual(self.formulas(self.write(), "Tie-outs"), {
            "B2": '=SUMIF(Assets!M2:M4,"line 17",Assets!K2:K4)',
            "B3": '=SUMIF(Assets!M2:M4,"line 19c",Assets!K2:K4)',
            "B4": "=SUM(Assets!K2:K4)",
            "B5": "=SUMIF(Assets!A2:A4,1,Assets!K2:K4)",
            **{f"E{r}": f"=B{r}-C{r}" for r in (2, 3, 4, 5)},
            **{f"F{r}": f'=IF(ABS(E{r})<0.005,"PASS","FAIL")'
               for r in (2, 3, 4, 5)},
        })


class ProvenanceTests(_WorkbookCase):
    def test_every_assets_input_names_its_scenario_field(self):
        sheet = self.write()["Assets"]
        base = "rental_properties[0].depreciable_assets"
        self.assertEqual(sheet["N2"].value, (
            f"Basis: {base}[0].basis. "
            f"Prior accumulated: {base}[0].prior_depreciation"))
        self.assertEqual(sheet["N3"].value, (
            f"Basis: {base}[1].basis. "
            f"Prior accumulated: {base}[1].prior_depreciation"))
        self.assertEqual(sheet["N4"].value, (
            f"Basis: {base}[2].basis. "
            f"Prior accumulated: none (placed in service this year)"))

    def test_an_acknowledged_prior_is_labelled_as_stated(self):
        ahead = dataclasses.replace(
            fx.APPLIANCE, prior_depreciation=9_500.0,
            acknowledges_prior_depreciation_as_stated=True)
        sheet = self.write(
            fx.scenario(rental_assets=(ahead,)),
            dw.PrintedFigures(activity_lines={RENTAL: 500}))["Assets"]
        self.assertEqual(sheet["J2"].value, 9_500)
        self.assertEqual(sheet["N2"].value, (
            "Basis: rental_properties[0].depreciable_assets[0].basis. "
            "Prior accumulated: rental_properties[0].depreciable_assets[0]."
            "prior_depreciation (stated, differs from the tables: "
            "acknowledges_prior_depreciation_as_stated)"))

    def test_no_numeric_input_is_without_a_label(self):
        workbook = self.write()
        assets = workbook["Assets"]
        checked = 0
        for r in range(2, assets.max_row + 1):
            for column in "IJ":
                self.assertIsInstance(
                    assets[f"{column}{r}"].value, (int, float))
                checked += 1
            self.assertTrue(assets[f"N{r}"].value)
        self.assertEqual(checked, 6)
        years = workbook["Year-by-year"]
        for r in range(2, years.max_row + 1):
            self.assertIsInstance(years[f"D{r}"].value, float)
            self.assertRegex(years[f"G{r}"].value, r"^Pub 946 Table A-\d")
        tie_outs = workbook["Tie-outs"]
        for r in range(2, tie_outs.max_row + 1):
            self.assertIsInstance(tie_outs[f"C{r}"].value, int)
            self.assertRegex(tie_outs[f"D{r}"].value, r" as printed$")

    def test_year_sources_name_the_table_and_month(self):
        sheet = self.write()["Year-by-year"]
        self.assertEqual(sheet["G2"].value, "Pub 946 Table A-6, month 1")
        self.assertEqual(sheet["G9"].value, "Pub 946 Table A-1")
        self.assertEqual(sheet["G12"].value, "Pub 946 Table A-1")


class MirrorRuleTests(_WorkbookCase):
    """Every formula cell, evaluated in Python from the cells as written to
    disk, equals the hand figure."""

    def value(self, workbook, sheet, coordinate):
        return dw.cell_value(workbook, sheet, coordinate)

    def test_assets_formulas_give_the_hand_figures(self):
        workbook = self.write()
        self.assertEqual(
            [self.value(workbook, "Assets", f"H{r}") for r in (2, 3, 4)],
            [0.03636, 0.192, 0.1429])
        self.assertEqual(
            [self.value(workbook, "Assets", f"K{r}") for r in (2, 3, 4)],
            [7_272, 1_920, 1_000])
        self.assertEqual(
            [self.value(workbook, "Assets", f"L{r}") for r in (2, 3, 4)],
            [50_602, 7_120, 1_000])

    def test_year_by_year_formulas_give_the_hand_figures(self):
        workbook = self.write()
        self.assertEqual(
            [self.value(workbook, "Year-by-year", f"E{r}")
             for r in range(2, 13)],
            [6_970, 7_272, 7_272, 7_272, 7_272, 7_272, 7_272,
             2_000, 3_200, 1_920, 1_000])
        self.assertEqual(
            [self.value(workbook, "Year-by-year", f"F{r}")
             for r in range(2, 13)],
            [6_970, 14_242, 21_514, 28_786, 36_058, 43_330, 50_602,
             2_000, 5_200, 7_120, 1_000])
        self.assertEqual(
            [self.value(workbook, "Year-by-year", f"A{r}") for r in (2, 9, 12)],
            ["Rental building", "Appliance set", "Furniture"])

    def test_every_formula_cell_matches_the_resolver(self):
        scenario = fx.scenario()
        workbook = self.write(scenario)
        [activity] = resolver.audit_trail(scenario)
        year_row = 2
        checked = 0
        for r, asset in enumerate(activity.assets, start=2):
            self.assertEqual(
                self.value(workbook, "Assets", f"K{r}"), asset.current.amount)
            checked += 1
            for year in asset.years:
                self.assertEqual(
                    self.value(workbook, "Year-by-year", f"E{year_row}"),
                    year.amount)
                self.assertEqual(
                    self.value(workbook, "Year-by-year", f"F{year_row}"),
                    year.taken_before + year.amount)
                year_row += 1
                checked += 2
        self.assertEqual(checked, 3 + 2 * 11)

    def test_ceiling_bound_asset_gets_the_remaining_basis(self):
        # 9,500 already taken of 10,000: the table's 1,920 is held to 500.
        ahead = dataclasses.replace(
            fx.APPLIANCE, prior_depreciation=9_500.0,
            acknowledges_prior_depreciation_as_stated=True)
        workbook = self.write(
            fx.scenario(rental_assets=(ahead,)),
            dw.PrintedFigures(activity_lines={RENTAL: 500}))
        self.assertEqual(self.value(workbook, "Assets", "K2"), 500)
        self.assertEqual(self.value(workbook, "Assets", "L2"), 10_000)
        # The year rows stay the table reconstruction.
        self.assertEqual(self.value(workbook, "Year-by-year", "E4"), 1_920)

    def test_prior_above_basis_floors_at_zero(self):
        over = dataclasses.replace(
            fx.APPLIANCE, prior_depreciation=10_500.0,
            acknowledges_prior_depreciation_as_stated=True)
        workbook = self.write(
            fx.scenario(rental_assets=(over,)), dw.PrintedFigures())
        self.assertEqual(self.value(workbook, "Assets", "K2"), 0)
        self.assertEqual(self.value(workbook, "Tie-outs", "F2"), "PASS")

    def test_table_rounding_trimmed_in_the_last_year(self):
        # 3-year on a basis of 11: 33.33%, 44.45%, 14.81%, 7.41% round to
        # 4 + 5 + 2 + 1 = 12, a dollar over basis. The ceiling trims the
        # last year to 0, in the formula as in the engine.
        small = DepreciableAsset(
            description="Small tool", date_placed_in_service=date(2022, 5, 2),
            basis=11.0, recovery_class="3-year", prior_depreciation=11.0,
            no_bonus_or_section_179_history=True, convention="half-year")
        workbook = self.write(
            fx.scenario(rental_assets=(small,)), dw.PrintedFigures())
        self.assertEqual(
            [self.value(workbook, "Year-by-year", f"E{r}")
             for r in (2, 3, 4, 5)], [4, 5, 2, 0])
        self.assertEqual(self.value(workbook, "Year-by-year", "F5"), 11)
        self.assertEqual(self.value(workbook, "Assets", "K2"), 0)

    def test_half_dollar_product_rounds_up(self):
        # 2.5 x 20% = 0.5 exactly: one dollar (Python's round() gives 0).
        tiny = DepreciableAsset(
            description="Tiny", date_placed_in_service=date(2025, 2, 3),
            basis=2.5, recovery_class="5-year",
            no_bonus_or_section_179_history=True)
        workbook = self.write(
            fx.scenario(rental_assets=(tiny,)),
            dw.PrintedFigures(
                f4562_line_19={"b": 1}, f4562_line_22=1,
                activity_lines={RENTAL: 1}))
        self.assertEqual(self.value(workbook, "Assets", "K2"), 1)
        self.assertEqual(
            [self.value(workbook, "Tie-outs", f"F{r}") for r in (2, 3, 4)],
            ["PASS", "PASS", "PASS"])


class ElapsedScheduleTests(_WorkbookCase):
    """A fully depreciated asset still on the books renders: no rate, zero
    deduction, and the reason beside it."""

    OLD = DepreciableAsset(
        description="Old computer", date_placed_in_service=date(2015, 4, 1),
        basis=3_000.0, recovery_class="5-year", prior_depreciation=3_000.0,
        no_bonus_or_section_179_history=True, convention="half-year")

    def test_renders_with_no_rate_and_a_zero_deduction(self):
        workbook = self.write(
            fx.scenario(rental_assets=(self.OLD,)), dw.PrintedFigures())
        years = workbook["Year-by-year"]
        self.assertEqual(years.max_row, 12)        # 2015..2025
        self.assertEqual(years["C7"].value, 2020)
        self.assertEqual(years["D7"].value, 0.0576)
        self.assertEqual(years["G7"].value, "Pub 946 Table A-1")
        for r in range(8, 13):                     # 2021..2025
            with self.subTest(row=r):
                self.assertIsNone(years[f"D{r}"].value)
                self.assertEqual(
                    years[f"G{r}"].value,
                    "recovery period ended (no cell in Pub 946 Table A-1)")
                self.assertEqual(
                    dw.cell_value(workbook, "Year-by-year", f"E{r}"), 0)
        self.assertEqual(dw.cell_value(workbook, "Year-by-year", "F12"), 3_000)
        self.assertEqual(dw.cell_value(workbook, "Assets", "H2"), 0)
        self.assertEqual(dw.cell_value(workbook, "Assets", "K2"), 0)
        self.assertEqual(dw.cell_value(workbook, "Tie-outs", "F2"), "PASS")


class TieOutTests(_WorkbookCase):
    def claims(self, workbook) -> list[str]:
        sheet = workbook["Tie-outs"]
        return [sheet[f"A{r}"].value for r in range(2, sheet.max_row + 1)]

    def verdicts(self, workbook) -> list[str]:
        sheet = workbook["Tie-outs"]
        return [dw.cell_value(workbook, "Tie-outs", f"F{r}")
                for r in range(2, sheet.max_row + 1)]

    def test_tie_out_row_set(self):
        self.assertEqual(self.claims(self.write()), [
            "Form 4562 line 17: MACRS deductions for assets placed in "
            "service before 2025",
            "Form 4562 line 19c, column (g): depreciation deduction",
            "Form 4562 line 22: total depreciation",
            "Schedule E line 18: depreciation, rental property #0 "
            "('100 Example Street')",
        ])

    def test_printed_sources(self):
        sheet = self.write()["Tie-outs"]
        self.assertEqual(
            [sheet[f"D{r}"].value for r in (2, 3, 4, 5)], [
                "Form 4562 line 17 as printed",
                "Form 4562 line 19c column (g) as printed",
                "Form 4562 line 22 as printed",
                "Schedule E line 18 as printed"])

    def test_printed_column_holds_the_printed_values(self):
        # Figures deliberately NOT the engine's: column C is what it was
        # handed, never a recomputation.
        printed = dw.PrintedFigures(
            f4562_line_17=9_199, f4562_line_19={"c": 1_003},
            f4562_line_22=10_202, activity_lines={RENTAL: 10_207})
        sheet = self.write(printed=printed)["Tie-outs"]
        self.assertEqual(
            [sheet[f"C{r}"].value for r in (2, 3, 4, 5)],
            [9_199, 1_003, 10_202, 10_207])
        self.assertEqual(
            [sheet[f"C{r}"].data_type for r in (2, 3, 4, 5)], ["n"] * 4)

    def test_computed_column_gives_the_hand_figures(self):
        workbook = self.write()
        self.assertEqual(
            [dw.cell_value(workbook, "Tie-outs", f"B{r}")
             for r in (2, 3, 4, 5)],
            [fx.LINE_17, fx.LINE_19C, fx.LINE_22, fx.SCH_E_LINE_18])

    def test_all_pass_on_the_true_printed_values(self):
        workbook = self.write()
        self.assertEqual(
            [dw.cell_value(workbook, "Tie-outs", f"E{r}")
             for r in (2, 3, 4, 5)], [0, 0, 0, 0])
        self.assertEqual(self.verdicts(workbook), ["PASS"] * 4)

    def test_a_wrong_printed_value_fails_its_row_only(self):
        workbook = self.write(
            printed=true_printed(f4562_line_22=fx.LINE_22 + 1))
        self.assertEqual(dw.cell_value(workbook, "Tie-outs", "E4"), -1)
        self.assertEqual(
            self.verdicts(workbook), ["PASS", "PASS", "FAIL", "PASS"])

    def test_a_one_cent_difference_fails(self):
        workbook = self.write(printed=true_printed(
            activity_lines={RENTAL: fx.SCH_E_LINE_18 + 0.01}))
        self.assertEqual(self.verdicts(workbook)[3], "FAIL")

    def test_a_sub_half_cent_difference_passes(self):
        workbook = self.write(printed=true_printed(
            activity_lines={RENTAL: fx.SCH_E_LINE_18 + 0.004}))
        self.assertEqual(self.verdicts(workbook)[3], "PASS")

    def test_no_4562_rows_when_form_not_emitted(self):
        # Assets on a rental and on a Schedule C business, none placed this
        # year: Form 4562 is not filed, and each activity still ties out.
        equipment = dataclasses.replace(fx.APPLIANCE, description="Equipment")
        scenario = fx.scenario(
            rental_assets=(fx.BUILDING,), business_assets=(equipment,))
        workbook = self.write(scenario, dw.PrintedFigures(activity_lines={
            RENTAL: 7_272, ("schedule_c_businesses", 0): 1_920}))
        self.assertEqual(self.claims(workbook), [
            "Schedule E line 18: depreciation, rental property #0 "
            "('100 Example Street')",
            "Schedule C line 13 (business 1): depreciation, Schedule C "
            "business #0 ('Consulting')",
        ])
        self.assertEqual(self.formulas(workbook, "Tie-outs")["B2"],
                         "=SUMIF(Assets!A2:A3,1,Assets!K2:K3)")
        self.assertEqual(self.formulas(workbook, "Tie-outs")["B3"],
                         "=SUMIF(Assets!A2:A3,2,Assets!K2:K3)")
        self.assertEqual(
            [workbook["Assets"][f"A{r}"].value for r in (2, 3)], [1, 2])
        self.assertEqual(
            [dw.cell_value(workbook, "Tie-outs", f"B{r}") for r in (2, 3)],
            [7_272, 1_920])
        self.assertEqual(self.verdicts(workbook), ["PASS", "PASS"])

    def test_an_unprinted_activity_line_compares_as_zero(self):
        workbook = self.write(printed=dw.PrintedFigures())
        self.assertEqual(workbook["Tie-outs"]["C2"].value, 0)
        self.assertEqual(self.verdicts(workbook), ["FAIL"])


class OverrideTests(_WorkbookCase):
    """An overridden activity prints the override: its tie-out reads the
    override input, and a second row ties the asset sum to the engine figure
    the override restates."""

    def setUp(self):
        super().setUp()
        self.workbook = self.write(
            fx.scenario(
                rental_assets=(fx.BUILDING, fx.APPLIANCE),
                rental_override=fx.override(9_000.0, 9_192.0)),
            dw.PrintedFigures(activity_lines={RENTAL: 9_000}))

    def test_override_inputs_sit_below_the_assets_with_provenance(self):
        sheet = self.workbook["Assets"]
        self.assertEqual(
            [sheet[f"{c}5"].value for c in "ABIJN"],
            ["Activity #", "Activity", "Override amount",
             "Restated engine amount", "Provenance"])
        self.assertEqual(sheet["A6"].value, 1)
        self.assertEqual(sheet["I6"].value, 9_000)
        self.assertEqual(sheet["J6"].value, 9_192)
        self.assertEqual(sheet["N6"].value, (
            "Override amount: rental_properties[0].depreciation_override."
            "amount (acknowledged; the forms use this figure). Restated "
            "engine amount: rental_properties[0].depreciation_override."
            "restates_engine_amount"))

    def test_tie_out_rows_and_formulas(self):
        sheet = self.workbook["Tie-outs"]
        self.assertEqual([sheet[f"A{r}"].value for r in (2, 3)], [
            "Schedule E line 18: depreciation, rental property #0 "
            "('100 Example Street') (depreciation_override in effect)",
            "Engine figure restated by the override, rental property #0 "
            "('100 Example Street')"])
        formulas = self.formulas(self.workbook, "Tie-outs")
        self.assertEqual(formulas["B2"], "=ROUND(Assets!I6,0)")
        self.assertEqual(formulas["B3"], "=SUMIF(Assets!A2:A3,1,Assets!K2:K3)")
        self.assertEqual(formulas["C3"], "=ROUND(Assets!J6,0)")
        self.assertEqual(sheet["C2"].value, 9_000)
        self.assertEqual(sheet["D3"].value, (
            "depreciation_override.restates_engine_amount (Assets!J6); not "
            "printed on a form"))

    def test_both_rows_pass(self):
        self.assertEqual(
            [dw.cell_value(self.workbook, "Tie-outs", f"B{r}")
             for r in (2, 3)], [9_000, 9_192])
        self.assertEqual(
            [dw.cell_value(self.workbook, "Tie-outs", f"F{r}")
             for r in (2, 3)], ["PASS", "PASS"])


class HostileTextTests(_WorkbookCase):
    def test_description_starting_with_equals_is_text(self):
        hostile = dataclasses.replace(
            fx.APPLIANCE, description='=HYPERLINK("http://example.invalid")')
        workbook = self.write(
            fx.scenario(rental_assets=(hostile,)), dw.PrintedFigures())
        cell = workbook["Assets"]["C2"]
        self.assertEqual(cell.value, '=HYPERLINK("http://example.invalid")')
        self.assertEqual(cell.data_type, "s")
        self.assertEqual(
            dw.cell_value(workbook, "Year-by-year", "A2"),
            '=HYPERLINK("http://example.invalid")')

    def test_duplicate_descriptions_get_distinct_rows(self):
        twin = dataclasses.replace(fx.APPLIANCE, basis=4_000.0,
                                   prior_depreciation=2_080.0)
        workbook = self.write(
            fx.scenario(rental_assets=(fx.APPLIANCE, twin)),
            dw.PrintedFigures(activity_lines={RENTAL: 2_688}))
        self.assertEqual(
            [workbook["Assets"][f"C{r}"].value for r in (2, 3)],
            ["Appliance set", "Appliance set"])
        # 19.2% of 10,000 and of 4,000.
        self.assertEqual(
            [dw.cell_value(workbook, "Assets", f"K{r}") for r in (2, 3)],
            [1_920, 768])
        self.assertEqual(dw.cell_value(workbook, "Tie-outs", "F2"), "PASS")

    def test_long_description_is_kept_whole(self):
        long = dataclasses.replace(fx.APPLIANCE, description="x" * 300)
        workbook = self.write(
            fx.scenario(rental_assets=(long,)), dw.PrintedFigures())
        self.assertEqual(workbook["Assets"]["C2"].value, "x" * 300)


class EvaluatorTests(unittest.TestCase):
    """The Python mirror's own arithmetic, on hand cases."""

    def setUp(self):
        self.workbook = openpyxl.Workbook()
        self.sheet = self.workbook.active
        self.sheet.title = "One"
        other = self.workbook.create_sheet("Two words")
        for coordinate, value in {
                "A1": 10, "A2": 2.5, "A3": "k", "A4": "j", "A5": "k",
                "B3": 1, "B4": 20, "B5": 300, "C1": "=A1*A2"}.items():
            self.sheet[coordinate] = value
        other["A1"] = 7

    def ev(self, formula):
        return dw.evaluate(self.workbook, "One", formula)

    def test_arithmetic_and_precedence(self):
        self.assertEqual(self.ev("=A1+A2*2"), 15)
        self.assertEqual(self.ev("=(A1+A2)*2"), 25)
        self.assertEqual(self.ev("=A1-A2-A2"), 5)
        self.assertEqual(self.ev("=-A1+3"), -7)
        self.assertEqual(self.ev("=A1/4"), 2.5)

    def test_round_is_half_up_not_half_even(self):
        self.assertEqual(self.ev("=ROUND(0.5,0)"), 1)
        self.assertEqual(self.ev("=ROUND(2.5,0)"), 3)
        self.assertEqual(self.ev("=ROUND(2.4999,0)"), 2)
        self.assertEqual(self.ev("=ROUND(A2*A1,0)"), 25)
        self.assertEqual(self.ev("=ROUND(1.005,1)"), 1.0)

    def test_min_max_abs(self):
        self.assertEqual(self.ev("=MIN(A1,A2)"), 2.5)
        self.assertEqual(self.ev("=MAX(0,A2-A1)"), 0)
        self.assertEqual(self.ev("=ABS(A2-A1)"), 7.5)

    def test_sum_and_sumif_over_ranges(self):
        self.assertEqual(self.ev("=SUM(B3:B5)"), 321)
        self.assertEqual(self.ev('=SUMIF(A3:A5,"k",B3:B5)'), 301)
        self.assertEqual(self.ev('=SUMIF(A3:A5,"z",B3:B5)'), 0)
        self.assertEqual(self.ev("=SUMIF(B3:B5,20,B3:B5)"), 20)

    def test_if_and_comparison(self):
        self.assertEqual(self.ev('=IF(ABS(A2)<0.005,"PASS","FAIL")'), "FAIL")
        self.assertEqual(self.ev('=IF(ABS(0)<0.005,"PASS","FAIL")'), "PASS")
        self.assertEqual(self.ev('=IF(0.005<0.005,"PASS","FAIL")'), "FAIL")

    def test_references_follow_formulas_sheets_and_anchors(self):
        self.assertEqual(self.ev("=C1+1"), 26)
        self.assertEqual(self.ev("='Two words'!A1*2"), 14)
        self.assertEqual(self.ev("=One!$A$1"), 10)

    def test_an_empty_cell_is_zero(self):
        self.assertEqual(self.ev("=A1+Z99"), 10)
        self.assertEqual(self.ev("=Z99"), 0)

    def test_text_in_arithmetic_is_refused(self):
        with self.assertRaises(dw.FormulaError):
            self.ev("=A3+1")

    def test_an_unknown_function_is_refused(self):
        with self.assertRaises(dw.FormulaError):
            self.ev("=VLOOKUP(A1,B3:B5,1)")

    def test_garbage_is_refused(self):
        for text in ("=A1+", "=A1 A2", "=#REF!", "A1+A2"):
            with self.subTest(text=text):
                with self.assertRaises(dw.FormulaError):
                    self.ev(text)



class FullRecoveryLengthTests(_WorkbookCase):
    """A building in the last year of its 27.5-year schedule: 28 year rows,
    every one evaluated from the cells as written."""

    def setUp(self):
        super().setUp()
        self.scenario = fx.scenario(rental_assets=(fx.OLD_BUILDING,))
        self.workbook = self.write(self.scenario, dw.PrintedFigures(
            activity_lines={RENTAL: fx.OLD_BUILDING_2025}))

    def test_twenty_eight_year_rows(self):
        sheet = self.workbook["Year-by-year"]
        self.assertEqual(sheet.max_row, 29)
        self.assertEqual(
            [sheet[f"C{r}"].value for r in (2, 29)], [1998, 2025])
        self.assertEqual(
            [sheet[f"D{r}"].value for r in (2, 3, 11, 12, 28, 29)],
            [0.02879, 0.03636, 0.03637, 0.03636, 0.03636, 0.02576])
        self.assertEqual(
            self.formulas(self.workbook, "Year-by-year")["E29"],
            "=ROUND(MAX(0,MIN(Assets!$I$2*D29,Assets!$I$2-F28)),0)")

    def test_every_year_row_gives_the_hand_figure(self):
        cache = {}
        value = lambda c: dw.cell_value(
            self.workbook, "Year-by-year", c, cache)
        expected = ([2_879] + [3_636] * 8 + [3_637, 3_636] * 9 + [2_576])
        self.assertEqual(len(expected), 28)
        self.assertEqual([value(f"E{r}") for r in range(2, 30)], expected)
        self.assertEqual(value("F28"), 97_424)
        self.assertEqual(value("F29"), 100_000)

    def test_the_last_cumulative_cell_is_readable_on_its_own(self):
        # No shared cache: one call walks the whole chain above it.
        self.assertEqual(
            dw.cell_value(self.workbook, "Year-by-year", "F29"), 100_000)

    def test_current_year_and_tie_out(self):
        self.assertEqual(
            dw.cell_value(self.workbook, "Assets", "K2"), 2_576)
        self.assertEqual(
            dw.cell_value(self.workbook, "Assets", "L2"), 100_000)
        self.assertEqual(
            dw.cell_value(self.workbook, "Tie-outs", "F2"), "PASS")


class WrittenAsStatedTests(_WorkbookCase):
    def test_basis_is_written_unrounded(self):
        # 20% of 1,234.56 = 246.912 -> 247.
        cents = dataclasses.replace(
            fx.FURNITURE, basis=1_234.56, recovery_class="5-year")
        workbook = self.write(
            fx.scenario(rental_assets=(cents,)), dw.PrintedFigures())
        self.assertEqual(workbook["Assets"]["I2"].value, 1_234.56)
        self.assertEqual(dw.cell_value(workbook, "Assets", "K2"), 247)

    def test_stated_prior_is_written_unrounded(self):
        cents = dataclasses.replace(fx.APPLIANCE, prior_depreciation=5_200.25)
        workbook = self.write(
            fx.scenario(rental_assets=(cents,)), dw.PrintedFigures())
        self.assertEqual(workbook["Assets"]["J2"].value, 5_200.25)

    def test_activity_label_and_claim_starting_with_equals_are_text(self):
        base = fx.scenario()
        rental = dataclasses.replace(
            base.rental_properties[0], address="=1+1")
        scenario = dataclasses.replace(base, rental_properties=[rental])
        workbook = self.write(scenario)
        label = workbook["Assets"]["B2"]
        self.assertEqual(label.data_type, "s")
        self.assertEqual(label.value, "rental property #0 ('=1+1')")
        # The orchestrator's labels always begin with the engine's own words;
        # the writer takes whatever trail it is handed, so hand it one whose
        # label IS the hostile text.
        [activity] = resolver.audit_trail(fx.scenario())
        bare = dataclasses.replace(activity, label="=1+1", assets=tuple(
            dataclasses.replace(a, activity_label="=1+1")
            for a in activity.assets))
        path = dw.write_depreciation_audit(
            (bare,), true_printed(), self.tmp / "bare.xlsx")
        written = openpyxl.load_workbook(path)["Assets"]["B2"]
        self.assertEqual(written.data_type, "s")
        self.assertEqual(written.value, "=1+1")
        # A label that IS the hostile text, not merely containing it.
        with unittest.mock.patch.object(
                dw, "_SECTION_LINE",
                {"rental_properties": "=1+1",
                 "schedule_c_businesses": "=1+1"}):
            hostile = self.write(scenario, name="hostile.xlsx")
        claim = hostile["Tie-outs"]["A5"]
        self.assertEqual(claim.data_type, "s")
        self.assertTrue(claim.value.startswith("=1+1: depreciation"))
        source = hostile["Tie-outs"]["D5"]
        self.assertEqual(source.data_type, "s")
        self.assertEqual(source.value, "=1+1 as printed")


if __name__ == "__main__":
    unittest.main()
