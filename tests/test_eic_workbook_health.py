"""Workbook-oracle health for the childless EIC (spec §10.3) and the
oracle-only answer adapter's mechanics (tests/eic_oracle_adapter.py) — all
static: openpyxl reads and the engine's own input writer, never a
recalculation, so nothing here reaches LibreOffice. All values synthetic or
read from params; nothing typed in."""
import csv
import functools
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

from tenforty.mappings.f1040 import (
    EIC_FORM_4797_NO_KEY, EIC_FORM_4797_QUESTION, F1040,
    eic_form_4797_no_answer)
from tenforty.oracle.engine import SpreadsheetEngine
from tenforty.orchestrator import ReturnOrchestrator
from tenforty.params.federal import load as load_params
from tests import eic_oracle_adapter as adapter
from tests.fixtures import eic_battery
from tests.helpers import REPO_ROOT, SPREADSHEETS_DIR

YEARS = tuple(sorted(adapter.EIC_WORKBOOK_CELLS))


@functools.cache
def _workbook(year):
    """The year's template, loaded once (read-only use only)."""
    return openpyxl.load_workbook(
        SPREADSHEETS_DIR / "federal" / str(year) / "1040.xlsx")


class WorkbookConstantsMatchParamsTests(unittest.TestCase):
    """§10.3 item 1: the EIC sheet's hardcoded constants equal the attested
    params. A mismatch means that year's workbook is not an oracle for the
    affected rows until adjudicated — never allowlisted."""

    def test_every_constant_cell(self):
        for year in YEARS:
            params = load_params(year)
            wb = _workbook(year)
            for (sheet, cell), field in adapter.EIC_WORKBOOK_CELLS[year][
                    "constants"].items():
                name, column = field if isinstance(field, tuple) else (
                    field, None)
                expected = getattr(params, name)
                if column is not None:
                    expected = expected[column]
                with self.subTest(year=year, cell=f"{sheet}!{cell}"):
                    self.assertEqual(wb[sheet][cell].value, expected)


# Ruling R-S7-10: line 27 as each template computes it. 2021-2024 carry no
# "no credit" gate at all; 2025 zeroes only through AI117 <- EIC_No_Box (O1),
# whose OR-chain has no Step 2 term. NoEICredit (A1) feeds neither, so the
# Step 2 stop is display-only in every year — the link that was never pinned.
LINE_27 = {
    2025: ("'1040'!$AF$113",
           '=IF(AX113<>"",ROUND(AX113,0),IF(AI117<>"",0,EarnedIncomeCredit))'),
    2024: ("'1040'!$Z$90",
           '=IF(AN90<>"",ROUND(AN90,0),EarnedIncomeCredit)'),
    2023: ("'1040'!$Z$87",
           '=IF(AN87<>"",ROUND(AN87,0),EarnedIncomeCredit)'),
    2022: ("'1040'!$Y$81",
           '=IF(AM81<>"",ROUND(AM81,0),EarnedIncomeCredit)'),
    2021: ("'1040'!$Y$72",
           '=IF(AM72<>"",ROUND(AM72,0),EarnedIncomeCredit)'),
}
LINE_27_GATE_2025 = {
    ("1040", "AI117"): '=IF(EIC_No_Box,"X","")',
    ("Line 27a (EIC)", "O1"): "=OR(G28,A157,G199,A216,A321,A430,O462)",
}


def _named(wb, name):
    sheet, cell = wb.defined_names[name].attr_text.replace("$", "").split("!")
    return sheet.strip("'"), cell


class LineTwentySevenTests(unittest.TestCase):
    """Ruling R-S7-10, per workbook year: what line 27 reads, and that the
    sheet's own verdict NoEICredit sees the Form 4797 answer. A relayout or a
    vendor fix (line 27 starting to honor NoEICredit) fails here — then
    revisit refuse_inconsistent_workbook_eic."""

    def test_the_eic_name_and_its_formula(self):
        for year, (target, formula) in sorted(LINE_27.items()):
            wb = _workbook(year)
            with self.subTest(year=year):
                self.assertEqual(wb.defined_names["EIC"].attr_text, target)
                sheet, cell = _named(wb, "EIC")
                self.assertEqual(wb[sheet][cell].value, formula)

    def test_the_2025_gate_has_no_step_2_term(self):
        wb = _workbook(2025)
        self.assertEqual(_named(wb, "EIC_No_Box"), ("Line 27a (EIC)", "O1"))
        for (sheet, cell), formula in LINE_27_GATE_2025.items():
            with self.subTest(cell=f"{sheet}!{cell}"):
                self.assertEqual(wb[sheet][cell].value, formula)

    def test_no_eic_credit_sees_the_form_4797_answer(self):
        self.assertEqual(set(LINE_27), set(EIC_FORM_4797_QUESTION))
        for year, q in sorted(EIC_FORM_4797_QUESTION.items()):
            wb = _workbook(year)
            with self.subTest(year=year):
                self.assertEqual(_named(wb, "NoEICredit"), (q.sheet, "A1"))
                self.assertIn(f'{q.q3_no}<>""', wb[q.sheet]["A1"].value)


class LiteralTableTests(unittest.TestCase):
    """Years whose workbook EIC Table holds LITERAL values (2024; 2025's is
    rebuilt from derived slopes, R-S7-8): the sheet's whole table equals the
    ingested published table (assets/eic_tables/<year>.csv) — every bin,
    every column. A vendor edit or a re-ingest drift fails here."""

    def test_the_sheet_table_is_the_published_table(self):
        # Each published bin is read from the sheet row that contains it (a
        # bin the publication splits for one column's footnote may be one
        # row on the sheet). Frozen exceptions per year (R-S8-10): formula
        # cells, pinned as text, and typed cells departing from the print.
        years = [y for y in YEARS
                 if adapter.EIC_WORKBOOK_CELLS[y].get("literal_table")]
        self.assertTrue(years)
        for year in years:
            cells = adapter.EIC_WORKBOOK_CELLS[year]
            formulas = cells.get("table_formulas", {})
            exceptions = cells.get("table_literal_exceptions", {})
            ws = _workbook(year)["EIC Table"]
            rows, found = [], {}
            for row in ws.iter_rows(min_row=3, max_col=10):
                lo, hi = row[0].value, row[1].value
                if isinstance(lo, (int, float)) and isinstance(hi, (int, float)):
                    rows.append((int(lo), int(hi), row[2:]))
                for cell in row[2:]:
                    if str(cell.value).startswith("="):
                        found[cell.coordinate] = cell.value
            with self.subTest(year=year, check="formula cells"):
                self.assertEqual(found, {k: "".join(v) if isinstance(v, tuple)
                                         else v for k, v in formulas.items()})
            with open(REPO_ROOT / "assets" / "eic_tables" / f"{year}.csv",
                      newline="") as fh:
                reader = list(csv.reader(fh))
            columns = reader[0][2:]
            mismatched = {}
            for record in reader[1:]:
                lo, hi = int(record[0]), int(record[1])
                s_lo, s_hi, sheet_cells = max(
                    (r for r in rows if r[0] <= lo), key=lambda r: r[0])
                self.assertLessEqual(hi, s_hi, (year, lo, hi))
                for column, printed, cell in zip(columns, record[2:],
                                                 sheet_cells):
                    if cell.coordinate in found:
                        continue
                    if cell.value != int(printed):
                        mismatched[(s_lo, column)] = (cell.value, int(printed))
            with self.subTest(year=year, check="literal cells"):
                self.assertEqual(mismatched, exceptions)


class AnswerCellStructureTests(unittest.TestCase):
    """Each eligibility question has exactly one INPUT cell (blank in the
    template) and a formula for its other answer; Step 2 Q3 (Form 4797) is
    the one question with two input cells (ruling R-S7-6a). A vendor
    relayout fails here.

    Every input — including the Q3 "No" cell production now writes — is
    blank in the TEMPLATE: production's write lands in the compute-time copy
    the engine makes, never in the template. That is not a contradiction;
    do not "fix" it."""

    def test_inputs_are_blank_and_the_derived_side_is_a_formula(self):
        for year in YEARS:
            cells = adapter.EIC_WORKBOOK_CELLS[year]
            wb = _workbook(year)
            for key, (sheet, cell) in cells["inputs"].items():
                with self.subTest(year=year, input=key):
                    self.assertIsNone(wb[sheet][cell].value)
            for cell, formula in cells["derived"].items():
                with self.subTest(year=year, derived=cell):
                    self.assertEqual(wb[cells["sheet"]][cell].value, formula)
            for key, (sheet, cell) in cells["diagnostics"].items():
                with self.subTest(year=year, diagnostic=key):
                    self.assertTrue(str(wb[sheet][cell].value)
                                    .startswith("="))
            for key, (cell, formula) in cells["stops"].items():
                with self.subTest(year=year, stop=key):
                    self.assertEqual(wb[cells["sheet"]][cell].value, formula)
            for cell, formula in cells.get("observations", {}).items():
                with self.subTest(year=year, observation=cell):
                    self.assertEqual(wb[cells["sheet"]][cell].value, formula)
            for cell, value in cells.get("table_observations", {}).items():
                with self.subTest(year=year, table_observation=cell):
                    self.assertEqual(wb["EIC Table"][cell].value, value)

    def test_the_form_4797_no_cell_is_productions(self):
        # The adapter answers Q3 in the very cell production writes.
        for year in YEARS:
            sheet, cell = adapter.EIC_WORKBOOK_CELLS[year]["inputs"][
                EIC_FORM_4797_NO_KEY]
            with self.subTest(year=year):
                self.assertEqual(
                    (sheet, cell),
                    (F1040.SHEET_MAP[year][EIC_FORM_4797_NO_KEY],
                     F1040.get_inputs(year)[EIC_FORM_4797_NO_KEY]))


class AdapterWritesEveryAnswerTests(unittest.TestCase):
    """The answers land in the input cells through the engine's own writer
    (no recalculation), and the derived formulas survive."""

    def _written(self, year, scenario):
        mapping = adapter.answering_mapping(year)
        with tempfile.TemporaryDirectory() as tmp:
            copy = Path(tmp) / "1040.xlsx"
            shutil.copy2(SPREADSHEETS_DIR / "federal" / str(year)
                         / "1040.xlsx", copy)
            SpreadsheetEngine()._write_inputs(
                copy, mapping.get_inputs(year), mapping.SHEET_MAP[year],
                adapter.answers(scenario))
            return openpyxl.load_workbook(copy)

    def test_favorable_facts_clear_every_input_and_state_no_children(self):
        year = 2025
        cells = adapter.EIC_WORKBOOK_CELLS[year]
        wb = self._written(year, eic_battery.battery_for(
            year, "single")["plateau"](year))
        special = {"eic_answer_qualifying_children": 0,
                   EIC_FORM_4797_NO_KEY: eic_form_4797_no_answer(year)}
        for key, (sheet, cell) in cells["inputs"].items():
            with self.subTest(input=key):
                self.assertEqual(wb[sheet][cell].value, special.get(key))
        for cell, formula in cells["derived"].items():
            self.assertEqual(wb[cells["sheet"]][cell].value, formula)

    def test_each_unfavorable_fact_marks_only_its_own_answer(self):
        # Pure: which answer each stated fact sets (the cell it lands in is
        # proven by the writes in this class).
        for kind, key in (
                ("fact_ssn_not_valid", "eic_answer_ssn_not_valid"),
                ("fact_nonresident_alien", "eic_answer_nonresident_alien"),
                ("fact_main_home_not_in_us", "eic_answer_main_home_not_us"),
                ("fact_qualifying_child_of_another",
                 "eic_answer_qc_of_another"),
                ("fact_claimable_as_dependent", "eic_answer_dependent_you")):
            with self.subTest(kind=kind):
                got = adapter.answers(eic_battery.battery_for(
                    2025, "single")[kind](2025))
                self.assertEqual({k for k, v in got.items() if v == "X"},
                                 {key})

    def test_an_unfavorable_fact_lands_in_its_cell(self):
        year = 2025
        sheet, cell = adapter.EIC_WORKBOOK_CELLS[year]["inputs"][
            "eic_answer_ssn_not_valid"]
        wb = self._written(year, eic_battery.battery_for(
            year, "single")["fact_ssn_not_valid"](year))
        self.assertEqual(wb[sheet][cell].value, "X")


class AdapterRoutesThroughTheEngineTests(unittest.TestCase):
    """Inside the adapter, the workbook path hands the engine the answering
    mapping and the answers (the engine is stopped before recalculation)."""

    def test_the_engine_receives_the_answers(self):
        class Stop(Exception):
            pass
        seen = {}

        def record(**kwargs):
            seen.update(kwargs)
            raise Stop
        scenario = eic_battery.battery_for(2025, "single")["plateau"](2025)
        with tempfile.TemporaryDirectory() as tmp:
            orch = ReturnOrchestrator(spreadsheets_dir=SPREADSHEETS_DIR,
                                      work_dir=Path(tmp))
            with mock.patch.object(orch.engine, "compute", record):
                with adapter.eic_answers_adapter(orch):
                    eff, _ = orch._build_effective_scenario(scenario)
                    with self.assertRaises(Stop):
                        orch._compute_1040_via_workbook(eff)
        self.assertIn("eic_answer_qualifying_children", seen["inputs"])
        self.assertIn("eic_answer_ssn_not_valid",
                      seen["mapping"].get_inputs(2025))
        self.assertIn("eic_diag_prompt_ao",
                      seen["mapping"].get_outputs(2025))
