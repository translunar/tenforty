"""Form 4562 copies on a return whose Schedule E fills columns B and C.

Schedule E prints up to three rentals in columns A / B / C, and every
activity that places property in service this year gets its own Form 4562
copy, named by recon key (`rental_0`, `rental_1`, `rental_2`, `sch_c_0`).
These tests hold the two together: each copy names its own property, each
copy's line 22 is its own column's line 18, a rental that places nothing
gets no copy while its column still prints, and the 40% answer is the
RETURN's on every column and every copy.

Every PDF test goes through the orchestrator's emit path and reads the
filled PDFs. Nothing calls `f4562.compute_for` or `sch_e.compute` with a
hand-picked 40% answer.

THE SHAPES. Rentals are "100 / 200 / 300 Example Street" in that order, so
they are columns A / B / C and recon keys rental_0 / rental_1 / rental_2.
The business is "Consulting" (sch_c_0). Figures are Publication 946 rates
times basis, typed from the publication:

  Table A-6 (27.5-year, mid-month)  placed in month 1, year 1:  3.485%
                                    placed in month 6, year 7:  3.636%
  Table A-1 (5-year, half-year)     year 1: 20.00%
  Tables A-2..A-5 (5-year, mid-quarter), year 1, by quarter placed:
                                    35.00% / 25.00% / 15.00% / 5.00%

  buildings        A  building 200,000 placed June 2019        7,272, no copy
                   B  building 100,000 placed January 2025     3,485
                   C  building  60,000 placed January 2025     2,091
    twin: A also places equipment 10,000 in February 2025 (alone and
    pooled: nothing in the last quarter -> half-year, 2,000). A's copy:
    line 17 7,272, row 19b 2,000, line 22 9,272.

  two rentals      A  equipment 9,000 placed March
                   B  equipment 1,000 placed November (alone: 100% late)
    pooled 1,000 of 10,000 -> the RETURN is half-year: A 1,800, B 200
    (mid-quarter would print B 50)

  tipped by a rental
                   A  equipment 6,000 March + 3,000 November (alone 33 1/3%)
                   B  equipment 3,000 placed December
    pooled 6,000 of 12,000 -> mid-quarter: A 2,100 + 150 = 2,250, B 150
    (half-year would print A 1,800)

  three rentals and a business
                   A  equipment 6,000 March + 3,000 November (alone 33 1/3%)
                   B  equipment 5,000 March + 2,000 November (alone 28 4/7%)
                   C  equipment 4,000 June  + 1,000 October  (alone 20%)
                   business  equipment 8,000 placed December
    pooled 14,000 of 29,000 -> mid-quarter on all four copies:
      A 2,100 + 150 = 2,250   B 1,750 + 100 = 1,850
      C 1,000 +  50 = 1,050   business 400
    (half-year would print 1,800 / 1,400 / 1,000 / 1,600)

  workbook         A  building 200,000 placed January 2019     7,272, no copy
                   B  equipment 9,000 placed March             1,800
                   C  equipment 1,000 placed November            200
    whole return: 1,800 + 200 + 7,272 = 9,272

Field paths are 2025 LITERALS: Form 4562's as anchored in
tests/test_f4562_per_activity_emit.py, Schedule E's column cells as typed
from the template in tests/test_sch_e_three_properties_emit.py and
tests/test_pdf_sch_e_mapping.py -- never looked up through a mapping.
"""

import dataclasses
import unittest
from datetime import date

import openpyxl

from tenforty.audit import depreciation_workbook as dw
from tenforty.forms.depreciation.resolver import mid_quarter_applies
from tenforty.models import DepreciableAsset, RentalProperty
from tests import _audit_workbook_fixtures as fx
from tests import test_depreciation_audit_multi_activity as multi
from tests import test_f4562_emit_gate as gate

YEAR = 2025
ADDRESSES = ("100 Example Street", "200 Example Street", "300 Example Street")
COPY_KEYS = ("f4562_rental_0", "f4562_rental_1", "f4562_rental_2")

_P1 = "topmostSubform[0].Page1[0]"
ACTIVITY_BOX = f"{_P1}.f1_2[0]"
LINE_17 = f"{_P1}.f1_25[0]"
ROW_19B_BASIS = f"{_P1}.SectionBTable[0].Line19b[0].f1_33[0]"
ROW_19B_CONVENTION = f"{_P1}.SectionBTable[0].Line19b[0].f1_35[0]"
ROW_19B_DEDUCTION = f"{_P1}.SectionBTable[0].Line19b[0].f1_37[0]"
LINE_22 = "topmostSubform[0].Page2[0].f2_2[0]"
# Schedule E page 1, columns A / B / C.
SCH_E_LINE_1A = (
    f"{_P1}.Table_Line1a[0].RowA[0].f1_3[0]",
    f"{_P1}.Table_Line1a[0].RowB[0].f1_4[0]",
    f"{_P1}.Table_Line1a[0].RowC[0].f1_5[0]")
SCH_E_LINE_18 = (
    f"{_P1}.Table_Expenses[0].Line18[0].f1_61[0]",
    f"{_P1}.Table_Expenses[0].Line18[0].f1_62[0]",
    f"{_P1}.Table_Expenses[0].Line18[0].f1_63[0]")
SCH_C_LINE_13 = f"{_P1}.Lines8-17[0].f1_22[0]"

_read = gate._read


def _rental(column: int, *assets) -> RentalProperty:
    """The rental printed in column ``column`` (0 = A)."""
    return RentalProperty(
        address=ADDRESSES[column], property_type=1, fair_rental_days=365,
        personal_use_days=0, rents_received=24_000.0,
        depreciable_assets=list(assets))


def _building(basis: float) -> DepreciableAsset:
    """A residential rental building placed in January of the return year."""
    return DepreciableAsset(
        description="Rental building", date_placed_in_service=date(YEAR, 1, 15),
        basis=basis, recovery_class="27.5-year")


def _scenario(*assets_by_column, businesses=()):
    return gate._scenario(
        rentals=[_rental(column, *assets)
                 for column, assets in enumerate(assets_by_column)],
        businesses=businesses)


def _buildings(*, column_a_places: bool = False):
    column_a = (gate._old_building(),)
    if column_a_places:
        column_a += (gate._new_equipment(),)
    return _scenario(
        column_a, (_building(100_000.0),), (_building(60_000.0),))


def _two_rentals():
    return _scenario(
        (fx.equipment(9_000.0, 3),), (fx.equipment(1_000.0, 11),))


def _tipped_by_a_rental():
    return _scenario(
        (fx.equipment(6_000.0, 3), fx.equipment(3_000.0, 11)),
        (fx.equipment(3_000.0, 12),))


_THREE_RENTALS = (
    (fx.equipment(6_000.0, 3), fx.equipment(3_000.0, 11)),
    (fx.equipment(5_000.0, 3), fx.equipment(2_000.0, 11)),
    (fx.equipment(4_000.0, 6), fx.equipment(1_000.0, 10)))


def _three_rentals_and_a_business():
    # 20,000 of receipts keeps the return well under the Form 8995
    # simple-path threshold whichever convention prints.
    return _scenario(
        *_THREE_RENTALS, businesses=[fx.business(
            fx.equipment(8_000.0, 12), gross_receipts=20_000.0)])


def _workbook_rentals():
    return _scenario(
        (fx.BUILDING,), (fx.equipment(9_000.0, 3),),
        (fx.equipment(1_000.0, 11),))


def _copy_keys(emitted) -> list[str]:
    return sorted(key for key in emitted if key.startswith("f4562"))


def _row_19b(pdf_path) -> tuple[str, str, str]:
    """(basis, convention, deduction) as printed on row 19b."""
    return (_read(pdf_path, ROW_19B_BASIS),
            _read(pdf_path, ROW_19B_CONVENTION),
            _read(pdf_path, ROW_19B_DEDUCTION))


def _line_18s(emitted) -> list[str]:
    """Schedule E line 18 as printed in columns A, B, C."""
    return [_read(emitted["sch_e"], path) for path in SCH_E_LINE_18]


class ColumnsBAndCCopiesTests(gate._EmitCase):
    """Rental A holds a prior-year building only; rentals B and C each place
    a building this year."""

    def test_columns_b_and_c_each_get_a_copy_named_by_recon_key(self):
        emitted = self._emit(_buildings())
        self.assertEqual(
            _copy_keys(emitted), ["f4562_rental_1", "f4562_rental_2"])
        self.assertEqual(
            emitted["f4562_rental_1"].name, "f4562_2025_rental_1.pdf")
        self.assertEqual(
            emitted["f4562_rental_2"].name, "f4562_2025_rental_2.pdf")
        self.assertEqual(
            sorted(p.name for p in emitted["1040"].parent.glob("f4562*")),
            ["f4562_2025_rental_1.pdf", "f4562_2025_rental_2.pdf"])

    def test_a_column_with_only_prior_year_assets_gets_no_copy_but_prints_line_18(
            self):
        emitted = self._emit(_buildings())
        self.assertNotIn("f4562_rental_0", emitted)
        self.assertNotIn("f4562", emitted)
        self.assertEqual(
            list(emitted["1040"].parent.glob("f4562*rental_0*")), [])
        self.assertEqual(_read(emitted["sch_e"], SCH_E_LINE_18[0]), "7272")
        # Its building is on neither copy: each copy's line 17 is blank and
        # its line 22 is its own building alone.
        for key, line_22 in (("f4562_rental_1", "3485"),
                             ("f4562_rental_2", "2091")):
            with self.subTest(copy=key):
                self.assertEqual(_read(emitted[key], LINE_17), "")
                self.assertEqual(_read(emitted[key], LINE_22), line_22)

    def test_a_column_that_places_property_gets_its_copy(self):
        """Twin of the test above: the same rental A, now also placing
        equipment this year, has a copy -- three in all."""
        emitted = self._emit(_buildings(column_a_places=True))
        self.assertEqual(_copy_keys(emitted), list(COPY_KEYS))
        self.assertEqual(
            emitted["f4562_rental_0"].name, "f4562_2025_rental_0.pdf")
        copy = emitted["f4562_rental_0"]
        self.assertEqual(_read(copy, ACTIVITY_BOX), ADDRESSES[0])
        self.assertEqual(_read(copy, LINE_17), "7272")
        self.assertEqual(_row_19b(copy), ("10000", "HY", "2000"))
        self.assertEqual(_read(copy, LINE_22), "9272")
        self.assertEqual(_read(emitted["sch_e"], SCH_E_LINE_18[0]), "9272")

    def test_each_copys_line_22_ties_to_its_columns_line_18(self):
        emitted = self._emit(_buildings())
        column_b, column_c = _line_18s(emitted)[1:]
        self.assertEqual([column_b, column_c], ["3485", "2091"])
        self.assertNotEqual(column_b, column_c)
        self.assertEqual(_read(emitted["f4562_rental_1"], LINE_22), column_b)
        self.assertEqual(_read(emitted["f4562_rental_2"], LINE_22), column_c)

    def test_each_copys_header_names_its_own_property(self):
        emitted = self._emit(_buildings())
        for column, key in ((1, "f4562_rental_1"), (2, "f4562_rental_2")):
            with self.subTest(copy=key):
                header = _read(emitted[key], ACTIVITY_BOX)
                self.assertEqual(header, ADDRESSES[column])
                self.assertEqual(
                    header, _read(emitted["sch_e"], SCH_E_LINE_1A[column]))

    def test_the_copies_share_the_returns_convention(self):
        """Rental B alone is over the 40% line; the return is under it."""
        scenario = _two_rentals()
        self.assertIs(mid_quarter_applies(scenario), False)
        self.assertIs(mid_quarter_applies(dataclasses.replace(
            scenario, rental_properties=scenario.rental_properties[1:])),
            True)
        emitted = self._emit(scenario)
        self.assertEqual(
            _row_19b(emitted["f4562_rental_1"]), ("1000", "HY", "200"))
        self.assertEqual(
            _row_19b(emitted["f4562_rental_0"]), ("9000", "HY", "1800"))
        self.assertEqual(_line_18s(emitted)[:2], ["1800", "200"])


class TwoRentalCopiesTests(gate._EmitCase):
    """Two rentals, both placing property this year."""

    def test_two_copies_keyed_rental_0_and_rental_1(self):
        emitted = self._emit(_two_rentals())
        self.assertEqual(
            _copy_keys(emitted), ["f4562_rental_0", "f4562_rental_1"])
        self.assertNotIn("f4562", emitted)
        self.assertEqual(
            sorted(p.name for p in emitted["1040"].parent.glob("f4562*")),
            ["f4562_2025_rental_0.pdf", "f4562_2025_rental_1.pdf"])

    def test_each_copy_names_its_own_address_and_its_own_column(self):
        emitted = self._emit(_two_rentals())
        headers = [_read(emitted[key], ACTIVITY_BOX) for key in COPY_KEYS[:2]]
        self.assertEqual(headers, list(ADDRESSES[:2]))
        self.assertEqual(headers, [
            _read(emitted["sch_e"], path) for path in SCH_E_LINE_1A[:2]])

    def test_the_columns_are_not_cross_wired(self):
        emitted = self._emit(_two_rentals())
        line_22s = [_read(emitted[key], LINE_22) for key in COPY_KEYS[:2]]
        self.assertEqual(line_22s, ["1800", "200"])
        # Column A is rental_0's and column B is rental_1's; C is blank.
        self.assertEqual(_line_18s(emitted), line_22s + [""])

    def test_a_rental_under_alone_prints_mid_quarter_when_another_rental_tips_the_return(
            self):
        scenario = _tipped_by_a_rental()
        self.assertIs(mid_quarter_applies(scenario), True)
        self.assertIs(mid_quarter_applies(dataclasses.replace(
            scenario, rental_properties=scenario.rental_properties[:1])),
            False)
        emitted = self._emit(scenario)
        self.assertEqual(
            _row_19b(emitted["f4562_rental_0"]), ("9000", "MQ", "2250"))
        self.assertEqual(
            _row_19b(emitted["f4562_rental_1"]), ("3000", "MQ", "150"))
        self.assertEqual(_line_18s(emitted)[:2], ["2250", "150"])


class ThreeRentalsAndABusinessTests(gate._EmitCase):
    """Columns A, B and C beside a Schedule C business, every one of the
    four placing property this year. No rental is over the 40% line alone;
    the return is."""

    COPIES = COPY_KEYS + ("f4562_sch_c_0",)

    def test_four_copies_with_stable_recon_keys(self):
        emitted = self._emit(_three_rentals_and_a_business())
        self.assertEqual(_copy_keys(emitted), list(self.COPIES))
        self.assertEqual(
            sorted(p.name for p in emitted["1040"].parent.glob("f4562*")),
            ["f4562_2025_rental_0.pdf", "f4562_2025_rental_1.pdf",
             "f4562_2025_rental_2.pdf", "f4562_2025_sch_c_0.pdf"])
        self.assertEqual(
            [_read(emitted[key], ACTIVITY_BOX) for key in self.COPIES],
            list(ADDRESSES) + ["Consulting"])

    def test_each_column_ties_to_its_own_copy(self):
        emitted = self._emit(_three_rentals_and_a_business())
        line_22s = [_read(emitted[key], LINE_22) for key in self.COPIES]
        self.assertEqual(line_22s, ["2250", "1850", "1050", "400"])
        self.assertEqual(_line_18s(emitted), line_22s[:3])
        self.assertEqual(
            _read(emitted["sch_c_1"], SCH_C_LINE_13), line_22s[3])
        self.assertEqual(
            [_read(emitted["sch_e"], path) for path in SCH_E_LINE_1A],
            [_read(emitted[key], ACTIVITY_BOX) for key in COPY_KEYS])

    def test_one_pooled_answer_on_every_column_and_every_copy(self):
        scenario = _three_rentals_and_a_business()
        self.assertIs(mid_quarter_applies(scenario), True)
        for column in range(3):
            with self.subTest(alone=ADDRESSES[column]):
                self.assertIs(mid_quarter_applies(dataclasses.replace(
                    scenario, schedule_c_businesses=[],
                    rental_properties=[scenario.rental_properties[column]])),
                    False)
        emitted = self._emit(scenario)
        self.assertEqual(
            [_row_19b(emitted[key]) for key in self.COPIES],
            [("9000", "MQ", "2250"), ("7000", "MQ", "1850"),
             ("5000", "MQ", "1050"), ("8000", "MQ", "400")])
        self.assertEqual(_line_18s(emitted), ["2250", "1850", "1050"])
        self.assertEqual(_read(emitted["sch_c_1"], SCH_C_LINE_13), "400")

    def test_without_the_business_every_column_prints_half_year(self):
        """Twin: the same three rentals with no business are under the 40%
        line (6,000 of 21,000), so the mid-quarter rows above are reachable
        only through the pooled answer."""
        scenario = _scenario(*_THREE_RENTALS)
        self.assertIs(mid_quarter_applies(scenario), False)
        emitted = self._emit(scenario)
        self.assertEqual(
            [_row_19b(emitted[key]) for key in COPY_KEYS],
            [("9000", "HY", "1800"), ("7000", "HY", "1400"),
             ("5000", "HY", "1000")])
        self.assertEqual(_line_18s(emitted), ["1800", "1400", "1000"])


class MultiRentalAuditWorkbookTests(multi._Case):
    """The depreciation audit workbook on a return with three rentals:
    column A holds a prior-year building (no copy), columns B and C place
    equipment."""

    RENTAL_LABELS = tuple(
        f"rental property #{index} ('{address}')"
        for index, address in enumerate(ADDRESSES))

    def test_one_tab_pair_per_rental_in_column_order(self):
        _emitted, workbook = self.emit(_workbook_rentals())
        self.assertEqual(workbook.sheetnames, [
            "Tie-outs", "Mid-quarter test",
            "Assets rental_0", "Years rental_0",
            "Assets rental_1", "Years rental_1",
            "Assets rental_2", "Years rental_2"])
        # A tab's formulas stay on its own pair of sheets. The current year
        # is the building's seventh (2019-2025: row 8 of its year-by-year
        # tab) and the equipment's first (row 2).
        for index, (description, deduction, year_row) in enumerate((
                ("Rental building", 7_272, 8), ("Equipment", 1_800, 2),
                ("Equipment", 200, 2))):
            with self.subTest(rental=index):
                sheet = workbook[f"Assets rental_{index}"]
                self.assertEqual(sheet.max_row, 2)
                self.assertEqual(sheet["B2"].value, self.RENTAL_LABELS[index])
                self.assertEqual(sheet["C2"].value, description)
                self.assertEqual(
                    self.value(workbook, f"Assets rental_{index}", "K2"),
                    deduction)
                self.assertEqual(
                    sheet["H2"].value,
                    f"='Years rental_{index}'!D{year_row}")

    def test_three_rentals_and_a_business_get_four_tab_pairs(self):
        _emitted, workbook = self.emit(_three_rentals_and_a_business())
        self.assertEqual(workbook.sheetnames, [
            "Tie-outs", "Mid-quarter test",
            "Assets rental_0", "Years rental_0",
            "Assets rental_1", "Years rental_1",
            "Assets rental_2", "Years rental_2",
            "Assets sch_c_0", "Years sch_c_0"])
        # The pooled sheet lays the one answer out across all four.
        sheet = workbook[multi.POOLED]
        self.assertEqual(
            [[sheet[f"{c}{r}"].value for c in "ACD"] for r in range(2, 6)],
            [["rental_0", 300_000, 900_000], ["rental_1", 200_000, 700_000],
             ["rental_2", 100_000, 500_000], ["sch_c_0", 800_000, 800_000]])
        self.assertEqual(
            [self.value(workbook, multi.POOLED, c) for c in ("C6", "D6")],
            [1_400_000, 2_900_000])
        self.assertEqual(sheet["C7"].value, "=C6*100>40*D6")
        self.assertIs(self.value(workbook, multi.POOLED, "C7"), True)

    def test_tie_out_rows_per_copy_and_per_column(self):
        emitted, workbook = self.emit(_workbook_rentals())
        sheet = workbook[multi.TIE]
        self.assertEqual(self.claims(workbook), [
            "Form 4562 copy rental_1: Form 4562 line 19b, column (g): "
            "depreciation deduction",
            "Form 4562 copy rental_1: Form 4562 line 22: total depreciation",
            "Form 4562 copy rental_2: Form 4562 line 19b, column (g): "
            "depreciation deduction",
            "Form 4562 copy rental_2: Form 4562 line 22: total depreciation",
            "Schedule E line 18: depreciation, " + self.RENTAL_LABELS[0],
            "Schedule E line 18: depreciation, " + self.RENTAL_LABELS[1],
            "Schedule E line 18: depreciation, " + self.RENTAL_LABELS[2],
            "Whole return: rental_0 has no Form 4562 copy (no placement "
            "this year): Schedule E line 18: depreciation, "
            + self.RENTAL_LABELS[0],
            "Whole return: Form 4562 copies' line 22 plus no-copy "
            "activities' schedule lines = total depreciation",
        ])
        # Each copy's line 22 is computed from its own tab and set against
        # the value printed on THAT copy.
        self.assertEqual(sheet["B3"].value, "=SUM('Assets rental_1'!K2:K2)")
        self.assertEqual(sheet["B5"].value, "=SUM('Assets rental_2'!K2:K2)")
        self.assertEqual(sheet["C3"].value, multi._printed(
            emitted["f4562_rental_1"], multi.F4562_LINE_22))
        self.assertEqual(sheet["C5"].value, multi._printed(
            emitted["f4562_rental_2"], multi.F4562_LINE_22))
        self.assertEqual([sheet["C3"].value, sheet["C5"].value], [1_800, 200])
        # Each column's line 18 row carries what Schedule E printed there.
        self.assertEqual(
            [sheet[f"C{row}"].value for row in (6, 7, 8)],
            [multi._printed(emitted["sch_e"], path)
             for path in SCH_E_LINE_18])
        self.assertEqual(
            [sheet[f"C{row}"].value for row in (6, 7, 8)],
            [7_272, 1_800, 200])
        self.assertEqual(
            [self.value(workbook, multi.TIE, f"F{row}")
             for row in (2, 3, 4, 5, 6, 7, 8, 10)], ["PASS"] * 8)
        self.assertEqual(sheet.max_row, 10)

    def test_only_the_rental_without_a_copy_is_listed_as_lacking_one(self):
        _emitted, workbook = self.emit(_workbook_rentals())
        self.assertEqual(
            [claim for claim in self.claims(workbook)
             if "has no Form 4562 copy" in claim],
            ["Whole return: rental_0 has no Form 4562 copy (no placement "
             "this year): Schedule E line 18: depreciation, "
             + self.RENTAL_LABELS[0]])
        # Twin: with every rental placing property, none is listed.
        _emitted, workbook = self.emit(_scenario(*_THREE_RENTALS))
        self.assertEqual(
            [claim for claim in self.claims(workbook)
             if "has no Form 4562 copy" in claim], [])

    def test_whole_return_row_includes_the_no_copy_rentals_schedule_line(
            self):
        _emitted, resolved, printed = self.emit_with_handoff(
            _workbook_rentals())

        def whole_return_row(figures, name):
            workbook = openpyxl.load_workbook(dw.write_depreciation_audit(
                resolved, figures, self.tmp / name))
            sheet = workbook[multi.TIE]
            self.assertEqual(sheet.max_row, 10)
            self.assertEqual(sheet["B10"].value, "=B3+B5+C9")
            self.assertEqual(sheet["C10"].value, "=C7+C8+C9")
            return [self.value(workbook, multi.TIE, f"{c}10") for c in "BCEF"]

        # The two copies' line 22s (1,800 + 200) plus the no-copy rental's
        # Schedule E line 18 (7,272) is the return's total depreciation.
        self.assertEqual(
            whole_return_row(printed, "true.xlsx"),
            [9_272, 9_272, 0, "PASS"])
        # Twin: column B prints one dollar more than its copy's line 22.
        lines = dict(printed.activity_lines)
        lines[("rental_properties", 1)] += 1
        self.assertEqual(
            whole_return_row(
                dataclasses.replace(printed, activity_lines=lines),
                "off.xlsx"),
            [9_272, 9_273, -1, "FAIL"])

    def test_the_pooled_sheet_rows_every_rental_and_decides_once(self):
        _emitted, workbook = self.emit(_workbook_rentals())
        sheet = workbook[multi.POOLED]
        self.assertEqual(
            [[sheet[f"{c}{r}"].value for c in "ABCD"] for r in (2, 3, 4)], [
                ["rental_0", self.RENTAL_LABELS[0], 0, 0],
                ["rental_1", self.RENTAL_LABELS[1], 0, 900_000],
                ["rental_2", self.RENTAL_LABELS[2], 100_000, 100_000]])
        self.assertEqual(sheet["C6"].value, "=C5*100>40*D5")
        self.assertEqual(
            [self.value(workbook, multi.POOLED, c) for c in ("C5", "D5")],
            [100_000, 1_000_000])
        self.assertIs(self.value(workbook, multi.POOLED, "C6"), False)

    def test_the_hand_off_keys_copies_by_recon_key_and_lines_by_column(self):
        emitted, _resolved, printed = self.emit_with_handoff(
            _workbook_rentals())
        self.assertEqual(printed.f4562_copies, {
            "rental_1": dw.CopyFigures(
                line_17=None, line_19={"b": 1_800}, line_22=1_800),
            "rental_2": dw.CopyFigures(
                line_17=None, line_19={"b": 200}, line_22=200)})
        self.assertEqual(dict(printed.activity_lines), {
            ("rental_properties", 0): 7_272,
            ("rental_properties", 1): 1_800,
            ("rental_properties", 2): 200})
        self.assertEqual(
            [row.recon_key for row in printed.convention.activities],
            ["rental_0", "rental_1", "rental_2"])
        self.assertIs(printed.convention.mid_quarter, False)
        for key in ("rental_1", "rental_2"):
            with self.subTest(copy=key):
                self.assertEqual(
                    printed.f4562_copies[key].line_22,
                    multi._printed(emitted[f"f4562_{key}"], LINE_22))


if __name__ == "__main__":
    unittest.main()
