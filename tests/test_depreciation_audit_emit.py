"""The audit workbook as emitted by the orchestrator: beside the forms on
every asset-mode return, tied to the values on the filled PDFs.

PDF field paths are 2025 LITERALS (Form 4562 line 17, line 19c column (g),
line 22; Schedule E line 18, property A), as pinned in
tests/test_f4562_prior_asset_line17.py and tests/test_depreciation_emit_pins.py
-- not looked up through the mapping.
"""
import dataclasses
import tempfile
from datetime import date
import unittest
from pathlib import Path
from unittest import mock

import openpyxl
from pypdf import PdfReader

from tenforty import pdf_packet
from tenforty.__main__ import _assemble_packets_and_prune
from tenforty.audit import depreciation_workbook as dw
from tenforty.forms import f4562 as form_4562
from tenforty.models import DepreciableAsset
from tenforty.orchestrator import ReturnOrchestrator
from tests import _audit_workbook_fixtures as fx
from tests.helpers import SPREADSHEETS_DIR

F4562_LINE_17 = "topmostSubform[0].Page1[0].f1_25[0]"
F4562_LINE_19C_DEDUCTION = (
    "topmostSubform[0].Page1[0].SectionBTable[0].Line19c[0].f1_43[0]")
F4562_LINE_22 = "topmostSubform[0].Page2[0].f2_2[0]"
SCH_E_LINE_18 = "topmostSubform[0].Page1[0].Table_Expenses[0].Line18[0].f1_61[0]"
SCH_C_LINE_13 = "topmostSubform[0].Page1[0].Lines8-17[0].f1_22[0]"


def _printed(pdf_path, field_path) -> int:
    fields = PdfReader(str(pdf_path)).get_fields() or {}
    return int(str(fields[field_path].get("/V")).replace(",", ""))


class _EmitCase(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.tmp = Path(tmp.name)
        self.out = self.tmp / "out"
        self.orchestrator = ReturnOrchestrator(
            spreadsheets_dir=SPREADSHEETS_DIR, work_dir=self.tmp / "work")

    def emit(self, scenario) -> dict:
        _results, emitted = self.orchestrator.run_full_return(
            scenario, self.out)
        return emitted


class AssetModeEmitTests(_EmitCase):
    def setUp(self):
        super().setUp()
        self.emitted = self.emit(fx.scenario())
        self.workbook = openpyxl.load_workbook(
            self.emitted["depreciation_audit"])

    def test_asset_mode_return_emits_the_workbook(self):
        self.assertEqual(
            self.emitted["depreciation_audit"],
            self.out / "depreciation_audit_2025.xlsx")
        self.assertTrue(self.emitted["depreciation_audit"].is_file())
        self.assertEqual(
            self.workbook.sheetnames, ["Tie-outs", "Assets", "Year-by-year"])

    def test_the_forms_are_still_emitted_beside_it(self):
        for key in ("1040", "sch_e", "f4562"):
            with self.subTest(key=key):
                self.assertTrue(self.emitted[key].is_file())

    def test_printed_column_matches_the_filled_pdf(self):
        sheet = self.workbook["Tie-outs"]
        on_paper = [
            _printed(self.emitted["f4562"], F4562_LINE_17),
            _printed(self.emitted["f4562"], F4562_LINE_19C_DEDUCTION),
            _printed(self.emitted["f4562"], F4562_LINE_22),
            _printed(self.emitted["sch_e"], SCH_E_LINE_18),
        ]
        self.assertEqual(
            on_paper, [fx.LINE_17, fx.LINE_19C, fx.LINE_22, fx.SCH_E_LINE_18])
        self.assertEqual(
            [sheet[f"C{r}"].value for r in (2, 3, 4, 5)], on_paper)
        self.assertEqual(sheet.max_row, 5)

    def test_all_tie_outs_pass_under_the_mirror(self):
        self.assertEqual(
            [dw.cell_value(self.workbook, "Tie-outs", f"F{r}")
             for r in (2, 3, 4, 5)], ["PASS"] * 4)

    def test_the_key_is_standalone_review_matter(self):
        self.assertEqual(
            pdf_packet.classify_key("depreciation_audit"), "standalone")
        for key in self.emitted:
            with self.subTest(key=key):
                self.assertIsNotNone(pdf_packet.classify_key(key))

    def test_cli_prune_retains_the_workbook(self):
        path = self.emitted["depreciation_audit"]
        combined, retained = _assemble_packets_and_prune(
            self.emitted, self.out, fx.YEAR)
        self.assertIn(path, retained)
        self.assertTrue(path.is_file())
        self.assertFalse(self.emitted["f4562"].exists())
        self.assertNotIn(path, combined.values())
        # Still a readable workbook, and in no packet.
        self.assertEqual(
            openpyxl.load_workbook(path).sheetnames[0], "Tie-outs")


class ScheduleCEmitTests(_EmitCase):
    def test_schedule_c_line_13_ties_out(self):
        equipment = dataclasses.replace(fx.APPLIANCE, description="Equipment")
        emitted = self.emit(fx.scenario(
            rental_assets=(fx.BUILDING,), business_assets=(equipment,)))
        self.assertNotIn("f4562", emitted)
        workbook = openpyxl.load_workbook(emitted["depreciation_audit"])
        sheet = workbook["Tie-outs"]
        on_paper = [
            _printed(emitted["sch_e"], SCH_E_LINE_18),
            _printed(emitted["sch_c_1"], SCH_C_LINE_13)]
        self.assertEqual(on_paper, [7_272, 1_920])
        self.assertEqual([sheet[f"C{r}"].value for r in (2, 3)], on_paper)
        self.assertEqual(sheet.max_row, 3)
        self.assertEqual(
            [dw.cell_value(workbook, "Tie-outs", f"F{r}") for r in (2, 3)],
            ["PASS", "PASS"])


class OverrideEmitTests(_EmitCase):
    def test_overridden_activity_ties_to_the_printed_override(self):
        emitted = self.emit(fx.scenario(
            rental_assets=(fx.BUILDING, fx.APPLIANCE),
            rental_override=fx.override(9_000.0, 9_192.0)))
        workbook = openpyxl.load_workbook(emitted["depreciation_audit"])
        self.assertEqual(_printed(emitted["sch_e"], SCH_E_LINE_18), 9_000)
        self.assertEqual(workbook["Tie-outs"]["C2"].value, 9_000)
        self.assertEqual(
            [dw.cell_value(workbook, "Tie-outs", f"F{r}") for r in (2, 3)],
            ["PASS", "PASS"])


class NoWorkbookTests(_EmitCase):
    def assert_no_workbook(self, emitted):
        self.assertNotIn("depreciation_audit", emitted)
        self.assertEqual(list(self.out.glob("*.xlsx")), [])
        # The return itself emitted.
        self.assertTrue(emitted["1040"].is_file())

    def test_stated_mode_return_emits_no_workbook(self):
        self.assert_no_workbook(self.emit(
            fx.scenario(rental_assets=(), rental_stated=4_000.0)))

    def test_no_depreciation_return_emits_no_workbook(self):
        self.assert_no_workbook(self.emit(fx.scenario(rental_assets=())))


class WholeEmitRefusalTests(_EmitCase):
    def test_unbuildable_return_writes_nothing(self):
        fractional = dataclasses.replace(
            fx.APPLIANCE, prior_depreciation=9_500.5,
            acknowledges_prior_depreciation_as_stated=True)
        with self.assertRaisesRegex(
                NotImplementedError, "depreciation audit workbook"):
            self.emit(fx.scenario(rental_assets=(fractional,)))
        written = list(self.out.glob("*")) if self.out.exists() else []
        self.assertEqual(written, [])

    def test_the_whole_dollar_twin_emits(self):
        whole = dataclasses.replace(
            fx.APPLIANCE, prior_depreciation=9_500.0,
            acknowledges_prior_depreciation_as_stated=True)
        emitted = self.emit(fx.scenario(rental_assets=(whole,)))
        self.assertTrue(emitted["depreciation_audit"].is_file())

    def test_unwritable_output_refuses(self):
        with mock.patch.object(
                openpyxl.Workbook, "save",
                side_effect=PermissionError("read-only directory")):
            with self.assertRaisesRegex(
                    NotImplementedError,
                    r"depreciation_audit_2025\.xlsx could not be written"):
                self.emit(fx.scenario())
        self.assertEqual(list(self.out.glob("*.xlsx")), [])



_SB = "topmostSubform[0].Page1[0].SectionBTable[0]"
# Column (g) of the row a building placed this year prints on: the
# residential rental row is 19h through the 2024 form and 19i on the 2025
# form (which added a 50-year row above it). Literals, as pinned in
# tests/test_f4562_prior_asset_line17.py.
RESIDENTIAL_DEDUCTION = {
    2021: (f"{_SB}.Line19h_1[0].f1_66[0]", "h"),
    2024: (f"{_SB}.Line19h_1[0].f1_66[0]", "h"),
    2025: (f"{_SB}.Line19i_1[0].f1_79[0]", "i"),
}


class Line19LetteringByFormYearTests(_EmitCase):
    """The tie-out claim cites the row the deduction actually prints on."""

    def _emit_building(self, year):
        # January, 27.5-year, basis 200,000: 3.485% = 6,970.
        building = DepreciableAsset(
            description="Rental building",
            date_placed_in_service=date(year, 1, 15), basis=200_000.0,
            recovery_class="27.5-year")
        return self.emit(fx.scenario(rental_assets=(building,), year=year))

    def test_claim_cites_the_printed_row(self):
        for year, (field, letter) in RESIDENTIAL_DEDUCTION.items():
            with self.subTest(year=year):
                self.out = self.tmp / f"out_{year}"
                emitted = self._emit_building(year)
                self.assertEqual(_printed(emitted["f4562"], field), 6_970)
                workbook = openpyxl.load_workbook(
                    emitted["depreciation_audit"])
                tie = workbook["Tie-outs"]
                self.assertEqual(
                    [tie[f"A{r}"].value for r in (2, 3)], [
                        f"Form 4562 line 19{letter}, column (g): "
                        f"depreciation deduction",
                        "Form 4562 line 22: total depreciation"])
                self.assertEqual(
                    tie["D2"].value,
                    f"Form 4562 line 19{letter} column (g) as printed")
                self.assertEqual(tie["C2"].value, 6_970)
                self.assertEqual(
                    workbook["Assets"]["M2"].value, f"line 19{letter}")
                self.assertEqual(
                    tie["B2"].value,
                    f'=SUMIF(Assets!M2:M2,"line 19{letter}",Assets!K2:K2)')
                self.assertEqual(
                    [dw.cell_value(workbook, "Tie-outs", f"F{r}")
                     for r in (2, 3, 4)], ["PASS"] * 3)


class PrintedRowLabelTests(unittest.TestCase):
    def test_2025_form_letters(self):
        self.assertEqual(
            [form_4562.printed_row_label(c, 2025) for c in (
                "3-year", "5-year", "7-year", "10-year", "15-year",
                "20-year", "27.5-year", "39-year")],
            ["a", "b", "c", "d", "e", "f", "i", "j"])

    def test_through_2024_there_is_no_19j(self):
        for year in (2021, 2022, 2023, 2024):
            with self.subTest(year=year):
                self.assertEqual(
                    [form_4562.printed_row_label(c, year) for c in (
                        "3-year", "5-year", "7-year", "10-year", "15-year",
                        "20-year", "27.5-year", "39-year")],
                    ["a", "b", "c", "d", "e", "f", "h", "i"])


if __name__ == "__main__":
    unittest.main()
