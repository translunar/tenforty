"""The 481(a) sheet: present only with a `form_3115:` block.

Claimed (stated on the Form 3115 statement) against allowable (the table
reconstruction through the year before the change, from Year-by-year).
tenforty does not compute or check the section 481(a) adjustment, so nothing
here refuses: an unmatched statement row renders input-only, and a stated
adjustment that disagrees shows FAIL.

Hand figures (tests/_audit_workbook_fixtures.py): allowable through 2024 is
43,330 for the building and 5,200 for the appliance.
  building   claimed 30,000 - 43,330 = -13,330
  appliance  claimed  5,900 -  5,200 =     700
  net                                  -12,630
"""
import dataclasses
import tempfile
import unittest
from datetime import date
from pathlib import Path

import openpyxl

from tenforty.audit import depreciation_workbook as dw
from tenforty.forms.depreciation import resolver
from tenforty.models import Form3115, Form3115Asset
from tenforty.orchestrator import ReturnOrchestrator
from tests import _audit_workbook_fixtures as fx
from tests import _f3115_fixtures as f3115_fx
from tests.helpers import SPREADSHEETS_DIR
from tests.test_depreciation_audit_workbook import true_printed

NET = -12_630.0
STATEMENT = (
    ("Rental building", date(2019, 1, 15), 30_000.0),
    ("Appliance set", date(2023, 6, 1), 5_900.0),
    # On the statement, but no such asset among the depreciable assets.
    ("Unlisted shed", date(2018, 5, 1), 1_200.0),
)


def figures(line_26=NET, assets=STATEMENT) -> dw.Form3115Figures:
    return dw.Form3115Figures(
        line_26=line_26, year_of_change=fx.YEAR, assets=tuple(assets))


class _Case(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.tmp = Path(tmp.name)

    def write(self, form_3115=None):
        path = dw.write_depreciation_audit(
            resolver.audit_trail(fx.scenario()),
            true_printed(form_3115=form_3115), self.tmp / "audit.xlsx")
        return openpyxl.load_workbook(path)


class SheetPresenceTests(_Case):
    def test_absent_without_a_form_3115(self):
        self.assertEqual(
            self.write().sheetnames, ["Tie-outs", "Assets", "Year-by-year"])

    def test_last_sheet_with_a_form_3115(self):
        self.assertEqual(
            self.write(figures()).sheetnames,
            ["Tie-outs", "Assets", "Year-by-year", "481(a)"])


class Section481aSheetTests(_Case):
    def setUp(self):
        super().setUp()
        self.workbook = self.write(figures())
        self.sheet = self.workbook["481(a)"]

    def test_header_row(self):
        self.assertEqual([c.value for c in self.sheet[1]], [
            "Asset", "In service", "Claimed", "Allowable", "Adjustment",
            "Note", "Provenance"])

    def test_one_row_per_statement_asset_with_the_claimed_input(self):
        self.assertEqual(
            [self.sheet[f"A{r}"].value for r in (2, 3, 4)],
            ["Rental building", "Appliance set", "Unlisted shed"])
        self.assertEqual(
            [self.sheet[f"C{r}"].value for r in (2, 3, 4)],
            [30_000, 5_900, 1_200])
        self.assertEqual(self.sheet["B3"].value.date(), date(2023, 6, 1))

    def test_claimed_provenance_names_the_real_field(self):
        self.assertEqual(
            [self.sheet[f"G{r}"].value for r in (2, 3, 4)], [
                f"Claimed: form_3115.assets[{n}]."
                f"depreciation_claimed_present_method" for n in (0, 1, 2)])

    def test_formula_strings(self):
        formulas = {
            cell.coordinate: cell.value
            for row in self.sheet.iter_rows() for cell in row
            if cell.data_type == "f"}
        self.assertEqual(formulas, {
            # Cumulative through 2024: the building's sixth year row, the
            # appliance's second.
            "D2": "='Year-by-year'!F7",
            "E2": "=C2-D2",
            "D3": "='Year-by-year'!F10",
            "E3": "=C3-D3",
            "E6": "=SUM(E2:E4)",
        })

    def test_mirror_gives_the_hand_figures(self):
        value = lambda c: dw.cell_value(self.workbook, "481(a)", c)
        self.assertEqual([value("D2"), value("D3")], [43_330, 5_200])
        self.assertEqual([value("E2"), value("E3")], [-13_330, 700])
        self.assertEqual(value("E6"), -12_630)

    def test_unmatched_row_is_input_only_with_a_visible_note(self):
        self.assertIsNone(self.sheet["D4"].value)
        self.assertIsNone(self.sheet["E4"].value)
        self.assertEqual(
            self.sheet["F4"].value, "no matching depreciable asset")
        self.assertIsNone(self.sheet["F2"].value)

    def test_net_and_stated_rows(self):
        self.assertEqual(
            self.sheet["A6"].value, "Net section 481(a) adjustment")
        self.assertEqual(
            self.sheet["A7"].value, "Stated section 481(a) adjustment")
        self.assertEqual(self.sheet["E7"].value, -12_630)
        self.assertEqual(
            self.sheet["G7"].value, "form_3115.section_481a_adjustment")
        self.assertIn(
            "the filer must include it in gross income",
            self.sheet["F7"].value)
        self.assertIn(
            "tenforty does not carry the section 481(a) adjustment onto the "
            "return's income", self.sheet["F7"].value)

    def test_tie_out_row(self):
        tie = self.workbook["Tie-outs"]
        self.assertEqual(tie.max_row, 6)
        self.assertEqual(
            tie["A6"].value,
            "Form 3115 line 26: net section 481(a) adjustment")
        self.assertEqual(tie["B6"].value, "='481(a)'!E6")
        self.assertEqual(tie["C6"].value, -12_630)
        self.assertEqual(tie["E6"].value, "=B6-C6")
        self.assertEqual(tie["F6"].value, '=IF(ABS(E6)<0.005,"PASS","FAIL")')
        self.assertIn("Form 3115 line 26 as printed", tie["D6"].value)
        self.assertIn("form_3115.section_481a_adjustment", tie["D6"].value)
        self.assertEqual(
            dw.cell_value(self.workbook, "Tie-outs", "F6"), "PASS")

    def test_no_schedule_e_line_3_row(self):
        tie = self.workbook["Tie-outs"]
        claims = [tie[f"A{r}"].value for r in range(2, tie.max_row + 1)]
        self.assertEqual(len(claims), 5)
        self.assertEqual([c for c in claims if "line 3" in c], [])


class StatedFigureDisagreementTests(_Case):
    def test_a_disagreeing_stated_adjustment_shows_fail_and_still_emits(self):
        workbook = self.write(figures(line_26=31_250.0))
        self.assertEqual(workbook["Tie-outs"]["C6"].value, 31_250)
        self.assertEqual(
            dw.cell_value(workbook, "Tie-outs", "E6"), -43_880)
        self.assertEqual(dw.cell_value(workbook, "Tie-outs", "F6"), "FAIL")
        # The depreciation rows are untouched by it.
        self.assertEqual(
            [dw.cell_value(workbook, "Tie-outs", f"F{r}")
             for r in (2, 3, 4, 5)], ["PASS"] * 4)

    def test_no_matching_assets_at_all_still_emits(self):
        workbook = self.write(figures(assets=(STATEMENT[2],)))
        sheet = workbook["481(a)"]
        self.assertEqual(sheet["F2"].value, "no matching depreciable asset")
        self.assertEqual(dw.cell_value(workbook, "481(a)", "E4"), 0)
        self.assertEqual(dw.cell_value(workbook, "Tie-outs", "F6"), "FAIL")

    def test_same_description_different_date_does_not_match(self):
        workbook = self.write(figures(assets=(
            ("Rental building", date(2019, 2, 15), 30_000.0),)))
        self.assertEqual(
            workbook["481(a)"]["F2"].value, "no matching depreciable asset")
        self.assertIsNone(workbook["481(a)"]["D2"].value)

    def test_asset_placed_in_the_year_of_change_has_nothing_allowable(self):
        workbook = self.write(figures(assets=(
            ("Furniture", date(2025, 3, 10), 0.0),), line_26=0.0))
        sheet = workbook["481(a)"]
        self.assertEqual(sheet["D2"].value, 0)
        self.assertEqual(sheet["D2"].data_type, "n")
        self.assertEqual(
            sheet["F2"].value,
            "no depreciation allowable before 2025: placed in service in 2025")
        self.assertEqual(dw.cell_value(workbook, "Tie-outs", "F6"), "PASS")


class Form3115EmitTests(unittest.TestCase):
    """Through the orchestrator: line 26 as the Form 3115 prints it."""

    def test_line_26_and_claimed_totals_reach_the_workbook(self):
        block = f3115_fx.base_block(section_481a_adjustment=-12_630.0)
        statement = [
            dict(f3115_fx.ASSET_BUILDING, description=d,
                 date_placed_in_service=p,
                 depreciation_claimed_present_method=c)
            for d, p, c in STATEMENT]
        form = Form3115(**{
            **block, "assets": [Form3115Asset(**a) for a in statement]})
        base = fx.scenario(form_3115=form)
        scenario = dataclasses.replace(base, config=dataclasses.replace(
            base.config, **f3115_fx.IDENTITY))
        with tempfile.TemporaryDirectory() as tmp:
            orchestrator = ReturnOrchestrator(
                spreadsheets_dir=SPREADSHEETS_DIR, work_dir=Path(tmp) / "work")
            _results, emitted = orchestrator.run_full_return(
                scenario, Path(tmp) / "out")
            self.assertIn("f3115", emitted)
            workbook = openpyxl.load_workbook(emitted["depreciation_audit"])
        self.assertEqual(
            workbook.sheetnames,
            ["Tie-outs", "Assets", "Year-by-year", "481(a)"])
        self.assertEqual(workbook["Tie-outs"]["C6"].value, -12_630)
        self.assertEqual(
            [workbook["481(a)"][f"C{r}"].value for r in (2, 3, 4)],
            [30_000, 5_900, 1_200])
        self.assertEqual(
            [dw.cell_value(workbook, "Tie-outs", f"F{r}")
             for r in (2, 3, 4, 5, 6)], ["PASS"] * 5)


if __name__ == "__main__":
    unittest.main()
