"""Ruling R-S7-6a: the workbook path answers EIC Step 2 question 3 (Form
4797) "No" exactly when the sheet asks — static and engine-stopped checks
only; nothing here recalculates (the oracle half is
tests/test_eic_form_4797_oracle.py).

Why a formula, not a bare "X": the Q3 "No" cell is in NoEICredit's (A1)
OR-chain unconditionally, so a literal "X" would deny every workbook-routed
EIC; the sheet only ever asks Q3 after Q2 ("investment income over the
limit") is answered Yes. Each year's layout differs, so the gate cell, the
Q3 cells and the A1 membership are pinned per year: a vendor relayout of
any of them fails here. All identities synthetic.
"""
import functools
import re
import shutil
import tempfile
import unittest
from pathlib import Path
from unittest import mock

import openpyxl

from tenforty.mappings.f1040 import (
    EIC_FORM_4797_NO_KEY, EIC_FORM_4797_QUESTION, F1040,
    eic_form_4797_no_answer)
from tenforty.models import FilingStatus
from tenforty.oracle.engine import SpreadsheetEngine
from tenforty.orchestrator import ReturnOrchestrator
from tests.helpers import SPREADSHEETS_DIR
from tests.test_eic_grid_gates import UNREACHABLE_ITEMS, _input_fields
from tests.test_eic_routing import _eic_possible_childless

YEARS = tuple(sorted(EIC_FORM_4797_QUESTION))

# Q2's "Yes" formula per year, as read from each template.
GATE_FORMULAS = {
    2025: '=IF(E53>N58,"X","")',
    2024: '=IF(E53>N58,"X","")',
    2023: '=IF(E52>N57,"X","")',
    2022: '=IF(E52>N57,"X","")',
    2021: '=IF(E53>N57,"X","")',
}


@functools.cache
def _template(year):
    return openpyxl.load_workbook(
        SPREADSHEETS_DIR / "federal" / str(year) / "1040.xlsx")


def _or_terms(formula: str) -> list[str]:
    """Top-level arguments of an '=OR(...)' formula."""
    assert formula.startswith("=OR(") and formula.endswith(")"), formula
    terms, depth, current = [], 0, ""
    for ch in formula[4:-1]:
        if ch == "," and depth == 0:
            terms.append(current)
            current = ""
            continue
        depth += {"(": 1, ")": -1}.get(ch, 0)
        current += ch
    return terms + [current]


class PerYearLayoutTests(unittest.TestCase):

    def test_years_match_the_workbook_path(self):
        self.assertEqual(set(YEARS), set(F1040.INPUTS))

    def test_the_q2_gate_cell(self):
        for year in YEARS:
            q = EIC_FORM_4797_QUESTION[year]
            with self.subTest(year=year):
                self.assertEqual(_template(year)[q.sheet][q.q2_yes].value,
                                 GATE_FORMULAS[year])

    def test_the_q3_cells(self):
        for year in YEARS:
            q = EIC_FORM_4797_QUESTION[year]
            ws = _template(year)[q.sheet]
            with self.subTest(year=year):
                # Both answers are blank inputs in the template ...
                self.assertIsNone(ws[q.q3_yes].value)
                self.assertIsNone(ws[q.q3_no].value)
                # ... on the Form 4797 question's row ...
                self.assertTrue(str(ws[f"D{q.q3_yes[1:]}"].value)
                                .startswith("Yes.  See Form 4797 filers"))
                # ... and the sheet asks Q3 exactly when Q2 is Yes.
                self.assertIn(
                    f'IF(AND({q.q2_yes}<>"",{q.q3_yes}=""),"Please answer '
                    f'question at Step 2, Question 3."', ws["AN1"].value)

    def test_the_q3_no_cell_is_in_no_eic_credit_unconditionally(self):
        # The fact the formula answer exists for: a bare "X" here denies.
        for year in YEARS:
            q = EIC_FORM_4797_QUESTION[year]
            with self.subTest(year=year):
                self.assertIn(f'{q.q3_no}<>""',
                              _or_terms(_template(year)[q.sheet]["A1"].value))


class TheAnswerTests(unittest.TestCase):

    def test_it_answers_no_only_when_q2_is_yes(self):
        for year in YEARS:
            gate = EIC_FORM_4797_QUESTION[year].q2_yes
            with self.subTest(year=year):
                self.assertEqual(eic_form_4797_no_answer(year),
                                 f'=IF({gate}<>"","X","")')

    def test_production_maps_it_to_the_q3_no_cell(self):
        for year in YEARS:
            q = EIC_FORM_4797_QUESTION[year]
            with self.subTest(year=year):
                self.assertEqual(F1040.get_inputs(year)[EIC_FORM_4797_NO_KEY],
                                 q.q3_no)
                self.assertEqual(F1040.SHEET_MAP[year][EIC_FORM_4797_NO_KEY],
                                 q.sheet)
                self.assertNotIn(q.q3_yes, F1040.get_inputs(year).values())

    def test_the_engine_writes_it_as_a_formula(self):
        for year in YEARS:
            q = EIC_FORM_4797_QUESTION[year]
            with self.subTest(year=year), \
                    tempfile.TemporaryDirectory() as tmp:
                copy = Path(tmp) / "1040.xlsx"
                shutil.copy2(SPREADSHEETS_DIR / "federal" / str(year)
                             / "1040.xlsx", copy)
                SpreadsheetEngine()._write_inputs(
                    copy, F1040.get_inputs(year), F1040.SHEET_MAP[year],
                    {EIC_FORM_4797_NO_KEY: eic_form_4797_no_answer(year)})
                written = openpyxl.load_workbook(copy)[q.sheet]
                self.assertEqual(written[q.q3_no].value,
                                 eic_form_4797_no_answer(year))
                self.assertEqual(written[q.q2_yes].value,
                                 GATE_FORMULAS[year])

    def test_the_workbook_path_hands_the_engine_the_answer(self):
        class Stop(Exception):
            pass
        for year in YEARS:
            seen = {}

            def record(**kwargs):
                seen.update(kwargs)
                raise Stop
            with self.subTest(year=year), \
                    tempfile.TemporaryDirectory() as tmp:
                orch = ReturnOrchestrator(spreadsheets_dir=SPREADSHEETS_DIR,
                                          work_dir=Path(tmp))
                eff, _ = orch._build_effective_scenario(
                    _eic_possible_childless(year, FilingStatus.SINGLE))
                with mock.patch.object(orch.engine, "compute", record), \
                        self.assertRaises(Stop):
                    orch._compute_1040_via_workbook(eff)
                self.assertEqual(seen["inputs"][EIC_FORM_4797_NO_KEY],
                                 eic_form_4797_no_answer(year))
                self.assertIs(seen["mapping"], F1040)


class Form4797UnreachablePremiseTests(unittest.TestCase):
    """The premise of answering "No" for every filer: tenforty cannot express
    a Form 4797. Red the day it can — then this answer must be revisited
    (a filer with a 4797 would need "Yes" and Worksheet 1)."""

    def test_no_input_can_produce_a_form_4797(self):
        pattern, pinned = UNREACHABLE_ITEMS["form_4797"]
        fields = {f"{cls}.{name}" for cls, names in _input_fields().items()
                  for name in names}
        self.assertEqual({f for f in fields if re.search(
            pattern, f.split(".", 1)[1], re.I)}, set(pinned))
        # The proof's own exemptions, reviewed for this answer: the K-1 gain
        # is governed by the acknowledges_no_section_1231_gain attestation
        # (nonzero gain + False refuses; True affirms zero gain).
        self.assertEqual(set(pinned), {
            "SCorpIncome.net_gain_loss_4797",
            "ScheduleK1.section_1231_gain",
            "TaxReturnConfig.acknowledges_no_section_1231_gain"})

    def test_no_form_4797_module_exists(self):
        root = Path(__file__).resolve().parents[1] / "tenforty"
        self.assertEqual(sorted(p.name for p in root.rglob("*4797*")), [])
