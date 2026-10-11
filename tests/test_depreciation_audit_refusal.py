"""`depreciation_audit_unbuildable`: a return whose audit workbook cannot be
built does not emit.

Two triggers. The MIRROR MISMATCH: a formula cell, evaluated in Python over
the cells as written, differs from the engine's figure for that cell (or
cannot be read at all). And an UNWRITABLE output file. Each has a firing
proof and a twin that emits.

The mirror mismatch has two scenario shapes, both a stated prior with cents
where the basis ceiling binds (the engine works in whole dollars of prior;
the workbook's formula reads the stated cell):
  - acknowledged (`acknowledges_prior_depreciation_as_stated`): the engine
    takes the stated prior rounded;
  - unacknowledged: the stated prior rounds to the table reconstruction, and
    the engine takes the reconstruction.

Two shapes that are NOT refusals, by ruling: a fully depreciated asset (it
renders with no rate), and a table the workbook has no rates for (there is
none: `EngineTableCoverageTests` enumerates every table the engine can read).
"""
import dataclasses
import tempfile
import time
import unittest
from datetime import date
from pathlib import Path
from unittest import mock

import openpyxl

from tenforty.attestations import enforce_scoped_refusals
from tenforty.audit import depreciation_workbook as dw
from tenforty.forms.depreciation import macrs, resolver
from tenforty.models import (
    PERSONAL_PROPERTY_CLASSES, REAL_PROPERTY_CLASSES, DepreciableAsset,
)
from tests import _audit_workbook_fixtures as fx

# The engine takes the acknowledged prior as whole dollars (9,501), leaving
# 499 of the 10,000 basis; the workbook's formula reads the stated cell
# (10,000 - 9,500.50 = 499.50, which rounds to 500).
FRACTIONAL_PRIOR = dataclasses.replace(
    fx.APPLIANCE, prior_depreciation=9_500.5,
    acknowledges_prior_depreciation_as_stated=True)
WHOLE_PRIOR = dataclasses.replace(FRACTIONAL_PRIOR, prior_depreciation=9_500.0)


class UnbuildableRefusalTests(unittest.TestCase):
    def test_mirror_mismatch_fires(self):
        scenario = fx.scenario(rental_assets=(FRACTIONAL_PRIOR,))
        # The engine itself has a figure: only the workbook is the problem.
        self.assertEqual(resolver.audit_trail(scenario)[0].engine_amount, 499)
        with self.assertRaises(NotImplementedError) as caught:
            enforce_scoped_refusals(scenario, "emit")
        message = str(caught.exception)
        self.assertIn("depreciation audit workbook", message)
        self.assertIn("asset 'Appliance set' on rental property #0", message)
        self.assertIn("Assets!K2", message)
        self.assertIn("gives 500", message)
        self.assertIn("engine's figure is 499", message)
        self.assertIn("`prior_depreciation`", message)

    def test_whole_dollar_prior_twin_is_silent(self):
        scenario = fx.scenario(rental_assets=(WHOLE_PRIOR,))
        self.assertEqual(resolver.audit_trail(scenario)[0].engine_amount, 500)
        enforce_scoped_refusals(scenario, "emit")
        self.assertEqual(dw.unbuildable_reasons(scenario), [])

    def test_default_fixture_is_silent(self):
        enforce_scoped_refusals(fx.scenario(), "emit")
        self.assertEqual(dw.unbuildable_reasons(fx.scenario()), [])

    def test_stated_mode_return_is_silent(self):
        scenario = fx.scenario(rental_assets=(), rental_stated=4_000.0)
        enforce_scoped_refusals(scenario, "emit")
        self.assertEqual(dw.unbuildable_reasons(scenario), [])

    def test_prior_above_basis_is_silent(self):
        over = dataclasses.replace(WHOLE_PRIOR, prior_depreciation=10_500.0)
        enforce_scoped_refusals(fx.scenario(rental_assets=(over,)), "emit")

    def test_fully_depreciated_asset_is_silent(self):
        old = DepreciableAsset(
            description="Old computer", date_placed_in_service=date(2015, 4, 1),
            basis=3_000.0, recovery_class="5-year", prior_depreciation=3_000.0,
            no_bonus_or_section_179_history=True, convention="half-year")
        enforce_scoped_refusals(fx.scenario(rental_assets=(old,)), "emit")

    def test_each_mismatched_cell_is_named(self):
        scenario = fx.scenario(
            rental_assets=(fx.BUILDING, FRACTIONAL_PRIOR))
        reasons = dw.unbuildable_reasons(scenario)
        # The asset's own cell, and the activity total that sums it.
        self.assertEqual(len(reasons), 2)
        self.assertIn("Assets!K3", reasons[0])
        self.assertIn("Tie-outs!B2", reasons[1])
        self.assertNotIn("Rental building", reasons[0])


class WriterRefusalTests(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.out = Path(tmp.name) / "audit.xlsx"

    def test_write_refuses_an_unbuildable_trail(self):
        trail = resolver.audit_trail(
            fx.scenario(rental_assets=(FRACTIONAL_PRIOR,)))
        with self.assertRaisesRegex(NotImplementedError, r"Assets!K2"):
            dw.write_depreciation_audit(trail, dw.PrintedFigures(), self.out)
        self.assertFalse(self.out.exists())

    def test_write_emits_the_twin(self):
        trail = resolver.audit_trail(
            fx.scenario(rental_assets=(WHOLE_PRIOR,)))
        dw.write_depreciation_audit(trail, dw.PrintedFigures(), self.out)
        self.assertTrue(self.out.is_file())

    def test_unwritable_output_refuses(self):
        trail = resolver.audit_trail(fx.scenario())
        with mock.patch.object(
                openpyxl.Workbook, "save",
                side_effect=PermissionError("read-only directory")):
            with self.assertRaises(NotImplementedError) as caught:
                dw.write_depreciation_audit(
                    trail, dw.PrintedFigures(), self.out)
        message = str(caught.exception)
        self.assertIn("depreciation audit workbook", message)
        self.assertIn("audit.xlsx could not be written", message)
        self.assertIn("read-only directory", message)
        self.assertFalse(self.out.exists())

    def test_a_missing_directory_refuses(self):
        # The real thing, unpatched: openpyxl cannot open the path.
        trail = resolver.audit_trail(fx.scenario())
        missing = self.out.parent / "no_such_dir" / "audit.xlsx"
        with self.assertRaisesRegex(
                NotImplementedError, r"audit\.xlsx could not be written"):
            dw.write_depreciation_audit(trail, dw.PrintedFigures(), missing)

    def test_an_existing_file_is_overwritten(self):
        self.out.write_bytes(b"stale")
        dw.write_depreciation_audit(
            resolver.audit_trail(fx.scenario()), dw.PrintedFigures(), self.out)
        self.assertEqual(
            openpyxl.load_workbook(self.out).sheetnames[0], "Tie-outs")


class MirrorIsFailClosedTests(unittest.TestCase):
    """A formula the Python mirror cannot read is a refusal, never a cell
    it skips -- including the cells that have no engine figure to compare
    (differences, verdicts, running totals)."""

    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.out = Path(tmp.name) / "audit.xlsx"
        self.trail = resolver.audit_trail(fx.scenario())

    def test_unrecognized_function_in_a_verdict_cell_refuses(self):
        # The verdict column has no engine figure; it is still evaluated.
        with mock.patch.object(
                dw, "VERDICT_FORMULA", '=IFERROR(E{row},"PASS")'):
            with self.assertRaises(NotImplementedError) as caught:
                dw.write_depreciation_audit(
                    self.trail, dw.PrintedFigures(), self.out)
        message = str(caught.exception)
        self.assertIn("Tie-outs!F2", message)
        self.assertIn("unsupported function IFERROR", message)
        self.assertFalse(self.out.exists())

    def test_unreadable_text_in_a_verdict_cell_refuses(self):
        with mock.patch.object(dw, "VERDICT_FORMULA", "=E{row}^2"):
            with self.assertRaisesRegex(
                    NotImplementedError, r"Tie-outs!F2.*cannot read formula"):
                dw.write_depreciation_audit(
                    self.trail, dw.PrintedFigures(), self.out)
        self.assertFalse(self.out.exists())

    def test_the_predicate_reports_it_before_any_form_is_prepared(self):
        with mock.patch.object(
                dw, "VERDICT_FORMULA", '=IFERROR(E{row},"PASS")'):
            reasons = dw.unbuildable_reasons(fx.scenario())
            with self.assertRaisesRegex(NotImplementedError, "IFERROR"):
                enforce_scoped_refusals(fx.scenario(), "emit")
        self.assertEqual(len(reasons), 1)
        self.assertIn("Tie-outs!F2", reasons[0])

    def test_every_formula_cell_is_evaluated(self):
        # The mirror walks the workbook, not a list of expected cells: its
        # count is the count of formula cells on the sheets.
        workbook, reasons = dw._build(self.trail, dw.PrintedFigures())
        on_sheets = sum(
            1 for sheet in workbook.worksheets for row in sheet.iter_rows()
            for cell in row if cell.data_type == "f")
        self.assertEqual(reasons, [])
        # Assets 3 x (H, K, L); Year-by-year 11 x (A, E, F); Tie-outs one
        # activity row x (B, E, F).
        self.assertEqual(on_sheets, 9 + 33 + 3)
        self.assertEqual(len(dw.formula_cells(workbook)), on_sheets)


class EngineTableCoverageTests(unittest.TestCase):
    """Every Pub 946 table the engine can read is one the workbook renders.
    (An engine table added without workbook support fails here.)"""

    def test_every_engine_table_has_a_grid(self):
        seen = set()
        for cls in REAL_PROPERTY_CLASSES:
            asset = DepreciableAsset(
                description="Probe", date_placed_in_service=date(2020, 6, 1),
                basis=1.0, recovery_class=cls)
            seen.add(macrs.table_identity(
                asset, return_year=2025, mid_quarter=False))
        for cls in PERSONAL_PROPERTY_CLASSES:
            stated = [("half-year", None)] + [
                ("mid-quarter", q) for q in (1, 2, 3, 4)]
            for convention, quarter in stated:
                month = 6 if quarter is None else quarter * 3 - 1
                asset = DepreciableAsset(
                    description="Probe",
                    date_placed_in_service=date(2020, month, 1), basis=1.0,
                    recovery_class=cls, convention=convention,
                    quarter=quarter)
                seen.add(macrs.table_identity(
                    asset, return_year=2025, mid_quarter=False))
            # Placed in the return year: computed, under either 40% answer.
            for mid_quarter in (False, True):
                for month in (2, 5, 8, 11):
                    asset = DepreciableAsset(
                        description="Probe",
                        date_placed_in_service=date(2025, month, 1),
                        basis=1.0, recovery_class=cls)
                    seen.add(macrs.table_identity(
                        asset, return_year=2025, mid_quarter=mid_quarter))
        self.assertEqual(
            seen, {"A-1", "A-2", "A-3", "A-4", "A-5", "A-6", "A-7a"})
        self.assertEqual(seen - dw.GRID_TABLES, set())
        self.assertEqual(dw.GRID_TABLES, seen)



class UnacknowledgedFractionalPriorTests(unittest.TestCase):
    """The second mismatch shape: no acknowledgment. Basis 8, five-year,
    placed 2021: the tables give 2 + 3 + 2 + 1 = 8 through 2024, so nothing
    is left for 2025 (table amount 1). A stated prior of 7.50 rounds to that
    8 and loads; the workbook's formula sees 8 - 7.50 = 0.50 and rounds it
    to 1."""

    ASSET = DepreciableAsset(
        description="Hand tool", date_placed_in_service=date(2021, 4, 5),
        basis=8.0, recovery_class="5-year", prior_depreciation=7.5,
        no_bonus_or_section_179_history=True, convention="half-year")

    def test_stated_cents_under_a_binding_ceiling_fire(self):
        scenario = fx.scenario(rental_assets=(self.ASSET,))
        [activity] = resolver.audit_trail(scenario)
        self.assertFalse(activity.assets[0].prior_acknowledged)
        self.assertEqual(activity.assets[0].current.table_amount, 1)
        self.assertEqual(activity.engine_amount, 0)
        with self.assertRaises(NotImplementedError) as caught:
            enforce_scoped_refusals(scenario, "emit")
        message = str(caught.exception)
        self.assertIn("asset 'Hand tool'", message)
        self.assertIn("Assets!K2 gives 1 where the engine's figure is 0",
                      message)
        self.assertIn("`prior_depreciation`", message)

    def test_whole_dollar_twin_is_silent(self):
        whole = dataclasses.replace(self.ASSET, prior_depreciation=8.0)
        enforce_scoped_refusals(fx.scenario(rental_assets=(whole,)), "emit")


class HalfDollarBasisTests(unittest.TestCase):
    """A basis ending in .50 on a fully recovered asset emits: the year rows
    floor at zero as the engine does. 3-year half-year, placed 2018, basis
    4,289.50: 1,430 + 1,907 + 635 + 318 = 4,290, the rounded basis. From 2022
    on the formula's basis less cumulative is -0.50, which must not round to
    -1."""

    def _asset(self, basis, prior):
        return DepreciableAsset(
            description="Recovered", date_placed_in_service=date(2018, 5, 1),
            basis=basis, recovery_class="3-year", prior_depreciation=prior,
            no_bonus_or_section_179_history=True, convention="half-year")

    def test_half_dollar_basis_emits(self):
        scenario = fx.scenario(
            rental_assets=(self._asset(4_289.50, 4_290.0),))
        self.assertEqual(dw.unbuildable_reasons(scenario), [])
        enforce_scoped_refusals(scenario, "emit")
        workbook, _ = dw._build(
            resolver.audit_trail(scenario), dw.PrintedFigures())
        cache = {}
        self.assertEqual(
            [dw.cell_value(workbook, "Year-by-year", f"E{r}", cache)
             for r in range(2, 10)],
            [1_430, 1_907, 635, 318, 0, 0, 0, 0])
        self.assertEqual(
            dw.cell_value(workbook, "Year-by-year", "F9", cache), 4_290)

    def test_whole_dollar_twin_emits(self):
        scenario = fx.scenario(
            rental_assets=(self._asset(4_289.00, 4_289.0),))
        self.assertEqual(dw.unbuildable_reasons(scenario), [])


class LongHistoryTests(unittest.TestCase):
    """The mirror's cost is linear in year rows. (Evaluated without a cache,
    each cumulative cell re-evaluates the one above it twice, and 28 rows
    would not finish.)"""

    def setUp(self):
        self.scenario = fx.scenario(rental_assets=(fx.OLD_BUILDING,))

    def test_predicate_is_fast_on_a_28_row_asset(self):
        start = time.perf_counter()
        reasons = dw.unbuildable_reasons(self.scenario)
        elapsed = time.perf_counter() - start
        self.assertEqual(reasons, [])
        self.assertLess(elapsed, 1.0)

    def test_writer_is_fast_on_a_28_row_asset(self):
        with tempfile.TemporaryDirectory() as tmp:
            start = time.perf_counter()
            dw.write_depreciation_audit(
                resolver.audit_trail(self.scenario), dw.PrintedFigures(),
                Path(tmp) / "audit.xlsx")
            elapsed = time.perf_counter() - start
        self.assertLess(elapsed, 1.0)

    def test_32_row_nonresidential_schedule_is_fast(self):
        # 39-year, placed 1994: 32 year rows.
        office = DepreciableAsset(
            description="Office", date_placed_in_service=date(1994, 7, 1),
            basis=390_000.0, recovery_class="39-year",
            prior_depreciation=0.0,
            acknowledges_prior_depreciation_as_stated=True)
        start = time.perf_counter()
        reasons = dw.unbuildable_reasons(
            fx.scenario(rental_assets=(office,)))
        elapsed = time.perf_counter() - start
        self.assertEqual(reasons, [])
        self.assertLess(elapsed, 1.0)


class FailClosedOn481aTests(unittest.TestCase):
    """The fail-closed walk reaches the 481(a) sheet."""

    FIGURES = dw.Form3115Figures(
        line_26=0.0, year_of_change=fx.YEAR,
        assets=(("Rental building", date(2019, 1, 15), 30_000.0),))

    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.out = Path(tmp.name) / "audit.xlsx"
        self.trail = resolver.audit_trail(fx.scenario())
        self.printed = dw.PrintedFigures(form_3115=self.FIGURES)

    def test_every_481a_formula_cell_is_walked(self):
        workbook, reasons = dw._build(self.trail, self.printed)
        self.assertEqual(reasons, [])
        walked = dw.formula_cells(workbook)
        self.assertEqual(
            [c for s, c in walked if s == "481(a)"], ["D2", "E2", "E4"])
        # Tie-outs: the activity row and the line 26 row, x (B, E, F).
        self.assertEqual(
            [c for s, c in walked if s == "Tie-outs"],
            ["B2", "E2", "F2", "B3", "E3", "F3"])

    def test_unreadable_formula_on_the_481a_sheet_refuses(self):
        with mock.patch.object(
                dw, "NET_481A_FORMULA", "=SUMPRODUCT(E2:E{last})"):
            with self.assertRaises(NotImplementedError) as caught:
                dw.write_depreciation_audit(
                    self.trail, self.printed, self.out)
        message = str(caught.exception)
        self.assertIn("481(a)!E4", message)
        self.assertIn("unsupported function SUMPRODUCT", message)
        self.assertFalse(self.out.exists())

    def test_the_sheet_as_written_is_readable(self):
        dw.write_depreciation_audit(self.trail, self.printed, self.out)
        self.assertTrue(self.out.is_file())


class MirrorToleranceTests(unittest.TestCase):
    """The mismatch threshold is half a cent, proved below a dollar (every
    scenario-driven mismatch differs by a whole dollar, which any threshold
    under a dollar would catch)."""

    def setUp(self):
        self.workbook = openpyxl.Workbook()
        sheet = self.workbook.active
        sheet.title = "S"
        sheet["A1"] = 100.004
        sheet["A2"] = 100.006
        sheet["B1"] = "=A1"
        sheet["B2"] = "=A2"
        sheet["B3"] = '="text"'
        sheet["B4"] = "=A1<A2"

    def _reasons(self, coordinate, expected):
        return dw._mirror_mismatches(self.workbook, [dw._Expectation(
            "S", coordinate, expected, "subject", "`field`")])

    def test_under_half_a_cent_is_not_a_mismatch(self):
        self.assertEqual(self._reasons("B1", 100), [])

    def test_over_half_a_cent_is_a_mismatch(self):
        [reason] = self._reasons("B2", 100)
        self.assertIn("S!B2 gives 100.006 where the engine's figure is 100",
                      reason)

    def test_text_where_a_figure_is_expected_is_a_mismatch(self):
        [reason] = self._reasons("B3", 0)
        self.assertIn("S!B3 gives text", reason)

    def test_true_false_where_a_figure_is_expected_is_a_mismatch(self):
        # True == 1 in Python; it is still not the engine's figure 1.
        [reason] = self._reasons("B4", 1)
        self.assertIn("S!B4 gives True", reason)

    def test_comparing_text_with_a_number_cannot_be_checked(self):
        self.workbook["S"]["B5"] = '=IF("a"<1,1,2)'
        [reason] = dw._mirror_mismatches(self.workbook, [])
        self.assertIn("S!B5 cannot be checked", reason)

    def test_a_formula_that_refers_to_itself_cannot_be_checked(self):
        self.workbook["S"]["B6"] = "=B6+1"
        [reason] = dw._mirror_mismatches(self.workbook, [])
        self.assertIn("S!B6 cannot be checked", reason)


if __name__ == "__main__":
    unittest.main()
